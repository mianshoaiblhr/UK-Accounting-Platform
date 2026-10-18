#!/usr/bin/env python3
"""Writes docs/architecture/v2-m3a-golden-workings.xlsx from tests/golden/v2-m3a-adjustments.json: every journal, release, provision balance,
reconciliation and P&L figure of the V2-M3a golden scenario recomputed here with Python's Decimal and datetime, independently of the platform,
next to the hand-written expected figure with a check column. For the named accountant's review (DEC-054). Exits 1 if any figure differs."""
import calendar, datetime as dt, json, pathlib, sys
from decimal import Decimal as D, ROUND_HALF_UP

root = pathlib.Path(__file__).resolve().parents[2]
G = json.loads((root / "tests/golden/v2-m3a-adjustments.json").read_text())
day = dt.date.fromisoformat
C2 = lambda x: x.quantize(D("0.01"), rounding=ROUND_HALF_UP)
ledger = []                                        # (date, account, debit - credit)
for p in G["seed"]:
    for code, amt in p["lines"]:
        a = D(amt); d = day(p["date"])
        ledger += [(d, "1200", a), (d, code, -a)] if p["direction"] == "RECEIPT" else [(d, code, a), (d, "1200", -a)]
def post(d, lines): ledger.extend((d, k, D(dr) - D(cr)) for k, dr, cr in lines)
def months_counted(start, d):
    if d < start: return 0
    span = (d.year - start.year) * 12 + d.month - start.month
    return span + 1 if d.day == calendar.monthrange(d.year, d.month)[1] else span
def due(i, d):
    s, e, amt = day(i["periodStart"]), day(i["periodEnd"]), D(i["amount"])
    if d < s: return D(0)
    if i["convention"] == "DAILY":
        total = (e - s).days + 1; return amt * D(min((d - s).days + 1, total)) / D(total)
    total = (e.year - s.year) * 12 + e.month - s.month + 1
    return amt * D(min(months_counted(s, d), total)) / D(total)
rows, bad = [], 0
def check(step, item, got, exp):
    global bad
    ok = D(str(got)) == D(str(exp)); bad += not ok
    rows.append([step, item, float(D(str(got))), float(D(str(exp))), "OK" if ok else "DIFFERS"])
def jcheck(step, got, exp):
    g = sorted((k, str(C2(D(dr))), str(C2(D(cr)))) for k, dr, cr in got); e = sorted((k, dr, cr) for k, dr, cr in exp)
    global bad
    ok = g == e; bad += not ok; rows.append([step, "journal lines", "; ".join(f"{k} {dr}/{cr}" for k, dr, cr in g), "; ".join(f"{k} {dr}/{cr}" for k, dr, cr in e), "OK" if ok else "DIFFERS"])
items = {i["key"]: i for i in G["items"]}
for i in G["items"]:
    a, pl, bal, d = D(i["amount"]), i["plAccount"], i["balanceAccount"], day(i["date"])
    if i["kind"] in ("ACCRUAL", "ACCRUED_INCOME"):
        j = [(pl, a, 0), (bal, 0, a)] if i["kind"] == "ACCRUAL" else [(bal, a, 0), (pl, 0, a)]
        post(d, j); jcheck(f"{i['key']} {i['kind'].lower()}", j, i["expect"]["journal"])
        rj = [(k, cr, dr) for k, dr, cr in j]; post(d + dt.timedelta(days=1), rj); jcheck(f"{i['key']} reversal {d + dt.timedelta(days=1)}", rj, i["expect"]["reversal"]["journal"])
    elif i["setup"] == "POSTED":
        j = [(bal, a, 0), (pl, 0, a)] if i["kind"] == "PREPAYMENT" else [(pl, a, 0), (bal, 0, a)]
        post(d, j); jcheck(f"{i['key']} set-up", j, i["expect"]["journal"])
released = {k: D(0) for k in items}
pv = G["provision"]; pbal = []
moves = sorted(pv["movements"], key=lambda m: m["date"])
def provision_until(d):
    for m in moves:
        if day(m["date"]) <= d and m not in pbal:
            a = D(m["amount"])
            j = [(pv["expenseAccount"], a, 0), (pv["account"], 0, a)] if m["kind"] == "MADE" else [(pv["account"], a, 0), (m.get("offsetAccount") or pv["expenseAccount"], 0, a)]
            post(day(m["date"]), j); jcheck(f"Provision {m['kind'].lower()} {m['date']}", j, m["journal"]); pbal.append(m)
for r in G["runs"]:
    d = day(r["asAt"]); provision_until(d); js = []
    for k, i in items.items():
        if i["kind"] not in ("PREPAYMENT", "DEFERRED_INCOME"): continue
        ch = C2(due(i, d)) - released[k]
        if ch > 0:
            check(f"Release {r['asAt']}", f"{k}", ch, r["lines"].get(k, "0")); released[k] += ch
            js += [(i["plAccount"], ch, 0), (i["balanceAccount"], 0, ch)] if i["kind"] == "PREPAYMENT" else [(i["balanceAccount"], ch, 0), (i["plAccount"], 0, ch)]
    post(d, js); jcheck(f"Release {r['asAt']}", js, r["journal"]); check(f"Release {r['asAt']}", "total", sum(D(x[1]) for x in js if x[1]), r["total"])
    for code, exp in G["reconciliation"][r["asAt"]].items():
        typ = "ASSET" if code in ("1110", "1120") else "LIABILITY"
        led = sum(v for (dd, k, v) in ledger if k == code and dd <= d) * (1 if typ == "ASSET" else -1)
        check(f"Reconciliation {r['asAt']}", f"{code} ledger", led, exp)
for rep in pv["reports"]:
    f, t = day(rep["from"]), day(rep["to"]); live = [m for m in moves]
    op = sum((D(m["amount"]) if m["kind"] == "MADE" else -D(m["amount"])) for m in live if day(m["date"]) < f)
    s = lambda k: sum(D(m["amount"]) for m in live if m["kind"] == k and f <= day(m["date"]) <= t)
    for name, v in [("opening", op), ("made", s("MADE")), ("used", s("USED")), ("released", s("RELEASED")), ("closing", op + s("MADE") - s("USED") - s("RELEASED"))]: check(f"Provision report {rep['from']}..{rep['to']}", name, v, rep[name])
for label, (f, t) in {"2026": ("2026-01-01", "2026-12-31"), "2027Q1": ("2027-01-01", "2027-03-31")}.items():
    for code, exp in G["pl"][label].items(): check(f"P&L {label}", code, sum(v for (dd, k, v) in ledger if k == code and day(f) <= dd <= day(t)), exp)
try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:
    print("openpyxl missing"); sys.exit(1)
wb = Workbook(); ws = wb.active; ws.title = "Read me"
for line in ["V2-M3a golden year-end adjustments - workings (SYNTHETIC)", "", G["purpose"], "", G["disclaimer"], "",
             "Every figure is recomputed by this workbook's generator with Python Decimal and calendar dates, independently of the platform, and compared with the hand-written expected figure. The platform's API test matches the same figures to the cent.",
             "", "Why: " + G["runs"][0]["why"], G["runs"][1]["why"], G["reconciliation"]["why"], G["pl"]["why"]]: ws.append([line])
ws.column_dimensions["A"].width = 150; ws["A1"].font = Font(bold=True, size=14)
for r in ws.iter_rows():
    for c in r: c.alignment = Alignment(wrap_text=True, vertical="top")
s3 = wb.create_sheet("Checks"); s3.append(["Step", "Figure", "Recomputed", "Expected (hand)", "Check"])
for r in rows: s3.append(r)
for i, w in enumerate([34, 20, 50, 50, 10], 1): s3.column_dimensions["ABCDE"[i - 1]].width = w
for c in s3[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3864")
for row in s3.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if isinstance(c.value, float): c.number_format = '#,##0.00;[Red]-#,##0.00'
    row[4].fill = PatternFill("solid", fgColor="C6EFCE" if row[4].value == "OK" else "FFC7CE")
out = root / "docs/architecture/v2-m3a-golden-workings.xlsx"; wb.save(out)
print(f"wrote {out.relative_to(root)}; {len(rows)} figures; differences: {bad}")
sys.exit(1 if bad else 0)
