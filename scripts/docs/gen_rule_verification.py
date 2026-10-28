#!/usr/bin/env python3
"""Writes docs/legal/rule-verification-worksheet.xlsx: what the named verifier (DEC-056) checks in the platform before real use (DEC-083):
the VAT codes and rate versions a practice enters, the draft FRS 105 / FRS 102 Section 1A packs, the size thresholds and the draft deadline rules (V7-M1). Every figure and
box mapping here is a PROPOSAL to be checked against the provision named; nothing in this workbook is verified, and neither engineering
nor AI verifies. Usage: python3 scripts/docs/gen_rule_verification.py (needs openpyxl)."""
import json
import pathlib
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

root = pathlib.Path(__file__).resolve().parents[2]
OUT = root / "docs/legal/rule-verification-worksheet.xlsx"
LEG = "https://www.legislation.gov.uk/ukpga/1994/23/"

# (code, name, treatment, direction, reverse charge, rate %, net boxes, VAT boxes, provision to check, source link, what to check)
VAT_CODES = [
    ("S20", "Standard-rated sales", "STANDARD", "SALES", "no", "20", "6", "1", "VATA 1994 s.2(1)", LEG + "section/2", "The standard rate and the date it has applied from"),
    ("R5", "Reduced-rate sales", "REDUCED", "SALES", "no", "5", "6", "1", "VATA 1994 s.29A and Sch 7A", LEG + "section/29A", "The reduced rate and which supplies the firm's clients make at it"),
    ("Z0", "Zero-rated sales", "ZERO", "SALES", "no", "0", "6", "", "VATA 1994 s.30 and Sch 8", LEG + "section/30", "Zero-rated supplies go in box 6 with no VAT"),
    ("EX", "Exempt sales", "EXEMPT", "SALES", "no", "", "6", "", "VATA 1994 s.31 and Sch 9", LEG + "section/31", "Exempt supplies are included in box 6 (Notice 700/12)"),
    ("OS", "Outside the scope of VAT (sales)", "OUTSIDE_SCOPE", "SALES", "no", "", "", "", "VATA 1994 ss.1 and 4", LEG + "section/4", "Outside-scope income is left out of every box"),
    ("S20P", "Standard-rated purchases", "STANDARD", "PURCHASES", "no", "20", "7", "4", "VATA 1994 ss.24-26", LEG + "section/26", "Input tax recoverable; blocked items (cars, entertainment) use their own code"),
    ("R5P", "Reduced-rate purchases", "REDUCED", "PURCHASES", "no", "5", "7", "4", "VATA 1994 s.29A and ss.24-26", LEG + "section/29A", "As S20P at the reduced rate"),
    ("Z0P", "Zero-rated purchases", "ZERO", "PURCHASES", "no", "0", "7", "", "VATA 1994 s.30", LEG + "section/30", "Zero-rated purchases go in box 7"),
    ("EXP", "Exempt purchases", "EXEMPT", "PURCHASES", "no", "", "7", "", "VATA 1994 s.31", LEG + "section/31", "Exempt purchases are included in box 7 (Notice 700/12)"),
    ("OSP", "Outside the scope of VAT (purchases)", "OUTSIDE_SCOPE", "PURCHASES", "no", "", "", "", "VATA 1994 ss.1 and 4", LEG + "section/4", "Wages, drawings and the like are left out of every box"),
    ("BLK", "Purchases with input tax blocked", "STANDARD", "PURCHASES", "no", "0", "7", "", "Value Added Tax (Input Tax) Order 1992, SI 1992/3222, arts 5 and 7", "https://www.legislation.gov.uk/uksi/1992/3222/article/5",
     "Business entertainment and cars: net value in box 7, no VAT reclaimed. Check that a 0% STANDARD code is the right way to hold it"),
    ("RC20", "Reverse charge: services from abroad", "STANDARD", "PURCHASES", "yes", "20", "6, 7", "1, 4", "VATA 1994 s.8; VAT Notice 741A", LEG + "section/8",
     "The customer accounts for output and input VAT; net value in boxes 6 and 7"),
    ("DRC20", "Domestic reverse charge: construction services", "STANDARD", "PURCHASES", "yes", "20", "7", "1, 4", "VATA 1994 s.55A; SI 2019/892", LEG + "section/55A",
     "The customer's net value goes in box 7 only, not box 6. Check before use"),
]
STEPS = [
    ("1", "Staging is running (DEC-085), with the firm's organisation and a synthetic company.", "Person with AWS access", ""),
    ("2", "Enter each VAT code on the 'VAT codes' sheet and its rate version, with the source URL, provision and the date it was read.", "A person with vat:manage who is NOT the verifier", "One set per practice: codes and rates are held by the organisation, not by each client"),
    ("3", "Open each rate version, check the rate, the dates and the boxes against the provision, and verify it with notes on what was checked.", "Shoaib, with tax:rule-verify", "The platform refuses a verifier who entered the rate (DEC-021)"),
    ("4", "Import content/standards/drafts/frs105-micro.json, frs102-1a-small.json and uk-size-thresholds.json with scripts/standards/import-pack.mjs.", "A person with accounts:rules-manage who is NOT the verifier", "Each lands UNVERIFIED"),
    ("5", "Check every line, map entry, rule and threshold against the standard and the Companies Act provisions, using docs/architecture/v2-draft-wording-review.xlsx; correct or remove rules, then verify each version.", "Shoaib, with accounts:rule-verify", "The platform refuses a verifier who imported the version (DEC-052)"),
    ("6", "Import content/compliance/drafts/uk-deadline-rules.json with scripts/compliance/import-deadline-rules.mjs, then check each rule on the 'Deadline rules' sheet and verify it, or supersede it with a corrected version first.", "Import: a person with compliance:rules-manage; verify: Shoaib, with compliance:rule-verify", "The platform refuses a verifier who entered the rule (DEC-088)"),
    ("7", "Enter the presentation rule for boxes 6 to 9 on the 'VAT presentation' sheet (Settings > VAT presentation rules, or POST vat-presentation-rules), then check it against Notice 700/12 and verify it.", "Enter: a person with vat:manage; verify: Shoaib, with tax:rule-verify", "Without a verified rule no VAT return can be filed or recorded (DEC-091)"),
    ("8", "Import content/ixbrl/drafts/frc-frs105.json with scripts/ixbrl/import-taxonomy.mjs, then check the entry point, the namespaces and every entry on the 'iXBRL mapping' sheet against the FRC taxonomy files; correct entries while UNVERIFIED, then verify the version.", "Import: a person with ixbrl:rules-manage; verify: Shoaib, with ixbrl:rule-verify", "The platform refuses a verifier who created the version (DEC-093)"),
    ("9", "Import content/tax/drafts/uk-ct-rules.json with scripts/tax/import-ct-rules.mjs, then check each rule on the 'Corporation Tax rules' sheet against its provision and verify it, or supersede it with a corrected version first.", "Import: a person with ct:manage; verify: Shoaib, with tax:rule-verify", "The platform refuses a verifier who entered the rule; CT periods cannot be derived until the period rule is verified (DEC-097)"),
    ("10", "Fill in the 'Sign-off' sheet with the platform's record of each verification.", "Shoaib", "Evidence for the go-live checklist (DEC-085, item 3)"),
]
LAW = [
    ("Value Added Tax Act 1994, ss.1, 2, 4, 8, 24-26, 29A, 30, 31 and 55A", "The charge, the rates, input tax, reverse charges", LEG + "contents"),
    ("Value Added Tax Regulations 1995, SI 1995/2518, regs 25 and 25A", "Returns and their filing", "https://www.legislation.gov.uk/uksi/1995/2518/regulation/25A"),
    ("Value Added Tax (Input Tax) Order 1992, SI 1992/3222", "Input tax blocks: business entertainment, cars", "https://www.legislation.gov.uk/uksi/1992/3222/contents"),
    ("VAT Notice 700/12", "How each box of the return is filled", "https://www.gov.uk/guidance/how-to-fill-in-and-submit-your-vat-return-vat-notice-70012"),
    ("Making Tax Digital for VAT: VAT (MTD) API, submit VAT return", "The value of each box as HMRC's API receives it; boxes 6 to 9 in whole pounds", "https://developer.service.hmrc.gov.uk/api-documentation/docs/api/service/vat-api/1.0"),
    ("FRC Taxonomies (FRS 102 suite)", "The concepts, entry point and namespaces of tagged UK accounts", "https://www.frc.org.uk/accountants/frc-taxonomies/"),
    ("Companies House: filing accounts as iXBRL (software filing)", "What Companies House accepts electronically", "https://www.gov.uk/government/publications/technical-interface-specifications-for-companies-house-software"),
    ("Companies Act 2006, ss.441-444 and 1068", "Duty to file accounts; the registrar's power to require electronic form", "https://www.legislation.gov.uk/ukpga/2006/46/section/441"),
    ("VAT Notice 700 (the VAT guide)", "Rounding of VAT (the platform rounds each line half-up, DEC-025)", "https://www.gov.uk/guidance/vat-guide-notice-700"),
    ("Companies Act 2006, ss.382-384 and 384A-384B", "Small companies and micro-entities: the size tests", "https://www.legislation.gov.uk/ukpga/2006/46/section/382"),
    ("Companies Act 2006, ss.442-443 and 853A", "Deadlines: filing accounts, the month-end rule, the confirmation statement", "https://www.legislation.gov.uk/ukpga/2006/46/section/442"),
    ("Corporation Tax Act 2009, s.10", "When a CT accounting period ends: 12 months from its start at the latest", "https://www.legislation.gov.uk/ukpga/2009/4/section/10"),
    ("Corporation Tax Act 2010, Part 3A (ss.18A-18N), inserted by Finance Act 2021 s.7 and Sch 1", "The small profits rate, the lower and upper limits, associated companies and marginal relief", "https://www.legislation.gov.uk/ukpga/2010/4/part/3A"),
    ("Finance Act 2021, s.6", "The main rate of 25% for the financial year 2023", "https://www.legislation.gov.uk/ukpga/2021/26/section/6"),
    ("HMRC: Rates and allowances for Corporation Tax", "HMRC's summary of the rates and limits by financial year", "https://www.gov.uk/government/publications/rates-and-allowances-corporation-tax"),
    ("HMRC Company Taxation Manual CTM01405", "Accounting periods: a period of account longer than 12 months", "https://www.gov.uk/hmrc-internal-manuals/company-taxation-manual/ctm01405"),
    ("Taxes Management Act 1970, s.59D; Finance Act 1998, Sch 18 para 14", "Deadlines: Corporation Tax payment and the Company Tax Return", "https://www.legislation.gov.uk/ukpga/1970/9/section/59D"),
    ("Small Companies and Groups (Accounts and Directors' Report) Regulations 2008, SI 2008/409", "Formats and notes for small and micro-entity accounts", "https://www.legislation.gov.uk/uksi/2008/409/contents"),
    ("FRS 105 and FRS 102 Section 1A (Financial Reporting Council)", "The standards the draft packs follow", "https://www.frc.org.uk/accountants/accounting-and-reporting-policy/uk-accounting-standards"),
    ("Card Protection Plan Ltd v Customs and Excise Commissioners (C-349/96)", "Single or multiple supplies: decides which code a sale takes", "https://eur-lex.europa.eu/legal-content/EN/TXT/?uri=CELEX:61996CJ0349"),
    ("Fiscale eenheid Koninklijke Ahold BV v Staatssecretaris van Financien (C-484/06)", "Rounding of VAT is a matter for national rules", "https://eur-lex.europa.eu/legal-content/EN/TXT/?uri=CELEX:62006CJ0484"),
]

INK = "172231"
HEAD, NOTE_FILL = PatternFill("solid", fgColor=INK), PatternFill("solid", fgColor="FFFBEB")
FONT, BOLD = Font(name="Arial", size=10, color=INK), Font(name="Arial", size=10, bold=True, color="FFFFFF")
TITLE, NOTE = Font(name="Arial", size=14, bold=True, color=INK), Font(name="Arial", size=9, italic=True, color="566170")
LINK = Font(name="Arial", size=10, color="1D5FBF", underline="single")
RULE, WRAP = Border(bottom=Side(style="thin", color="E1E6EC")), Alignment(wrap_text=True, vertical="top")


def table(ws, title, note, headers, rows, widths, link_cols=(), input_cols=()):
    ws["A1"], ws["A1"].font = title, TITLE
    ws["A2"], ws["A2"].font = note, NOTE
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font, cell.fill, cell.alignment = BOLD, HEAD, WRAP
    for r, row in enumerate(rows, 5):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font, cell.alignment, cell.border = FONT, WRAP, RULE
            if c in link_cols and isinstance(v, str) and v.startswith("http"):
                cell.hyperlink, cell.font = v, LINK
            if c in input_cols:
                cell.fill = NOTE_FILL
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"
    return ws


def pack_summary(name):
    d = json.loads((root / "content/standards/drafts" / name).read_text())
    if "sets" in d:
        return f"{len(d['sets'])} threshold sets"
    return f"{len(d['lines'])} lines, {len(d['mapEntries'])} map entries, {len(d['rules'])} wording rules, {len(d['validationRules'])} validation rules"


wb = Workbook()
ws = wb.active
ws.title = "Steps"
table(ws, "Rule verification worksheet (DEC-083)",
      "Generated by scripts/docs/gen_rule_verification.py. Nothing here is verified: verification is done in the platform, by Shoaib, against the provisions named.",
      ["Step", "What", "Who", "Note"], STEPS, [6, 80, 34, 52])

vat = wb.create_sheet("VAT codes")
rows = [r + ("", "", "") for r in VAT_CODES]
table(vat, "VAT codes and rates to enter (proposals to check)",
      "Rates and boxes are engineering's proposals, written without reading the sources. Yellow columns are for the verifier.",
      ["Code", "Name", "Treatment", "Direction", "Reverse charge", "Rate %", "Net boxes", "VAT boxes", "Provision to check", "Source", "What to check",
       "Verified rate and boxes", "Verified on", "Verifier's notes"],
      rows, [8, 34, 14, 12, 10, 8, 10, 10, 34, 40, 50, 22, 14, 40], link_cols=(10,), input_cols=(12, 13, 14))

std = wb.create_sheet("Standards and thresholds")
srows = [
    ("FRS 105 micro-entity pack", "content/standards/drafts/frs105-micro.json", pack_summary("frs105-micro.json"), "v2-draft-wording-review.xlsx: FRS105 wording, lines, map", "", "", ""),
    ("FRS 102 Section 1A small-entity pack", "content/standards/drafts/frs102-1a-small.json", pack_summary("frs102-1a-small.json"), "v2-draft-wording-review.xlsx: FRS102-1A wording, lines, map", "", "", ""),
    ("Companies Act size thresholds", "content/standards/drafts/uk-size-thresholds.json", pack_summary("uk-size-thresholds.json"), "v2-draft-wording-review.xlsx: Size thresholds", "", "", ""),
]
table(std, "Draft standards and thresholds to import and verify", "Each lands UNVERIFIED on import (DEC-052); the review workbook holds every line for checking.",
      ["Item", "File", "Contents", "Where to review it", "Verified version id", "Verified on", "Verifier's notes"], srows, [36, 46, 52, 50, 22, 14, 40], input_cols=(5, 6, 7))

# V7-M1 (DEC-088): the draft deadline rules, read from the pack itself so the sheet cannot drift from it
pack = json.loads((root / "content/compliance/drafts/uk-deadline-rules.json").read_text())
drows = [(r["kind"], ", ".join(r["legalForms"]), r["daysBeforeMonths"], r["months"], r["daysAfterMonths"], "yes" if r["monthEndRule"] else "no", r["effectiveFrom"],
          r["source"]["provision"], r["source"]["url"], r["verifierNote"], "", "", "") for r in pack["rules"]]
ded = wb.create_sheet("Deadline rules")
table(ded, "Draft deadline rules (V7-M1, DEC-088)",
      "Due date = anchor + days before, then + months (month-end rule: a month-end anchor stays a month end), then + days after. Import with scripts/compliance/import-deadline-rules.mjs; each lands UNVERIFIED.",
      ["Kind", "Legal forms", "Days before", "Months", "Days after", "Month-end rule", "Governs anchors from", "Provision to check", "Source", "What the rule does not model",
       "Verified (yes / corrected)", "Verified on", "Verifier's notes"],
      drows, [22, 30, 8, 8, 8, 10, 14, 60, 44, 60, 18, 14, 40], link_cols=(9,), input_cols=(11, 12, 13))

# V5-M2 (DEC-091): the presentation rule of boxes 6 to 9; engineering's proposal, written without reading the source
PRES = [("TRUNCATE_PENCE", "Whole pounds: the pence are left out of boxes 6, 7, 8 and 9 (boxes 1 to 5 keep pounds and pence)", "2026-01-01",
         "VAT Notice 700/12, the guidance on boxes 6 to 9", "https://www.gov.uk/guidance/how-to-fill-in-and-submit-your-vat-return-vat-notice-70012",
         "Whether pence are left out (truncated) or rounded, and from which date; the alternative method is ROUND_TO_POUND", "", "", "")]
pres = wb.create_sheet("VAT presentation")
table(pres, "Presentation rule for boxes 6 to 9 (V5-M2, DEC-091)",
      "One current version for the organisation; it applies to a return by its period end. Enter it with its source; it lands UNVERIFIED.",
      ["Method", "What it does", "Effective from", "Provision to check", "Source", "What to check", "Verified (yes / corrected)", "Verified on", "Verifier's notes"],
      PRES, [18, 60, 14, 40, 60, 60, 18, 14, 40], link_cols=(5,), input_cols=(7, 8, 9))

# V3-M1 (DEC-093): the draft taxonomy mapping, read from the pack itself so the sheet cannot drift from it
xpack = json.loads((root / "content/ixbrl/drafts/frc-frs105.json").read_text())
xrows = [(e["sourceKind"], e["sourceKey"], e["concept"], e["dataType"], e["periodType"], "yes" if e["negate"] else "no", "" if e["decimals"] is None else e["decimals"],
          "; ".join(f'{d["axis"]} = {d["member"]}' for d in e["dimensions"]), "yes" if e["mandatory"] else "no", e["verifierNote"], "", "", "") for e in xpack["entries"]]
xs = wb.create_sheet("iXBRL mapping")
table(xs, "Draft taxonomy mapping for FRS 105 accounts (V3-M1, DEC-093)",
      f"Entry point {xpack['version']['entryPoint']}; namespaces " + ", ".join(f"{k} = {v}" for k, v in xpack['version']['namespaces'].items()) + ". Every concept is a proposal written without reading the taxonomy.",
      ["Source kind", "Source", "Concept", "Data type", "Period", "Negate", "Decimals", "Dimensions", "Mandatory", "What to check", "Verified (yes / corrected)", "Verified on", "Verifier's notes"],
      xrows, [12, 18, 52, 12, 10, 8, 9, 50, 10, 60, 18, 14, 40], input_cols=(11, 12, 13))

# V4-M1 (DEC-097): the draft Corporation Tax rules, read from the pack itself so the sheet cannot drift from it
cpack = json.loads((root / "content/tax/drafts/uk-ct-rules.json").read_text())
crows = [(r["kind"], r["effectiveFrom"], r.get("effectiveTo", ""), "; ".join(f"{k} = {v['numerator']}/{v['denominator']}" if isinstance(v, dict) else f"{k} = {v}" for k, v in r["parameters"].items()),
          r["source"]["provision"], r["source"]["url"], r["verifierNote"], "", "", "") for r in cpack["rules"]]
cts = wb.create_sheet("Corporation Tax rules")
table(cts, "Draft Corporation Tax rules (V4-M1, DEC-097)",
      "Import with scripts/tax/import-ct-rules.mjs; each lands UNVERIFIED. Rates are percentages, limits are pounds for a 12-month period with no associated companies.",
      ["Kind", "Effective from", "Effective to", "Parameters", "Provision to check", "Source", "What to check", "Verified (yes / corrected)", "Verified on", "Verifier's notes"],
      crows, [20, 14, 14, 60, 60, 50, 70, 18, 14, 40], link_cols=(6,), input_cols=(8, 9, 10))

sign = wb.create_sheet("Sign-off")
sign_rows = ([(r[0], "Rate version", "", "", "") for r in VAT_CODES] + [(r[0], "Standard version", "", "", "") for r in srows]
             + [(r["kind"], "Deadline rule", "", "", "") for r in pack["rules"]] + [("Boxes 6 to 9", "Presentation rule", "", "", ""), ("FRS 105 taxonomy mapping", "Taxonomy version", "", "", "")]
             + [(r["kind"], "Corporation Tax rule", "", "", "") for r in cpack["rules"]])
table(sign, "Sign-off", "Filled in by the verifier from the platform's records.", ["Item", "Kind", "Platform record id", "Verified by", "Date"],
      sign_rows, [40, 18, 40, 20, 14], input_cols=(3, 4, 5))

law = wb.create_sheet("Legal references")
table(law, "Legislation, guidance and case law", "Not opened from this environment and not legal advice; docs/legal/m4b-vat-sources.xlsx has the full VAT source list.",
      ["Provision or case", "What it decides here", "Link"], LAW, [70, 60, 70], link_cols=(3,))
wb.save(OUT)
print(OUT)
