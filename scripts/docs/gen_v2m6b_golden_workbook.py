#!/usr/bin/env python3
"""Writes docs/architecture/v2-m6b-golden-workings.xlsx from tests/golden/v2-m6b-workflow.json: the control status and the state of every step of
the V2-M6b golden year-end workflow (DEC-069), re-derived here phase by phase from the scenario's facts and an independent implementation of the
DEC-069 rules, next to the hand-written expectations with a check column. For the named accountant's review (DEC-054). Exits 1 if anything
differs."""
import json, pathlib, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

root = pathlib.Path(__file__).resolve().parents[2]
G = json.loads((root / "tests/golden/v2-m6b-workflow.json").read_text())
rows, bad = [], 0
def check(step, item, got, exp):
    global bad
    ok = got == exp
    bad += not ok; rows.append([step, item, json.dumps(got), json.dumps(exp), "OK" if ok else "DIFFERS"])

# DEC-069 as amended by DEC-070: step 12 waits for steps 1 to 11; steps 12 to 14 are done through their own records (none in this scenario)
DEPS = {1: [], **{n: [1] for n in range(2, 10)}, 10: list(range(1, 10)), 11: [10], 12: list(range(1, 12)), 13: [12], 14: [13]}
OWN_RECORD = {12, 13, 14}

# ── the facts of each phase, from the scenario (what the platform's controls will read) ──
facts = [
    dict(drafts=1, rc08="OPEN", settings=False, run=False, vat=False, suspense=0, loans_flagged=0, checklist_complete=False, accounts_complete=True),
    dict(drafts=0, rc08="OPEN", settings=False, run=False, vat=False, suspense=0, loans_flagged=0, checklist_complete=False, accounts_complete=True),
    dict(drafts=0, rc08="SIGNED_OFF", settings=True, run=True, vat=False, suspense=0, loans_flagged=0, checklist_complete=True, accounts_complete=True),
    dict(drafts=0, rc08="SIGNED_OFF", settings=True, run=True, vat=False, suspense=0, loans_flagged=0, checklist_complete=True, accounts_complete=True),
    dict(drafts=1, rc08="SIGNED_OFF", settings=True, run=True, vat=False, suspense=0, loans_flagged=0, checklist_complete=True, accounts_complete=True),
]
# the sign-off log: (phase it is done in, action, step)
log = [(1, "SIGN_OFF", 1)] + [(3, "SIGN_OFF", n) for n in range(2, 12)] + [(4, "REOPEN", 9)]

def controls(f):
    r = lambda ok: "READY" if ok else "ATTENTION"
    return {1: r(f["drafts"] == 0), 2: r(f["rc08"] in ("CLEAR", "SIGNED_OFF")), 3: r(f["settings"]), 4: r(f["settings"]), 5: r(f["run"]), 6: "READY",
            7: r(f["loans_flagged"] == 0), 8: "NOT_APPLICABLE" if not f["vat"] else "READY", 9: r(f["suspense"] == 0), 10: r(f["checklist_complete"]), 11: r(f["accounts_complete"]),
            # 12: the document is complete (no review yet); 13: no live review, the period open, no approval facts; 14: no client approval
            12: r(f["accounts_complete"]), 13: "ATTENTION", 14: "ATTENTION"}

for i, (phase, f) in enumerate(zip(G["phases"], facts)):
    c = controls(f)
    signed = {}
    for p, action, step in log:
        if p <= i: signed[step] = action == "SIGN_OFF"
    state = {}
    ok = lambda s: c[s] in ("READY", "NOT_APPLICABLE")
    for n in range(1, 15):
        if signed.get(n): state[n] = "DONE" if ok(n) else "LAPSED"
        else: state[n] = None
    for n in range(1, 15):
        if state[n] is None:
            state[n] = "BLOCKED" if any(state[d] != "DONE" for d in DEPS[n]) else ("TODO" if ok(n) else "ATTENTION")
    current = next(n for n in range(1, 15) if state[n] != "DONE")
    for n in range(1, 15):
        check(phase["name"], f"step {n}", [c[n], state[n]], phase["steps"][n - 1])
    check(phase["name"], "current step", current, phase["current"])
    # the refusals the phase expects
    for rf in phase.get("refusals", []):
        n, action = rf["step"], rf.get("action", "sign-off")
        if action == "reopen": got = "step_not_signed_off" if not signed.get(n) else "allowed"
        else: got = "step_own_record" if n in OWN_RECORD else "step_blocked" if any(state[d] != "DONE" for d in DEPS[n]) else "step_not_ready" if not ok(n) else "allowed"
        check(phase["name"], f"refusal {action} step {n}", got, rf["code"])

wb = Workbook(); ws = wb.active; ws.title = "Read me"
for line in ["V2-M6b golden year-end workflow - workings (SYNTHETIC)", "", G["purpose"], "", G["disclaimer"], "",
             "Each phase's control statuses are derived from the scenario's facts and each step's state from an independent implementation of the DEC-069 rules as amended by DEC-070 (order, sign-offs, lapses; steps 12 to 14 through their own records), then compared with the hand-written expectations. The platform's API test matches the same states exactly.",
             "", *[f"{k}: {v}" for k, v in G["workings"].items()], "", *[f"Revision: {r}" for r in G.get("revisions", [])]]:
    ws.append([line])
ws.column_dimensions["A"].width = 150; ws["A1"].font = Font(bold=True, size=14)
for r in ws.iter_rows():
    for cc in r: cc.alignment = Alignment(wrap_text=True, vertical="top")
s2 = wb.create_sheet("Phases"); s2.append(["Phase", "Current step"] + [f"Step {n}" for n in range(1, 15)])
for p in G["phases"]: s2.append([p["name"], p["current"]] + [f"{c or '-'} / {s}" for c, s in p["steps"]])
s3 = wb.create_sheet("Checks"); s3.append(["Phase", "Item", "Recomputed", "Expected (hand)", "Check"])
for r in rows: s3.append(r)
for sh, widths in ((s2, [40, 8] + [18] * 14), (s3, [40, 26, 34, 34, 10])):
    for i, w in enumerate(widths): sh.column_dimensions[chr(65 + i)].width = w
    for cc in sh[1]: cc.font = Font(bold=True, color="FFFFFF"); cc.fill = PatternFill("solid", fgColor="1F3864")
    for row in sh.iter_rows(min_row=2):
        for cc in row: cc.alignment = Alignment(wrap_text=True, vertical="top")
for row in s3.iter_rows(min_row=2): row[4].fill = PatternFill("solid", fgColor="C6EFCE" if row[4].value == "OK" else "FFC7CE")
o = root / "docs/architecture/v2-m6b-golden-workings.xlsx"; wb.save(o)
print(f"wrote {o.relative_to(root)}; {len(rows)} checks; differences: {bad}")
for r in rows:
    if r[4] == "DIFFERS": print("DIFFERS", r)
sys.exit(1 if bad else 0)
