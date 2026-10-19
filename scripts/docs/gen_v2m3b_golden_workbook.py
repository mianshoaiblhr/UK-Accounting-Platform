#!/usr/bin/env python3
"""Writes docs/architecture/v2-m3b-golden-workings.xlsx from tests/golden/v2-m3b-close.json: every closing journal, trial balance, balance sheet and
P&L figure of the V2-M3b golden scenario (DEC-060) recomputed here with Python's Decimal, independently of the platform, next to the hand-written
expected figure with a check column. For the named accountant's review (DEC-054). Exits 1 if any figure differs."""
import datetime as dt, json, pathlib, sys
from decimal import Decimal as D

root = pathlib.Path(__file__).resolve().parents[2]
G = json.loads((root / "tests/golden/v2-m3b-close.json").read_text())
day = dt.date.fromisoformat
TYPE = lambda code: "ASSET" if code[0] in "01" else "LIABILITY" if code[0] == "2" else "EQUITY" if code[0] == "3" else "INCOME" if code[0] == "4" else "EXPENSE"
PL = lambda code: TYPE(code) in ("INCOME", "EXPENSE")
MAP = {"0030": "BS.FIXED_ASSETS.TANGIBLE", "1110": "BS.CURRENT_ASSETS.DEBTORS", "2110": "BS.CREDITORS_WITHIN_ONE_YEAR", "3000": "BS.CAPITAL.SHARE_CAPITAL", "3200": "BS.CAPITAL.RESERVES",
       "4000": "PL.TURNOVER", "4100": "PL.OTHER_INCOME", "6100": "PL.ADMIN_EXPENSES", "6110": "PL.ADMIN_EXPENSES", "6300": "PL.ADMIN_EXPENSES", "6310": "PL.ADMIN_EXPENSES", "6600": "PL.ADMIN_EXPENSES"}
CLOSING = {"YEAR_END_CLOSE", "YEAR_END_CLOSE_REVERSAL"}
ledger = []                                         # (date, account, debit - credit, source, journal ref)
rows, bad = [], 0
def post(d, lines, source="MANUAL", ref=""): ledger.extend((d, k, D(dr) - D(cr), source, ref) for k, dr, cr in lines)
def check(step, item, got, exp):
    global bad
    ok = D(str(got)) == D(str(exp)); bad += not ok
    rows.append([step, item, float(D(str(got))), float(D(str(exp))), "OK" if ok else "DIFFERS"])
def jcheck(step, got, exp):
    global bad
    g = [(k, f"{D(dr):.2f}", f"{D(cr):.2f}") for k, dr, cr in got]; e = [tuple(x) for x in exp]
    ok = g == e; bad += not ok          # order matters: account code order, retained earnings last (DEC-060)
    rows.append([step, "journal lines (in order)", "; ".join(f"{k} {dr}/{cr}" for k, dr, cr in g), "; ".join(f"{k} {dr}/{cr}" for k, dr, cr in e), "OK" if ok else "DIFFERS"])
def flag(step, item, got, exp):
    global bad
    ok = got == exp; bad += not ok; rows.append([step, item, str(got), str(exp), "OK" if ok else "DIFFERS"])
def net(code, frm, to, closing=True):
    return sum((v for (d, k, v, s, _) in ledger if k == code and frm <= d <= to and (closing or s not in CLOSING)), D(0))
LOW = day("0001-01-01")
codes = lambda: sorted({k for (_, k, _, _, _) in ledger})
year_start = lambda d: next(day(s) for s, e in G["periods"] if day(s) <= d <= day(e))
live = []                                           # live closes: (start, end)

def pl_figures(frm, to):
    lines = {}
    for c in codes():
        if not PL(c): continue
        n = net(c, frm, to, closing=False)
        amt = -n if TYPE(c) == "INCOME" else n
        lines[MAP[c]] = lines.get(MAP[c], D(0)) + amt
    v = lambda k: lines.get(k, D(0))
    gross = v("PL.TURNOVER") - v("PL.COST_OF_SALES")
    op = gross - v("PL.ADMIN_EXPENSES") + v("PL.OTHER_INCOME")
    ledger_net = -sum((net(c, frm, to, closing=False) for c in codes() if PL(c)), D(0))
    return lines, gross, op, ledger_net
def check_pl(step, spec, frm, to):
    lines, gross, op, ledger_net = pl_figures(frm, to)
    for k, exp in spec["lines"].items(): check(step, k, lines.get(k, D(0)), exp)
    if "grossProfit" in spec: check(step, "gross profit", gross, spec["grossProfit"])
    if "operatingProfit" in spec: check(step, "operating profit", op, spec["operatingProfit"])
    check(step, "profit for the period", op, spec["profitForPeriod"])
    if "incomeLessExpensesFromLedger" in spec: check(step, "income less expenses from the ledger (closes left out)", ledger_net, spec["incomeLessExpensesFromLedger"])
def bs_figures(asof):
    lines = {}
    for c in codes():
        if PL(c): continue
        n = net(c, LOW, asof, closing=False)
        amt = n if TYPE(c) == "ASSET" else -n
        lines[MAP[c]] = lines.get(MAP[c], D(0)) + amt
    cum = -sum((net(c, LOW, asof, closing=False) for c in codes() if PL(c)), D(0))
    ytd = -sum((net(c, year_start(asof), asof, closing=False) for c in codes() if PL(c)), D(0))
    lines["PL.BROUGHT_FORWARD"], lines["PL.YEAR_TO_DATE"] = cum - ytd, ytd
    v = lambda k: lines.get(k, D(0))
    net_assets = v("BS.FIXED_ASSETS.TANGIBLE") + v("BS.CURRENT_ASSETS.DEBTORS") - v("BS.CREDITORS_WITHIN_ONE_YEAR")
    equity = v("BS.CAPITAL.SHARE_CAPITAL") + v("BS.CAPITAL.RESERVES") + cum
    closed = -sum((v2 for (d, k, v2, s, _) in ledger if k == "3200" and d <= asof and s in CLOSING), D(0))
    return lines, net_assets, equity, closed
def check_bs(step, spec):
    asof = day(spec["asOf"]); lines, na, eq, closed = bs_figures(asof)
    for k, exp in spec["lines"].items(): check(step, k, lines.get(k, D(0)), exp)
    check(step, "net assets", na, spec["netAssets"]); check(step, "capital and reserves", eq, spec["equity"])
    if "closedToRetainedEarnings" in spec: check(step, "closed to retained earnings", closed, spec["closedToRetainedEarnings"])
def check_tb(step, spec):
    asof = day(spec["asOf"]); td = tc = D(0)
    present = codes()
    flag(step, "accounts listed", [r[0] for r in spec["rows"]], present)
    for code, edr, ecr in spec["rows"]:
        n = net(code, LOW, asof); dr, cr = (n, D(0)) if n > 0 else (D(0), -n); td += dr; tc += cr
        check(step, f"{code} debit", dr, edr); check(step, f"{code} credit", cr, ecr)
    check(step, "total debit", td, spec["totalDebit"]); check(step, "total credit", tc, spec["totalCredit"])
def close(start, end):
    """The closing journal: every income and expense account's net for the year reversed, in code order, the difference to retained earnings last."""
    lines = []
    for c in codes():
        if not PL(c): continue
        n = net(c, start, end)                      # the whole ledger: an earlier reversed close nets to nothing
        if n != 0: lines.append((c, -n if n < 0 else D(0), n if n > 0 else D(0)))
    profit = sum((dr - cr for _, dr, cr in lines), D(0))
    if profit != 0: lines.append(("3200", -profit if profit < 0 else D(0), profit if profit > 0 else D(0)))
    return lines, profit
def earlier_unclosed(start):
    return any(PL(k) and d < start and not any(s <= d <= e for s, e in live) and s2 not in CLOSING for (d, k, _, s2, _) in ledger)
def closed_year(d): return any(s <= d <= e for s, e in live)

# seed and the V2-M3a accrual (auto-reversed the next day)
for j in G["seed"]: post(day(j["date"]), j["lines"], ref=j["ref"])
a = G["accrual"]; post(day(a["date"]), a["expect"]["journal"], "ACCRUAL", "ACCRUAL"); post(day(a["expect"]["reversal"]["date"]), a["expect"]["reversal"]["journal"], "ACCRUAL_AUTO_REVERSAL", "ACCRUAL_REV")
for k, v in G["workings"].items(): rows.append(["Workings", k, v, "", "-"])
B = G["before"]
check_pl("Before: P&L 2026", B["profitAndLoss2026"], day("2026-01-01"), day("2026-12-31"))
check_pl("Before: P&L 2025-01-01..2027-06-30", B["profitAndLossAll"], day("2025-01-01"), day("2027-06-30"))
for bs in B["balanceSheets"]: check_bs(f"Before: balance sheet {bs['asOf']}", bs)
check_tb("Before: trial balance 2026-12-31", B["trialBalance2026"])

C = G["closes"]
pv = C["preview2026First"]; l, p = close(day(pv["yearStart"]), day(pv["yearEnd"]))
check("Preview 2026 (2025 not closed)", "profit", p, pv["profit"]); flag("Preview 2026 (2025 not closed)", "warnings", ["earlier_years_not_closed"] if earlier_unclosed(day(pv["yearStart"])) else [], pv["warnings"])
for key in ("close2025", "close2026"):
    c = C[key]; s, e = day(c["yearStart"]), day(c["yearEnd"])
    if "warnings" in c: flag(f"Close {c['yearStart'][:4]}", "warnings", ["earlier_years_not_closed"] if earlier_unclosed(s) else [], c["warnings"])
    l, p = close(s, e); jcheck(f"Close {c['yearStart'][:4]}", l, c["journal"]); check(f"Close {c['yearStart'][:4]}", "profit", p, c["profit"])
    post(day(c["journalDate"]), [(k, dr, cr) for k, dr, cr in l], "YEAR_END_CLOSE", key); live.append((s, e))

A = G["after"]
check_pl("After: P&L 2026 (same as before)", B["profitAndLoss2026"], day("2026-01-01"), day("2026-12-31"))
check_pl("After: P&L 2025-01-01..2027-06-30 (same as before)", B["profitAndLossAll"], day("2025-01-01"), day("2027-06-30"))
for bs in B["balanceSheets"]: check_bs(f"After: balance sheet {bs['asOf']} (same as before)", {**bs, "closedToRetainedEarnings": A["closedToRetainedEarnings"][bs["asOf"]]})
check_tb("After: trial balance 2026-12-31", A["trialBalance2026"])
check("After: close 2026", "year income less expenses, close included", -sum((net(c, day("2026-01-01"), day("2026-12-31")) for c in codes() if PL(c)), D(0)), A["closeView2026"]["yearIncomeLessExpensesIncludingClose"])
gl = A["generalLedger4000"]; f, t = day(gl["from"]), day(gl["to"])
for mode, incl in (("all", True), ("excludeClosing", False)):
    g = gl[mode]; opening = -net("4000", LOW, f - dt.timedelta(days=1), closing=incl)
    check(f"After: general ledger 4000 ({mode})", "opening balance", opening, g["openingBalance"])
    run = opening; items = []
    for (d, k, v, s, ref) in ledger:
        if k == "4000" and f <= d <= t and (incl or s not in CLOSING):
            run -= v; items.append([{"close2026": "CLOSE2026"}.get(ref, ref), f"{max(v, D(0)):.2f}", f"{max(-v, D(0)):.2f}", f"{run:.2f}"])
    flag(f"After: general ledger 4000 ({mode})", "lines (ref, debit, credit, balance)", items, g["items"])
rp = A["refusedPosting"]; flag("After: posting income or expense into 2026", "refused", "year_closed" if any(PL(k) for k, _, _ in rp["lines"]) and closed_year(day(rp["date"])) else "posted", rp["code"])
flag("After: closing 2026 again", "refused", "year_already_closed" if any(s <= day("2026-12-31") and day("2026-01-01") <= e for s, e in live) else "posted", A["closeAgain2026"]["code"])

R = G["reopen"]
orig = [(k, f"{v if v > 0 else 0:.2f}", f"{(-v) if v < 0 else 0:.2f}") for (d, k, v, s, ref) in ledger if ref == "close2026"]
mirror = [(k, cr, dr) for k, dr, cr in orig]; jcheck("Reverse close 2026 (mirror)", mirror, R["reverse2026"]["journal"])
post(day(R["reverse2026"]["journalDate"]), mirror, "YEAR_END_CLOSE_REVERSAL", "rev2026"); live.remove((day("2026-01-01"), day("2026-12-31")))
lp = R["latePosting"]; flag("After the reversal: late posting", "accepted", "posted" if not closed_year(day(lp["date"])) else "year_closed", "posted"); post(day(lp["date"]), lp["lines"], ref="LATE")
l, p = close(day("2026-01-01"), day("2026-12-31")); jcheck("Re-close 2026", l, R["reclose2026"]["journal"]); check("Re-close 2026", "profit", p, R["reclose2026"]["profit"])
post(day("2026-12-31"), l, "YEAR_END_CLOSE", "reclose2026"); live.append((day("2026-01-01"), day("2026-12-31")))
check_pl("After the re-close: P&L 2026", R["profitAndLoss2026"], day("2026-01-01"), day("2026-12-31"))
check_bs("After the re-close: balance sheet 2026-12-31", R["balanceSheet2026"])

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:
    print("openpyxl missing"); sys.exit(1)
wb = Workbook(); ws = wb.active; ws.title = "Read me"
for line in ["V2-M3b golden year-end closing journal - workings (SYNTHETIC)", "", G["purpose"], "", G["disclaimer"], "",
             "Every figure is recomputed by this workbook's generator with Python Decimal, independently of the platform, and compared with the hand-written expected figure. The platform's API test matches the same figures to the cent.",
             "", "Rules recomputed (DEC-060): the closing journal reverses each income and expense account's net for the year, in account-code order, with the difference (the year's profit) to retained earnings last.",
             "Every P&L figure leaves closing journals out; the balance sheet's lines leave them out too, so they are the same before and after the close, while the trial balance includes them.",
             "While a close is live nothing posts income or expense into its year; to change the year the close is reversed by its mirror and posted again.", "", G["after"]["note"]]: ws.append([line])
ws.column_dimensions["A"].width = 150; ws["A1"].font = Font(bold=True, size=14)
for r in ws.iter_rows():
    for c in r: c.alignment = Alignment(wrap_text=True, vertical="top")
s2 = wb.create_sheet("Ledger"); s2.append(["Date", "Account", "Debit", "Credit", "Source", "Journal"])
for (d, k, v, s, ref) in ledger: s2.append([d.isoformat(), k, float(max(v, D(0))), float(max(-v, D(0))), s, ref])
for i, w in enumerate([12, 10, 14, 14, 26, 14], 1): s2.column_dimensions["ABCDEF"[i - 1]].width = w
s3 = wb.create_sheet("Checks"); s3.append(["Step", "Figure", "Recomputed", "Expected (hand)", "Check"])
for r in rows: s3.append(r)
for i, w in enumerate([44, 40, 60, 60, 10], 1): s3.column_dimensions["ABCDE"[i - 1]].width = w
for sh in (s2, s3):
    for c in sh[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3864")
for row in s3.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if isinstance(c.value, float): c.number_format = '#,##0.00;[Red]-#,##0.00'
    row[4].fill = PatternFill("solid", fgColor="C6EFCE" if row[4].value == "OK" else "FFC7CE" if row[4].value == "DIFFERS" else "FFFFFF")
for row in s2.iter_rows(min_row=2):
    for c in row[2:4]: c.number_format = '#,##0.00'
out = root / "docs/architecture/v2-m3b-golden-workings.xlsx"; wb.save(out)
n = sum(1 for r in rows if r[4] != "-")
print(f"wrote {out.relative_to(root)}; {n} figures; differences: {bad}")
for r in rows:
    if r[4] == "DIFFERS": print("DIFFERS", r)
sys.exit(1 if bad else 0)
