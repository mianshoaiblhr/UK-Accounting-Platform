#!/usr/bin/env python3
"""Writes docs/architecture/v2-m6c-golden-workings.xlsx from tests/golden/v2-m6c-approval.json: the control status and the state of every step of
the V2-M6c golden (DEC-070) - partner review, client approval, locking and the filing-ready package - re-derived here phase by phase from the
scenario's facts and an independent implementation of the DEC-069 and DEC-070 rules, next to the hand-written expectations with a check column.
For the named accountant's review (DEC-054). Exits 1 if anything differs."""
import json, pathlib, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

root = pathlib.Path(__file__).resolve().parents[2]
G = json.loads((root / "tests/golden/v2-m6c-approval.json").read_text())
rows, bad = [], 0
def check(phase, item, got, exp):
    global bad
    ok = got == exp
    bad += not ok; rows.append([phase, item, json.dumps(got), json.dumps(exp), "OK" if ok else "DIFFERS"])

# DEC-069 as amended by DEC-070
DEPS = {1: [], **{n: [1] for n in range(2, 10)}, 10: list(range(1, 10)), 11: [10], 12: list(range(1, 12)), 13: [12], 14: [13]}
OWN = {12, 13, 14}

# ── the scenario, phase by phase: the accounts fingerprint (A, then B after the fact change) and the records that exist ──
# reviews: (fingerprint, by); approval: fingerprint; package: fingerprint
S = [
    dict(fp="A", reviews=[], period="OPEN", approval=None, package=None),
    dict(fp="A", reviews=[("A", "owner")], period="OPEN", approval=None, package=None),
    dict(fp="B", reviews=[("A", "owner")], period="OPEN", approval=None, package=None),
    dict(fp="B", reviews=[("A", "owner"), ("B", "owner")], period="CLOSED", approval=None, package=None),
    dict(fp="B", reviews=[("A", "owner"), ("B", "owner")], period="LOCKED", approval="B", package=None),
    dict(fp="B", reviews=[("A", "owner"), ("B", "owner")], period="LOCKED", approval="B", package="B"),
]
STEP11_BY = "partner"          # steps 1 to 10 by the accountant, step 11 by the partner
APPROVAL_FACTS = {"director": "@director", "approvedOn": "2026-03-16"}
RULES_VERIFIED = True          # the synthetic version and threshold set are verified by a second person in the setup
ROLES = {"owner": {"accounts:review", "accounts:prepare", "period:lock"}, "partner": {"accounts:review", "accounts:prepare", "period:lock"},
         "accountant": {"accounts:prepare"}}

def controls(s):
    c = {n: "READY" for n in range(1, 12)}; c[8] = "NOT_APPLICABLE"
    latest = s["reviews"][-1][0] if s["reviews"] else None
    live = latest == s["fp"]
    c[12] = "READY" if (live if latest else s["fp"] is not None) else "ATTENTION"
    if s["approval"]: c[13] = "READY" if s["approval"] == s["fp"] else "ATTENTION"
    else: c[13] = "READY" if live and s["period"] in ("CLOSED", "LOCKED") and APPROVAL_FACTS else "ATTENTION"
    if s["package"]: c[14] = "READY" if s["package"] == s["fp"] else "ATTENTION"
    else: c[14] = "READY" if s["approval"] and s["approval"] == s["fp"] and RULES_VERIFIED else "ATTENTION"
    return c

def states(s, c):
    signed = {n: True for n in range(1, 12)}
    signed[12], signed[13], signed[14] = bool(s["reviews"]), bool(s["approval"]), bool(s["package"])
    ok = lambda n: c[n] in ("READY", "NOT_APPLICABLE")
    st = {n: ("DONE" if ok(n) else "LAPSED") for n in range(1, 15) if signed[n]}
    for n in range(1, 15):
        if n not in st: st[n] = "BLOCKED" if any(st[d] != "DONE" for d in DEPS[n]) else ("TODO" if ok(n) else "ATTENTION")
    return st, next((n for n in range(1, 15) if st[n] != "DONE"), None)

def refusal(s, st, r):
    a, by = r["action"], r["by"]
    locked = s["approval"] is not None
    if a == "partner-review":
        if locked: return "year_approved"
        if "accounts:review" not in ROLES[by]: return 403
        if st[12] == "BLOCKED": return "step_blocked"
        if s["reviews"] and s["reviews"][-1][0] == s["fp"]: return "already_reviewed"
        if by == STEP11_BY: return "second_person_required"
        return "allowed"
    if a == "client-approval":
        if not {"accounts:review", "period:lock"} <= ROLES[by]: return 403
        if locked: return "year_approved"
        if st[13] == "BLOCKED": return "step_blocked"
        if st[13] != "TODO": return "step_not_ready"
        if r.get("director", APPROVAL_FACTS["director"]) != APPROVAL_FACTS["director"] or r.get("approvedOn", APPROVAL_FACTS["approvedOn"]) != APPROVAL_FACTS["approvedOn"]: return "approval_mismatch"
        if r.get("evidence") == "missing": return "evidence_invalid"
        return "allowed"
    if a == "package":
        if s["package"]: return "package_exists"
        if st[14] == "BLOCKED": return "step_blocked"
        return "allowed" if st[14] == "TODO" else "step_not_ready"
    if a in ("sign-off", "reopen"):
        if r["step"] in OWN: return "step_own_record"
        return "year_approved" if locked else "allowed"
    if a in ("fact", "flag-reopen", "unlock"): return "year_approved" if locked else "allowed"
    if a == "journal": return "period_locked" if s["period"] == "LOCKED" else "allowed"
    raise ValueError(a)

for s, phase in zip(S, G["phases"]):
    c = controls(s); st, current = states(s, c)
    for n in range(1, 15): check(phase["name"], f"step {n}", [c[n], st[n]], phase["steps"][n - 1])
    check(phase["name"], "current step", current, phase["current"])
    if "periodStatus" in phase: check(phase["name"], "period status", s["period"], phase["periodStatus"])
    for r in phase.get("refusals", []):
        exp = r.get("code", r.get("status"))
        check(phase["name"], f"refusal {r['action']}{' step ' + str(r['step']) if 'step' in r else ''} by {r['by']}" + (f" ({', '.join(k for k in r if k in ('director', 'approvedOn', 'evidence'))})" if any(k in r for k in ('director', 'approvedOn', 'evidence')) else ""), refusal(s, st, r), exp)

# the package's wording, from the rules and the facts as they stand at the end (employees changed from 1 to 2)
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]
facts = dict(G["setup"]["facts"]); facts[G["setup"]["factChange"][0]] = G["setup"]["factChange"][1]
y, m, d = map(int, facts["approval.date"].split("-"))
values = {"employees.average": facts["employees.average"], "approval.date": f"{d} {MONTHS[m - 1]} {y}", "approval.director": G["company"]["director"]}
wording = []
for rule in G["setup"]["rules"]:
    t = rule["body"]
    for k, v in values.items(): t = t.replace("{{fact:" + k + "}}", v)
    wording.extend(t.split("\n"))   # each line of the wording is its own paragraph (V2-M5a)
check("package", "wording", wording, G["package"]["wording"])
check("package", "approval", {"director": G["company"]["director"], "approvedOn": G["setup"]["approval"]["approvedOn"]}, G["package"]["approval"])

wb = Workbook(); ws = wb.active; ws.title = "Read me"
for line in ["V2-M6c golden approval, locking and package - workings (SYNTHETIC)", "", G["purpose"], "", G["disclaimer"], "",
             "Each phase's control statuses and step states are derived from the scenario's records (the fingerprint of the accounts, the partner reviews, the approval, the package, the period status) with an independent implementation of DEC-069 as amended by DEC-070, then compared with the hand-written expectations; each refusal is re-derived from the same rules. The platform's API test checks the same states and codes.",
             "", *[f"{k}: {v}" for k, v in G["workings"].items()], "", f"Correction: {G['package']['correction']}"]:
    ws.append([line])
ws.column_dimensions["A"].width = 150; ws["A1"].font = Font(bold=True, size=14)
for r in ws.iter_rows():
    for cc in r: cc.alignment = Alignment(wrap_text=True, vertical="top")
s2 = wb.create_sheet("Phases"); s2.append(["Phase", "Current step"] + [f"Step {n}" for n in range(1, 15)])
for p in G["phases"]: s2.append([p["name"], p["current"]] + [f"{c or '-'} / {s}" for c, s in p["steps"]])
s3 = wb.create_sheet("Records"); s3.append(["Phase", "Fingerprint now", "Partner reviews (fingerprint, by)", "Period", "Client approval", "Package"])
for s, p in zip(S, G["phases"]): s3.append([p["name"], s["fp"], "; ".join(f"{f} by {b}" for f, b in s["reviews"]) or "-", s["period"], s["approval"] or "-", s["package"] or "-"])
s4 = wb.create_sheet("Checks"); s4.append(["Phase", "Item", "Recomputed", "Expected (hand)", "Check"])
for r in rows: s4.append(r)
for sh, widths in ((s2, [44, 8] + [18] * 14), (s3, [44, 14, 30, 10, 14, 10]), (s4, [44, 40, 40, 40, 10])):
    for i, w in enumerate(widths): sh.column_dimensions[chr(65 + i)].width = w
    for cc in sh[1]: cc.font = Font(bold=True, color="FFFFFF"); cc.fill = PatternFill("solid", fgColor="1F3864")
    for row in sh.iter_rows(min_row=2):
        for cc in row: cc.alignment = Alignment(wrap_text=True, vertical="top")
for row in s4.iter_rows(min_row=2): row[4].fill = PatternFill("solid", fgColor="C6EFCE" if row[4].value == "OK" else "FFC7CE")
o = root / "docs/architecture/v2-m6c-golden-workings.xlsx"; wb.save(o)
print(f"wrote {o.relative_to(root)}; {len(rows)} checks; differences: {bad}")
for r in rows:
    if r[4] == "DIFFERS": print("DIFFERS", r)
sys.exit(1 if bad else 0)
