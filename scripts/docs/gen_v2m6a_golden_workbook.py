#!/usr/bin/env python3
"""Writes docs/architecture/v2-m6a-golden-workings.xlsx from tests/golden/v2-m6a-review.json: every flag of the V2-M6a golden review checklist
(DEC-068) recomputed here from the hand-written postings with Python's Decimal and calendar dates, independently of the platform, next to the
hand-written expected flags with a check column. For the named accountant's review (DEC-054). Exits 1 if anything differs."""
import datetime as dt, json, pathlib, sys
from decimal import Decimal as D, ROUND_HALF_UP
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

root = pathlib.Path(__file__).resolve().parents[2]
G = json.loads((root / "tests/golden/v2-m6a-review.json").read_text())
V = json.loads((root / G["linesFrom"]).read_text())["version"]
S, E = G["settings"], G["expected"]
day = dt.date.fromisoformat
rows, bad = [], 0
def check(step, item, got, exp):
    global bad
    ok = got == exp
    bad += not ok; rows.append([step, item, json.dumps(got, sort_keys=True), json.dumps(exp, sort_keys=True), "OK" if ok else "DIFFERS"])
f2 = lambda x: str(D(x).quantize(D("0.01"), rounding=ROUND_HALF_UP))

# ── the chart the golden touches: type, subtype, V1 reporting line ──
CHART = {"1100": ("ASSET", "CURRENT_ASSET_RECEIVABLE", "BS.CURRENT_ASSETS.DEBTORS"), "1110": ("ASSET", "CURRENT_ASSET_RECEIVABLE", "BS.CURRENT_ASSETS.DEBTORS"),
         "1200": ("ASSET", "CURRENT_ASSET_BANK", "BS.CURRENT_ASSETS.CASH"), "9999": ("ASSET", "CURRENT_ASSET_OTHER", "BS.CURRENT_ASSETS.DEBTORS"),
         "2100": ("LIABILITY", "CURRENT_LIABILITY_PAYABLE", "BS.CREDITORS_WITHIN_ONE_YEAR"), "2110": ("LIABILITY", "CURRENT_LIABILITY_ACCRUAL", "BS.CREDITORS_WITHIN_ONE_YEAR"),
         "2230": ("LIABILITY", "CURRENT_LIABILITY_OTHER", "BS.CREDITORS_WITHIN_ONE_YEAR"), "3000": ("EQUITY", "SHARE_CAPITAL", "BS.CAPITAL.SHARE_CAPITAL"),
         "4000": ("INCOME", "SALES", "PL.TURNOVER"), "4100": ("INCOME", "OTHER_INCOME", "PL.OTHER_INCOME"), "6100": ("EXPENSE", "OVERHEADS", "PL.ADMIN_EXPENSES"),
         "6200": ("EXPENSE", "OVERHEADS", "PL.ADMIN_EXPENSES"), "6310": ("EXPENSE", "OVERHEADS", "PL.ADMIN_EXPENSES"), "6400": ("EXPENSE", "OVERHEADS", "PL.ADMIN_EXPENSES"),
         "6600": ("EXPENSE", "OVERHEADS", "PL.ADMIN_EXPENSES")}
DEBIT_NORMAL = {"ASSET", "EXPENSE"}

# ── postings: (ref, date, kind, [(code, debit - credit)]) ──
P = []
ob = G["openingBalance"]; P.append(("OB", ob["date"], "OPENING_BALANCE", [(k, D(d) - D(c)) for k, d, c in ob["lines"]]))
for c in G["cash"]:
    a = sum(D(x) for _, x in c["lines"])
    if c["direction"] == "RECEIPT": P.append((c["ref"], c["date"], "CASH_RECEIPT", [("1200", a)] + [(k, -D(x)) for k, x in c["lines"]]))
    else: P.append((c["ref"], c["date"], "CASH_PAYMENT", [(k, D(x)) for k, x in c["lines"]] + [("1200", -a)]))
for j in G["journals"]: P.append((j["ref"], j["date"], "MANUAL", [(k, D(d) - D(c)) for k, d, c in j["lines"]]))
i = G["invoice"]; P.append((i["ref"], i["issueDate"], "SALES_INVOICE", [("1100", sum(D(x) for _, x in i["lines"]))] + [(k, -D(x)) for k, x in i["lines"]]))
for b in G["bills"]: P.append((b["ref"], b["issueDate"], "PURCHASE_BILL", [(k, D(x)) for k, x in b["lines"]] + [("2100", -sum(D(x) for _, x in b["lines"]))]))
for ref, d, _, lines in P: check("Ledger", f"{ref} balances", f2(sum(v for _, v in lines)), "0.00")
YS, YE, PS, PE = "2026-01-01", "2026-12-31", "2025-01-01", "2025-12-31"
def bal(code, frm, to): return sum((v for _, d, _, lines in P if frm <= d <= to for k, v in lines if k == code), D(0))
natural = lambda code, v: v if CHART[code][0] in DEBIT_NORMAL else -v

out = {k: [] for k in E["controls"]}
# RC-01
for ref, d, kind, lines in P:
    amt = sum(v for _, v in lines if v > 0)
    if kind == "MANUAL" and YS <= d <= YE and amt >= D(S["unexplainedAmount"]): out["RC-01"].append({"subject": f"journal:{ref}", "figures": {"date": d, "source": kind, "amount": f2(amt)}})
# RC-02: statutory lines with accounts and the retained profit, through the V2-M1 golden map
mp = {(a, b): t for a, b, t in V["mapEntries"]}
target = lambda code: mp.get((CHART[code][2], CHART[code][1])) or mp.get((CHART[code][2], None))
def lines_of(frm, to):
    vals = {}
    for l in V["lines"]:
        if l["kind"] != "ACCOUNTS": continue
        side = l["naturalSide"]; total = D(0)
        for code in CHART:
            if target(code) != l["key"]: continue
            v = bal(code, frm, to) if l["statement"] == "PL" else bal(code, "0001-01-01", to)
            total += v if side == "DEBIT" else -v
        vals[l["key"]] = total
    vals["B.PLR"] = -sum((bal(code, "0001-01-01", to) for code in CHART if CHART[code][2].startswith("PL.")), D(0))
    return vals
cur, prv = lines_of(YS, YE), lines_of(PS, PE)
for l in V["lines"]:
    k = l["key"]
    if k not in cur: continue
    ch = cur[k] - prv[k]
    if abs(ch) >= D(S["materialAmount"]) and (prv[k] == 0 or abs(ch) / abs(prv[k]) * 100 >= D(S["materialPercent"])):
        out["RC-02"].append({"subject": f"line:{k}", "figures": {"line": k, "current": f2(cur[k]), "comparative": f2(prv[k]), "change": f2(ch), "percent": None if prv[k] == 0 else f2(abs(ch) / abs(prv[k]) * 100)}})
# RC-03 / RC-04 / RC-07
for code in sorted(CHART):
    t = CHART[code][0]
    if t in ("INCOME", "EXPENSE"):
        n = natural(code, bal(code, YS, YE))
        if n < 0: out["RC-03"].append({"subject": f"account:{code}", "figures": {"account": code, "balance": f2(n)}})
    elif t in ("ASSET", "LIABILITY"):
        n = natural(code, bal(code, "0001-01-01", YE))
        if n < 0: out["RC-04"].append({"subject": f"account:{code}", "figures": {"account": code, "balance": f2(n)}})
    if code == "9999" and bal(code, "0001-01-01", YE) != 0: out["RC-07"].append({"subject": f"account:{code}", "figures": {"account": code, "balance": f2(bal(code, "0001-01-01", YE))}})
# RC-05: overdue by more than the setting at the year end (nothing is settled in the golden)
for party, name, docs in (("CUSTOMER", G["company"]["customer"], [G["invoice"]]), ("SUPPLIER", G["company"]["supplier"], G["bills"])):
    over = sum((sum(D(x) for _, x in d["lines"]) for d in docs if (day(YE) - day(d["dueDate"])).days > S["oldItemDays"]), D(0))
    if over: out["RC-05"].append({"subject": f"{party.lower()}:{name}", "figures": {"party": party, "currency": "GBP", "overdue": f2(over), "beyondDays": S["oldItemDays"]}})
# RC-06: the loan account (credit-positive balance, day ends)
moves = sorted((d, -v) for _, d, _, lines in P for k, v in lines if k == "2230" and YS <= d <= YE)
b, ends = -bal("2230", "0001-01-01", "2025-12-31"), {}
for d, v in moves: b += v; ends[d] = b
lowest, days, cur_b, prev = D(0), 0, -bal("2230", "0001-01-01", "2025-12-31"), day(YS)
for d in sorted(ends) + [None]:
    end = day(YE) if d is None else day(d) - dt.timedelta(days=1)
    if end >= prev:
        lowest = min(lowest, cur_b)
        if cur_b < 0: days += (end - prev).days + 1
    if d is not None: cur_b, prev = ends[d], day(d)
flags = (["OVERDRAWN_AT_END"] if cur_b < 0 else []) + (["OVERDRAWN_IN_PERIOD"] if lowest < 0 else [])
out["RC-06"].append({"subject": "loan:2230", "figures": {"closing": f2(cur_b), "lowest": f2(lowest), "daysOverdrawn": days, "flags": flags}})
# RC-08: bank ledger accounts with a balance and no bank account set up
for code in sorted(CHART):
    if CHART[code][1] == "CURRENT_ASSET_BANK" and bal(code, "0001-01-01", YE) != 0:
        out["RC-08"].append({"subject": f"account:{code}", "figures": {"account": code, "reason": "NO_BANK_ACCOUNT", "ledgerBalance": f2(bal(code, "0001-01-01", YE))}})
# RC-09: manual journals, purchase bills and cash payments of at least the setting without evidence
ev = {j["ref"] for j in G["journals"] if j["evidence"]}
types = {"MANUAL": "journal", "PURCHASE_BILL": "purchase_document", "CASH_PAYMENT": "cash_transaction"}
for ref, d, kind, lines in sorted(P, key=lambda p: (p[1], types.get(p[2], ""))):
    amt = sum(v for _, v in lines if v > 0)
    if kind in types and YS <= d <= YE and amt >= D(S["evidenceAmount"]) and ref not in ev:
        out["RC-09"].append({"subject": f"{types[kind]}:{ref}", "figures": {"type": types[kind], "date": d, "amount": f2(amt)}})

for ctl, exp in E["controls"].items():
    check("Counts", ctl, len(out[ctl]), E["counts"][ctl])
    for g, e in zip(out[ctl], exp): check(ctl, e["subject"], g, e)
check("Workings", "profit 2026", f2(-sum((bal(c, YS, YE) for c in CHART if CHART[c][2].startswith("PL.")), D(0))), "22550.00")

wb = Workbook(); ws = wb.active; ws.title = "Read me"
for line in ["V2-M6a golden review checklist - workings (SYNTHETIC)", "", G["purpose"], "", G["disclaimer"], "",
             "Every flag and figure is recomputed by this workbook's generator from the hand-written postings with Python Decimal and calendar dates, independently of the platform, and compared with the hand-written expected flags. The platform's API test matches the same flags exactly.",
             "", f"Settings: {json.dumps(S)}", "", *[f"{k}: {v}" for k, v in G["workings"].items()]]:
    ws.append([line])
ws.column_dimensions["A"].width = 150; ws["A1"].font = Font(bold=True, size=14)
for r in ws.iter_rows():
    for cc in r: cc.alignment = Alignment(wrap_text=True, vertical="top")
s1 = wb.create_sheet("Ledger"); s1.append(["Ref", "Date", "Source", "Account", "Debit", "Credit"])
for ref, d, kind, lines in sorted(P, key=lambda p: p[1]):
    for k, v in lines: s1.append([ref, d, kind, k, float(max(v, D(0))), float(max(-v, D(0)))])
s2 = wb.create_sheet("Flags"); s2.append(["Control", "Subject", "Figures"])
for ctl, items in out.items():
    for it in items: s2.append([ctl, it["subject"], json.dumps(it["figures"])])
s3 = wb.create_sheet("Checks"); s3.append(["Step", "Item", "Recomputed", "Expected (hand)", "Check"])
for r in rows: s3.append(r)
for sh, widths in ((s1, [10, 12, 18, 10, 14, 14]), (s2, [10, 34, 110]), (s3, [12, 34, 70, 70, 10])):
    for i, w in enumerate(widths): sh.column_dimensions["ABCDEF"[i]].width = w
    for cc in sh[1]: cc.font = Font(bold=True, color="FFFFFF"); cc.fill = PatternFill("solid", fgColor="1F3864")
    for row in sh.iter_rows(min_row=2):
        for cc in row: cc.alignment = Alignment(wrap_text=True, vertical="top")
for row in s3.iter_rows(min_row=2): row[4].fill = PatternFill("solid", fgColor="C6EFCE" if row[4].value == "OK" else "FFC7CE")
o = root / "docs/architecture/v2-m6a-golden-workings.xlsx"; wb.save(o)
print(f"wrote {o.relative_to(root)}; {len(rows)} checks; differences: {bad}")
for r in rows:
    if r[4] == "DIFFERS": print("DIFFERS", r)
sys.exit(1 if bad else 0)
