#!/usr/bin/env python3
"""Regenerates docs/architecture/v1-compliance-matrix.md and .xlsx from v1-compliance-matrix.json (the source of truth).
Usage: python3 scripts/docs/gen_v1_matrix.py [--commit SHA] [--ci URL]   (needs openpyxl for the workbook)
Statuses are changed only in the JSON, together with the evidence (files and passing tests) that justifies the change.
Every row carries an owner, dependencies, acceptance criteria and deferred scope (DEC-007 / second batch of 2026-10-09)."""
import argparse, collections, datetime, json, pathlib

root = pathlib.Path(__file__).resolve().parents[2]
src = root / "docs/architecture/v1-compliance-matrix.json"
ap = argparse.ArgumentParser(); ap.add_argument("--commit", default=""); ap.add_argument("--ci", default="")
a = ap.parse_args()
d = json.loads(src.read_text())
names = d["statuses"]
rows = d["rows"]
cnt = collections.Counter(r["s"] for r in rows)
total = len(rows); applicable = total - cnt["N"]
ms = collections.OrderedDict((k, [r for r in rows if r["m"] == k]) for k in d["milestones"])
esc = lambda s: (s or "").replace("|", "\\|").replace("\n", " ")
mt = lambda k: d["milestones"][k]["title"]

md = [f"# {d['title']} (plan version {d.get('version', 1)})", "",
      f"Source requirements: {d['source']}.", "",
      "Generated from `v1-compliance-matrix.json`" + (f" at commit `{a.commit}`" if a.commit else "") + (f"; CI: <{a.ci}>" if a.ci else "") + ". **Statuses change only with evidence (code and passing tests); a requirement is never marked implemented because a table or flag exists.** Owners are roles, not people: *Engineering* builds, *product owner* decides scope and policy, *qualified tax reviewer* / *DPO / legal* verify rules the product cannot verify itself.", "",
      "## Plan history (append-only)", ""] + [f"* {h}" for h in d.get("history", [])] + ["", "## Summary", "", "| Status | Rows |", "|---|---|"]
for k in "IPMN": md.append(f"| {names[k]} | {cnt[k]}" + (f" ({cnt[k]*100//applicable}% of applicable)" if k != "N" and applicable else "") + " |")
md += [f"| **Total** | **{total}** |", "", "## Milestones", "", "| Milestone | Scope | Status | Owner | Depends on | Rows | Implemented |", "|---|---|---|---|---|---|---|"]
for k, rs in ms.items():
    m = d["milestones"][k]
    md.append(f"| {k} | {esc(m['title'])} | {esc(m['status'])} | {esc(m['owner'])} | {esc('; '.join(m['depends']) or '-')} | {len(rs)} | {sum(1 for r in rs if r['s'] == 'I')} |")
md += ["", d["acceptance"], ""]
for k, rs in ms.items():
    m = d["milestones"][k]
    md += [f"## {k} - {m['title']}", "", f"**Status:** {m['status']}  ", f"**Owner:** {m['owner']}  ", f"**Depends on:** {'; '.join(m['depends']) or '-'}  ", f"**Acceptance:** {m['acceptance']}  ", f"**Deferred scope:** {m['deferred']}", "",
           "| ID | Reference | Requirement | Status | Owner | Dependencies | Acceptance criteria | Deferred scope | Evidence | Notes |", "|---|---|---|---|---|---|---|---|---|---|"]
    for r in rs: md.append(f"| **{r['id']}** | {esc(r['ref'])} | {esc(r['req'])} | **{names[r['s']]}** | {esc(r.get('owner'))} | {esc(r.get('deps'))} | {esc(r.get('accept'))} | {esc(r.get('deferred'))} | {esc(r['ev'])} | {esc(r['note'])} |")
    md.append("")
(root / "docs/architecture/v1-compliance-matrix.md").write_text("\n".join(md))

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print("openpyxl missing: workbook not written"); raise SystemExit(0)
wb = Workbook(); ws = wb.active; ws.title = "V1 matrix"
ws.append(["ID", "Reference", "Requirement", "Milestone", "Status", "Owner", "Dependencies", "Acceptance criteria", "Deferred scope", "Evidence", "Notes"])
fills = {"I": "C6EFCE", "P": "FFEB9C", "M": "FFC7CE", "N": "D9D9D9"}
for r in rows:
    ws.append([r["id"], r["ref"], r["req"], r["m"], names[r["s"]], r.get("owner", ""), r.get("deps", ""), r.get("accept", ""), r.get("deferred", ""), r["ev"], r["note"]]); ws.cell(ws.max_row, 5).fill = PatternFill("solid", fgColor=fills[r["s"]])
for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3864")
for i, w in enumerate([14, 22, 60, 12, 24, 28, 28, 70, 36, 70, 50], 1): ws.column_dimensions[get_column_letter(i)].width = w
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "B2"; ws.auto_filter.ref = ws.dimensions
s = wb.create_sheet("Milestones"); s.append(["Milestone", "Scope", "Status", "Owner", "Depends on", "Acceptance", "Deferred scope", "Rows", "Implemented"])
for k, rs in ms.items():
    m = d["milestones"][k]; s.append([k, m["title"], m["status"], m["owner"], "; ".join(m["depends"]) or "-", m["acceptance"], m["deferred"], len(rs), sum(1 for r in rs if r["s"] == "I")])
for c in s[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3864")
for i, w in enumerate([22, 60, 40, 36, 44, 40, 50, 8, 12], 1): s.column_dimensions[get_column_letter(i)].width = w
for row in s.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
sm = wb.create_sheet("Summary"); sm.append(["Status", "Rows"])
for k in "IPMN": sm.append([names[k], cnt[k]])
sm.append(["Total", total]); sm.append([]); sm.append(["Plan version", d.get("version", 1)]); sm.append(["Commit", a.commit]); sm.append(["CI", a.ci]); sm.append(["Generated", datetime.date.today().isoformat()])
h = wb.create_sheet("Plan history"); h.append(["Entry"])
for x in d.get("history", []): h.append([x])
h.column_dimensions["A"].width = 160
sm.column_dimensions["A"].width = 22; sm.column_dimensions["B"].width = 80
wb.save(root / "docs/architecture/v1-compliance-matrix.xlsx")
print(total, dict(cnt))
