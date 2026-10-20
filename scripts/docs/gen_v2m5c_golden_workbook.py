#!/usr/bin/env python3
"""Writes docs/architecture/v2-m5c-golden-workings.xlsx from tests/golden/v2-m5c-cash-flow.json: every line, total, cash balance and proof of
the V2-M5c golden cash flow statement (DEC-067) recomputed here with Python's Decimal from the hand-written postings, independently of the
platform, next to the hand-written expected figures with a check column; plus the direct-method check that operating cash equals receipts less
payments. For the named accountant's review (DEC-054). Exits 1 if anything differs."""
import json, pathlib, sys
from decimal import Decimal as D
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

root = pathlib.Path(__file__).resolve().parents[2]
G = json.loads((root / "tests/golden/v2-m5c-cash-flow.json").read_text())
CF, E = G["cashFlow"], G["expected"]
rows, bad = [], 0
def check(step, item, got, exp):
    global bad
    ok = got == exp
    bad += not ok; rows.append([step, item, str(got), str(exp), "OK" if ok else "DIFFERS"])
f2 = lambda x: f"{D(x):.2f}"

# ── the chart the golden uses: reporting line and subtype of each account code it touches (V1 default chart, plus 6950) ──
CHART = {"0030": ("BS.FIXED_ASSETS.TANGIBLE", "FIXED_ASSET_TANGIBLE"), "0031": ("BS.FIXED_ASSETS.TANGIBLE", "FIXED_ASSET_DEPRECIATION"),
         "1000": ("BS.CURRENT_ASSETS.STOCK", "CURRENT_ASSET_STOCK"), "1110": ("BS.CURRENT_ASSETS.DEBTORS", "CURRENT_ASSET_RECEIVABLE"),
         "1200": ("BS.CURRENT_ASSETS.CASH", "CURRENT_ASSET_BANK"), "2110": ("BS.CREDITORS_WITHIN_ONE_YEAR", "CURRENT_LIABILITY_ACCRUAL"),
         "2220": ("BS.CREDITORS_WITHIN_ONE_YEAR", "CURRENT_LIABILITY_TAX"), "2300": ("BS.CREDITORS_AFTER_ONE_YEAR", "LONG_TERM_LIABILITY"),
         "3000": ("BS.CAPITAL.SHARE_CAPITAL", "SHARE_CAPITAL"), "3200": ("BS.CAPITAL.RESERVES", "RETAINED_EARNINGS"), "3300": ("BS.CAPITAL.RESERVES", "DIVIDENDS"),
         "4000": ("PL.TURNOVER", "SALES"), "4900": ("PL.INTEREST_RECEIVABLE", "INTEREST_INCOME"), "5000": ("PL.COST_OF_SALES", "COST_OF_SALES"),
         "6100": ("PL.ADMIN_EXPENSES", "OVERHEADS"), "6900": ("PL.DEPRECIATION", "DEPRECIATION"), "6950": ("PL.OTHER_EXPENSES", "OTHER_EXPENSE"),
         "7000": ("PL.INTEREST_PAYABLE", "INTEREST_EXPENSE"), "7100": ("PL.TAXATION", "TAX_EXPENSE")}
def target(code):
    line, sub = CHART[code]
    m = {(a, b): t for a, b, t in CF["map"]}
    return m.get((line, sub)) or m.get((line, None))
is_pl = lambda code: CHART[code][0].startswith("PL.")

# ── journals: (date, [(code, debit - credit)]); cash transactions move 1200 ──
journals = []
for p in G["cash"]:
    amt = sum(D(a) for _, a in p["lines"])
    if p["direction"] == "RECEIPT": journals.append((p["date"], [("1200", amt)] + [(k, -D(a)) for k, a in p["lines"]]))
    else: journals.append((p["date"], [(k, D(a)) for k, a in p["lines"]] + [("1200", -amt)]))
for j in G["journals"]: journals.append((j["date"], [(k, D(dr) - D(cr)) for k, dr, cr in j["lines"]]))
for d, lines in journals: check("Ledger", f"journal {d} balances", f2(sum(v for _, v in lines)), "0.00")
# the 2025 close is a closing journal: left out of every figure here (DEC-060), so it is not simulated

classes = {c["key"]: c for c in CF["classes"]}
def year(start, end):
    eff, cash_part, gross_in, gross_out = {}, {}, {}, {}
    cash_move = D(0)
    for d, lines in journals:
        if not (start <= d <= end): continue
        touches_cash = any(target(k) == "CASH" for k, _ in lines)
        for k, v in lines:
            if target(k) == "CASH": cash_move += v; continue
            eff[k] = eff.get(k, D(0)) - v
            if touches_cash:
                cash_part[k] = cash_part.get(k, D(0)) - v
                if -v > 0: gross_in[k] = gross_in.get(k, D(0)) - v
                else: gross_out[k] = gross_out.get(k, D(0)) - v
    opening = sum(v for d, lines in journals if d < start for k, v in lines if target(k) == "CASH")
    out = {}
    out["PROFIT"] = sum((e for k, e in eff.items() if is_pl(k) and target(k) is None), D(0))
    for key, c in classes.items():
        members = [k for k in eff if target(k) == key]
        E_ = sum((eff[k] for k in members), D(0)); C_ = sum((cash_part.get(k, D(0)) for k in members), D(0))
        if c["section"] == "OPERATING": out[key] = E_
        else:
            out[f"ADJ:{key}"] = E_ - C_
            if c.get("inflowName"):
                out[f"{key}:IN"] = sum((gross_in.get(k, D(0)) for k in members), D(0)); out[f"{key}:OUT"] = sum((gross_out.get(k, D(0)) for k in members), D(0))
            else: out[key] = C_
    out["_accounts"] = {k: (eff[k], cash_part.get(k, D(0))) for k in eff}
    out["_cash"] = (opening, cash_move, opening + cash_move)
    out["_full_profit"] = sum((e for k, e in eff.items() if is_pl(k)), D(0))
    return out

cur, cmp_ = year("2026-01-01", "2026-12-31"), year("2025-01-01", "2025-12-31")
for line in E["lines"]:
    k = line["key"]
    check("Line", f"{k} amount", f2(cur[k]), line["amount"]); check("Line", f"{k} comparative", f2(cmp_[k]), line["comparative"])
    for code, amt in line.get("accounts", []):
        e, c = cur["_accounts"][code]
        got = e - c if k.startswith("ADJ:") else e
        check("Line accounts", f"{k} {code}", f2(got), amt)
for k in E["omitted"]: check("Omitted (zero both years)", k, (f2(cur[k]), f2(cmp_[k])), ("0.00", "0.00"))
listed = {l["key"] for l in E["lines"]} | set(E["omitted"])
for k in cur:
    if not k.startswith("_"): check("Every line accounted for", k, k in listed, True)
section_of = lambda k: "OPERATING" if k == "PROFIT" or k.startswith("ADJ:") else classes[k.split(":")[0]]["section"]
for sec, t in E["totals"].items():
    for which, y in (("amount", cur), ("comparative", cmp_)):
        check("Total", f"{sec} {which}", f2(sum(v for k, v in y.items() if not k.startswith("_") and section_of(k) == sec)), t[which])
for which, y in (("amount", cur), ("comparative", cmp_)):
    o, m, c = y["_cash"]
    check("Cash", f"net change {which}", f2(m), E["netChange"][which]); check("Cash", f"opening {which}", f2(o), E["opening"][which]); check("Cash", f"closing {which}", f2(c), E["closing"][which])
    total = sum(v for k, v in y.items() if not k.startswith("_"))
    check("Proof", f"statement equals cash movement ({which})", f2(total - m), "0.00")
    check("Profit", f"full profit ({which})", f2(y["_full_profit"]), E["fullProfit"][which])
check("Proof", "cash movement 2026", f2(cur["_cash"][1]), E["proof"]["cashMovement"])
# the direct method, from the cash journals' operating counterparts
def direct(start, end):
    t = D(0)
    for d, lines in journals:
        if not (start <= d <= end) or not any(target(k) == "CASH" for k, _ in lines): continue
        for k, v in lines:
            tg = target(k)
            if tg != "CASH" and (tg is None or classes[tg]["section"] == "OPERATING"): t -= v
    return t
check("Direct method", "operating cash 2026", f2(direct("2026-01-01", "2026-12-31")), E["totals"]["OPERATING"]["amount"])
check("Direct method", "operating cash 2025", f2(direct("2025-01-01", "2025-12-31")), E["totals"]["OPERATING"]["comparative"])

wb = Workbook(); ws = wb.active; ws.title = "Read me"
for line in ["V2-M5c golden cash flow statement - workings (SYNTHETIC)", "", G["purpose"], "", G["disclaimer"], "",
             "Every line, total and cash balance is recomputed by this workbook's generator with Python Decimal from the hand-written postings, independently of the platform, and compared with the hand-written expected statement. The direct-method check recomputes operating cash from the cash journals. The platform's API test matches the same statement exactly.",
             "", *[f"{k}: {v}" for k, v in G["workings"].items()]]:
    ws.append([line])
ws.column_dimensions["A"].width = 150; ws["A1"].font = Font(bold=True, size=14)
for r in ws.iter_rows():
    for cc in r: cc.alignment = Alignment(wrap_text=True, vertical="top")
s1 = wb.create_sheet("Ledger"); s1.append(["Date", "Account", "Debit", "Credit", "Cash-flow target"])
for d, lines in sorted(journals, key=lambda j: j[0]):
    for k, v in lines: s1.append([d, k, float(max(v, D(0))), float(max(-v, D(0))), target(k) or "profit"])
s2 = wb.create_sheet("Statement"); s2.append(["Section", "Line", "Caption", "2026", "2025"])
for l in E["lines"]: s2.append([l["section"], l["key"], l["name"], float(cur[l["key"]]), float(cmp_[l["key"]])])
for sec, t in E["totals"].items(): s2.append([sec, "TOTAL", t["name"], float(D(t["amount"])), float(D(t["comparative"]))])
for k in ("netChange", "opening", "closing"): s2.append(["CASH", k, E[k]["name"], float(D(E[k]["amount"])), float(D(E[k]["comparative"]))])
s3 = wb.create_sheet("Accounts"); s3.append(["Account", "Target", "E 2026", "C 2026", "N 2026", "E 2025", "C 2025", "N 2025"])
for code in sorted(set(cur["_accounts"]) | set(cmp_["_accounts"])):
    a, b = cur["_accounts"].get(code, (D(0), D(0))), cmp_["_accounts"].get(code, (D(0), D(0)))
    s3.append([code, target(code) or "profit", float(a[0]), float(a[1]), float(a[0] - a[1]), float(b[0]), float(b[1]), float(b[0] - b[1])])
s4 = wb.create_sheet("Checks"); s4.append(["Step", "Item", "Recomputed", "Expected (hand)", "Check"])
for r in rows: s4.append(r)
for sh, widths in ((s1, [12, 10, 14, 14, 16]), (s2, [12, 16, 48, 14, 14]), (s3, [10, 14, 12, 12, 12, 12, 12, 12]), (s4, [26, 44, 30, 30, 10])):
    for i, w in enumerate(widths): sh.column_dimensions["ABCDEFGH"[i]].width = w
    for cc in sh[1]: cc.font = Font(bold=True, color="FFFFFF"); cc.fill = PatternFill("solid", fgColor="1F3864")
    for row in sh.iter_rows(min_row=2):
        for cc in row:
            cc.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(cc.value, float): cc.number_format = "#,##0.00;-#,##0.00"
for row in s4.iter_rows(min_row=2): row[4].fill = PatternFill("solid", fgColor="C6EFCE" if row[4].value == "OK" else "FFC7CE")
out = root / "docs/architecture/v2-m5c-golden-workings.xlsx"; wb.save(out)
print(f"wrote {out.relative_to(root)}; {len(rows)} checks; differences: {bad}")
for r in rows:
    if r[4] == "DIFFERS": print("DIFFERS", r)
sys.exit(1 if bad else 0)
