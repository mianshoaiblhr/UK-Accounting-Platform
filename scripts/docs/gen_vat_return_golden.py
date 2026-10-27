#!/usr/bin/env python3
"""Writes docs/architecture/v5-m1-golden-workings.xlsx: the golden workings of the synthetic quarter in tests/api/vat-return.test.ts (V5-M1,
DEC-090), for review by the named accountant (the DEC-054 pattern). Every box is a formula over the journal lines, so a reviewer can trace each
figure; the expected column holds what the platform returned in the test. The rate (20%) and the box mappings are SYNTHETIC test fixtures
verified in the test by a second person, not the real UK rate rules. Usage: python3 scripts/docs/gen_vat_return_golden.py (needs openpyxl)."""
import pathlib
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

root = pathlib.Path(__file__).resolve().parents[2]
OUT = root / "docs/architecture/v5-m1-golden-workings.xlsx"
HEAD = PatternFill("solid", fgColor="1F3A5F")
INPUT = PatternFill("solid", fgColor="FFF4CC")
EXPECT = PatternFill("solid", fgColor="E8F1E4")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
PERIOD = ("2026-01-01", "2026-03-31")

# (journal, journal date, tax point, description, account, VAT code, side, net, rate %, net boxes, VAT box)
LINES = [
    ("J1", "2026-02-10", "2026-02-10", "Synthetic sale", "4000 Sales", "SX", "SALES", 1000.00, 20, "6", "1"),
    ("J2", "2026-02-20", "2026-02-20", "Synthetic purchase", "6310 Expenses", "PX", "PURCHASES", 400.00, 20, "7", "4"),
    ("J3", "2026-04-05", "2026-03-30", "Synthetic sale entered after the period", "4000 Sales", "SX", "SALES", 100.00, 20, "6", "1"),
]


def header(ws, row, titles, widths):
    for i, (t, w) in enumerate(zip(titles, widths), start=1):
        c = ws.cell(row=row, column=i, value=t)
        c.font, c.fill, c.alignment, c.border = Font(bold=True, color="FFFFFF"), HEAD, Alignment(wrap_text=True, vertical="top"), BOX
        ws.column_dimensions[get_column_letter(i)].width = w


wb = Workbook()
about = wb.active
about.title = "About"
rows = [
    ("V5-M1 golden workings: the VAT return of a synthetic quarter (DEC-090)",),
    ("",),
    ("Period", f"{PERIOD[0]} to {PERIOD[1]} (quarters ending March, June, September, December)"),
    ("Source", "tests/api/vat-return.test.ts; the platform's figures are in the 'Expected' column of 'Boxes'"),
    ("Status", "FOR REVIEW by the named accountant (DEC-054 pattern). Not a real return; nothing is filed (DEC-082, DEC-086)"),
    ("Rates", "SYNTHETIC: 20% on the test codes SX (sales: net box 6, VAT box 1) and PX (purchases: net box 7, VAT box 4), entered by one person and verified by another in the test"),
    ("Yellow cells", "Inputs (the journal lines)"),
    ("Green cells", "What the platform returned in the test"),
    ("What to check", "That each box equals the formula over the lines; that the reconciliation difference is the late sale's VAT; that both flags are raised for the reason given"),
]
for r in rows:
    about.append(r)
about["A1"].font = Font(bold=True, size=13)
about.column_dimensions["A"].width = 18
about.column_dimensions["B"].width = 120

ws = wb.create_sheet("Lines")
titles = ["Journal", "Journal date", "Tax point", "Description", "Account", "VAT code", "Side", "Net", "Rate %", "VAT", "Net box", "VAT box", "Tax point in period", "Journal date in period"]
header(ws, 1, titles, [9, 12, 12, 38, 16, 9, 11, 11, 8, 11, 9, 9, 14, 14])
for i, l in enumerate(LINES, start=2):
    for j, v in enumerate(l[:9], start=1):
        c = ws.cell(row=i, column=j, value=v)
        c.fill, c.border = INPUT, BOX
    ws.cell(row=i, column=10, value=f"=ROUND(H{i}*I{i}/100,2)").border = BOX
    ws.cell(row=i, column=11, value=l[9]).border = BOX
    ws.cell(row=i, column=12, value=l[10]).border = BOX
    ws.cell(row=i, column=13, value=f'=AND(C{i}>="{PERIOD[0]}",C{i}<="{PERIOD[1]}")').border = BOX
    ws.cell(row=i, column=14, value=f'=AND(B{i}>="{PERIOD[0]}",B{i}<="{PERIOD[1]}")').border = BOX
    for col in (8, 10):
        ws.cell(row=i, column=col).number_format = "#,##0.00"
last = len(LINES) + 1
R = f"$2:${last}"

bx = wb.create_sheet("Boxes")
header(bx, 1, ["Box", "What it holds", "Working", "Amount", "Expected (platform)", "Agrees"], [6, 44, 70, 14, 18, 9])
inp = lambda col: f"Lines!{col}$2:{col}${last}"
BOXES = [
    (1, "VAT due on sales", f'=SUMIFS({inp("J")},{inp("L")},"1",{inp("M")},TRUE)', "SUM of VAT on lines whose VAT box is 1 and tax point is in the period", "220.00"),
    (2, "VAT due on acquisitions from the EU", "=0", "No line maps to box 2", "0.00"),
    (3, "Total VAT due (box 1 + box 2)", "=D2+D3", "Box 1 + box 2", "220.00"),
    (4, "VAT reclaimed on purchases", f'=SUMIFS({inp("J")},{inp("L")},"4",{inp("M")},TRUE)', "SUM of VAT on lines whose VAT box is 4 and tax point is in the period", "80.00"),
    (5, "Net VAT to pay (or reclaim)", "=ABS(D4-D5)", "Box 3 - box 4; payable because box 3 is larger", "140.00"),
    (6, "Total sales excluding VAT", f'=SUMIFS({inp("H")},{inp("K")},"6",{inp("M")},TRUE)', "SUM of net on lines whose net box is 6 and tax point is in the period (exact; whole pounds are V5-M2)", "1100.00"),
    (7, "Total purchases excluding VAT", f'=SUMIFS({inp("H")},{inp("K")},"7",{inp("M")},TRUE)', "SUM of net on lines whose net box is 7 and tax point is in the period", "400.00"),
    (8, "Supplies to the EU", "=0", "No line maps to box 8", "0.00"),
    (9, "Acquisitions from the EU", "=0", "No line maps to box 9", "0.00"),
]
for i, (n, what, formula, working, expected) in enumerate(BOXES, start=2):
    bx.cell(row=i, column=1, value=n)
    bx.cell(row=i, column=2, value=what)
    bx.cell(row=i, column=3, value=working)
    bx.cell(row=i, column=4, value=formula).number_format = "#,##0.00"
    e = bx.cell(row=i, column=5, value=float(expected))
    e.number_format, e.fill = "#,##0.00", EXPECT
    bx.cell(row=i, column=6, value=f"=ROUND(D{i}-E{i},2)=0")
    for col in range(1, 7):
        bx.cell(row=i, column=col).border = BOX
        bx.cell(row=i, column=col).alignment = Alignment(wrap_text=True, vertical="top")

rc = wb.create_sheet("Reconciliation")
header(rc, 1, ["Item", "Working", "Amount", "Expected (platform)", "Agrees"], [44, 80, 14, 18, 9])
REC = [
    ("Movement on the VAT accounts by journal date", f'SUMIFS of VAT (credit-positive: sales +, purchases -) on lines dated in the period',
     f'=SUMIFS({inp("J")},{inp("G")},"SALES",{inp("N")},TRUE)-SUMIFS({inp("J")},{inp("G")},"PURCHASES",{inp("N")},TRUE)', "120.00"),
    ("VAT lines by tax point", "The same, on lines whose tax point is in the period", f'=SUMIFS({inp("J")},{inp("G")},"SALES",{inp("M")},TRUE)-SUMIFS({inp("J")},{inp("G")},"PURCHASES",{inp("M")},TRUE)', "140.00"),
    ("Difference", "Movement - VAT lines: the VAT of J3, dated after the period with a tax point inside it", "=C2-C3", "-20.00"),
]
for i, (item, working, formula, expected) in enumerate(REC, start=2):
    rc.cell(row=i, column=1, value=item)
    rc.cell(row=i, column=2, value=working)
    rc.cell(row=i, column=3, value=formula).number_format = "#,##0.00"
    e = rc.cell(row=i, column=4, value=float(expected))
    e.number_format, e.fill = "#,##0.00", EXPECT
    rc.cell(row=i, column=5, value=f"=ROUND(C{i}-D{i},2)=0")
    for col in range(1, 6):
        rc.cell(row=i, column=col).border = BOX
        rc.cell(row=i, column=col).alignment = Alignment(wrap_text=True, vertical="top")

fl = wb.create_sheet("Flags")
header(fl, 1, ["Flag key", "Raised?", "Why", "Explained in the test as"], [36, 9, 70, 60])
FLAGS = [
    ("reconciliation_difference", "Yes", "The difference on 'Reconciliation' is not zero (-20.00)", "The sale was invoiced in March and entered in April (synthetic)"),
    ("late_posting:{J3}:{VAT line}", "Yes", "J3 is dated 2026-04-05, after the period, with a tax point (2026-03-30) inside it", "The sale was invoiced in March and entered in April (synthetic)"),
    ("superseded_rate:{rateId}", "No", "Neither rate version has been superseded", ""),
    ("repayment", "No", "Box 4 (80.00) is not more than boxes 1 and 2 (220.00)", ""),
    ("zero_return", "No", "Boxes are not all zero", ""),
]
for i, row in enumerate(FLAGS, start=2):
    for j, v in enumerate(row, start=1):
        c = fl.cell(row=i, column=j, value=v)
        c.border, c.alignment = BOX, Alignment(wrap_text=True, vertical="top")

st = wb.create_sheet("Steps")
header(st, 1, ["Step", "Who (test role)", "Result in the platform"], [34, 28, 80])
STEPS = [
    ("Open the return (date 2026-02-15)", "Accountant (vat:prepare)", "Period 2026-01-01 to 2026-03-31; linked to the VAT obligation; DRAFT"),
    ("Prepare", "Accountant", "PREPARED on the fingerprint of J1 and J2"),
    ("Post J3 (dated 5 April, tax point 30 March)", "Accountant", "Fingerprint changes; the preparation lapses (DRAFT); two flags raised"),
    ("Prepare again; explain both flags", "Accountant", "PREPARED; no unexplained flag"),
    ("Review", "Partner (vat:review), not the preparer", "REVIEWED"),
    ("Client approval (2026-04-20, director contact, signed approval scanned clean)", "Partner (vat:review and period:lock)", "APPROVED; return frozen; period locked; obligation READY"),
    ("Post a sale with tax point 2026-03-15", "Accountant", "Refused: vat_period_locked"),
]
for i, row in enumerate(STEPS, start=2):
    for j, v in enumerate(row, start=1):
        c = st.cell(row=i, column=j, value=v)
        c.border, c.alignment = BOX, Alignment(wrap_text=True, vertical="top")

so = wb.create_sheet("Sign-off")
header(so, 1, ["Reviewer", "Date", "Agrees (yes/no)", "Notes"], [28, 14, 16, 70])
for col in range(1, 5):
    so.cell(row=2, column=col).border = BOX
wb.save(OUT)
print(f"wrote {OUT.relative_to(root)}")
