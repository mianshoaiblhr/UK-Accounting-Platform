#!/usr/bin/env python3
"""Writes docs/architecture/v2-m4-golden-workings.xlsx from tests/golden/v2-m4-directors.json: every loan-account statement, category total, lowest
balance, day count, review flag and related-party figure of the V2-M4 golden scenario (DEC-062) recomputed here with Python's Decimal and calendar
dates, independently of the platform, next to the hand-written expected figure with a check column. For the named accountant's review (DEC-054).
Exits 1 if any figure differs."""
import datetime as dt, json, pathlib, re, sys
from decimal import Decimal as D

root = pathlib.Path(__file__).resolve().parents[2]
G = json.loads((root / "tests/golden/v2-m4-directors.json").read_text())
day = dt.date.fromisoformat
BANK, CONTROL = {"1200"}, {"1100", "2100", "2200"}
TYPE = lambda c: "ASSET" if c[0] in "01" else "LIABILITY" if c[0] == "2" else "EQUITY" if c[0] == "3" else "INCOME" if c[0] == "4" else "EXPENSE"
LOANS = {p["account"]: k for k, p in G["people"].items()}
rows, bad = [], 0
def check(step, item, got, exp):
    global bad
    ok = (D(str(got)) == D(exp)) if isinstance(exp, str) and re.fullmatch(r"-?\d+\.\d+", exp) else got == exp
    bad += not ok; rows.append([step, item, str(got), str(exp), "OK" if ok else "DIFFERS"])

journals = []                                            # (ref, date, [(code, debit - credit)], source)
def add(ref, date, lines, source): journals.append((ref, day(date), [(k, D(v)) for k, v in lines], source))
for p in G["postings"]:
    if p["kind"] == "cash":
        amt = sum(D(a) for _, a in p["lines"])
        if p["direction"] == "RECEIPT": add(p["ref"], p["date"], [("1200", amt)] + [(k, -D(a)) for k, a in p["lines"]], "CASH_RECEIPT")
        else: add(p["ref"], p["date"], [(k, D(a)) for k, a in p["lines"]] + [("1200", -amt)], "CASH_PAYMENT")
    else: add(p["ref"], p["date"], [(k, D(dr) - D(cr)) for k, dr, cr in p["lines"]], "MANUAL")
docs = {d["ref"]: d for d in G["relatedCompany"]["documents"]}
for d in G["relatedCompany"]["documents"]:
    k = d["kind"]
    if k == "invoice": a = sum(D(x) for _, x in d["lines"]); add(d["ref"], d["date"], [("1100", a)] + [(c, -D(x)) for c, x in d["lines"]], "SALES_INVOICE")
    if k == "creditNote": a = sum(D(x) for _, x in d["lines"]); add(d["ref"], d["date"], [(c, D(x)) for c, x in d["lines"]] + [("1100", -a)], "SALES_CREDIT_NOTE")
    if k == "receipt": add(d["ref"], d["date"], [("1200", D(d["amount"])), ("1100", -D(d["amount"]))], "SALES_RECEIPT")
    if k == "bill": a = sum(D(x) for _, x in d["lines"]); add(d["ref"], d["date"], [(c, D(x)) for c, x in d["lines"]] + [("2100", -a)], "PURCHASE_BILL")
    if k == "payment": add(d["ref"], d["date"], [("2100", D(d["amount"])), ("1200", -D(d["amount"]))], "PURCHASE_PAYMENT")
    if k == "cash": a = sum(D(x) for _, x in d["lines"]); add(d["ref"], d["date"], [(c, D(x)) for c, x in d["lines"]] + [("1200", -a)], "CASH_PAYMENT")
for ref, _, lines, _ in journals: check("Ledger", f"{ref} balances", sum(v for _, v in lines), "0.00")

def category(code):
    if code in BANK: return "CASH_AND_BANK"
    if code in LOANS: return "OTHER_LOAN_ACCOUNT"
    if code == "3300": return "DIVIDENDS"
    if TYPE(code) == "EXPENSE": return "EXPENSES"
    if TYPE(code) == "INCOME": return "INCOME"
    return "OTHER"
def statement(acct, frm, to):
    f, t = day(frm), day(to)
    bal = -sum((v for (_, d, ls, _) in journals for (k, v) in ls if k == acct and d < f), D(0))   # credit-positive
    opening = bal; moves = []
    for ref, d, ls, _ in sorted(journals, key=lambda j: j[1]):
        if not (f <= d <= t): continue
        m = -sum((v for k, v in ls if k == acct), D(0))
        if m == 0: continue
        cats = {category(k) for k, v in ls if k != acct and k != "2200"}
        bal += m; moves.append((ref, d, m, cats.pop() if len(cats) == 1 else "MIXED", bal))
    # day-end balances within the range
    lowest, lowest_on, days, b, i, x = None, None, 0, opening, 0, f
    while x <= t:
        while i < len(moves) and moves[i][1] == x: b = moves[i][4]; i += 1
        if lowest is None or b < lowest: lowest, lowest_on = b, x
        days += b < 0; x += dt.timedelta(days=1)
    return opening, moves, bal, lowest, lowest_on, days
def flags(opening, moves, closing, lowest, th):
    f = (["OVERDRAWN_AT_END"] if closing < 0 else []) + (["OVERDRAWN_IN_PERIOD"] if lowest < 0 else [])
    large = [m[0] for m in moves if th and abs(m[2]) >= D(th["largeMovement"])]
    if th and large: f.append("LARGE_MOVEMENT")
    if th and lowest < 0 and -lowest >= D(th["overdrawn"]): f.append("OVERDRAWN_ABOVE_THRESHOLD")
    return f, large
for k, person in G["people"].items():
    e = G["loanAccounts"][k]["year"]; step = f"Loan account {person['name']} {e['from']}..{e['to']}"
    o, mv, c, lo, lo_on, days = statement(person["account"], e["from"], e["to"])
    check(step, "opening", o, e["opening"]); check(step, "closing", c, e["closing"])
    check(step, "movements (ref, date, amount, category, balance)", [[r, d.isoformat(), f"{m:.2f}", cat, f"{b:.2f}"] for r, d, m, cat, b in mv], e["movements"])
    cats = {}
    for _, _, m, cat, _ in mv: cats[cat] = cats.get(cat, D(0)) + m
    for cat, exp in e["byCategory"].items(): check(step, f"category {cat}", cats.get(cat, D(0)), exp)
    check(step, "categories listed", sorted(cats), sorted(e["byCategory"]))
    check(step, "lowest balance", lo, e["lowest"]); check(step, "lowest on", lo_on.isoformat(), e["lowestOn"]); check(step, "days overdrawn", days, e["daysOverdrawn"])
    fl, _ = flags(o, mv, c, lo, None); check(step, "flags without thresholds", fl, e["flags"])
    check(step, "not evaluated without thresholds", ["LARGE_MOVEMENT", "OVERDRAWN_ABOVE_THRESHOLD"], e["notEvaluated"])
    fl, large = flags(o, mv, c, lo, G["thresholds"]); check(step, "flags with thresholds", fl, e["flagsWithThresholds"]); check(step, "large movements", large, e["largeMovements"])
h = G["loanAccounts"]["A"]["secondHalf"]; step = f"Loan account Alex Director {h['from']}..{h['to']}"
o, mv, c, lo, lo_on, days = statement("2230", h["from"], h["to"])
check(step, "opening", o, h["opening"]); check(step, "closing", c, h["closing"]); check(step, "movements", [m[0] for m in mv], h["movements"])
check(step, "lowest balance", lo, h["lowest"]); check(step, "lowest on", lo_on.isoformat(), h["lowestOn"]); check(step, "days overdrawn", days, h["daysOverdrawn"])
check(step, "flags with thresholds", flags(o, mv, c, lo, G["thresholds"])[0], h["flags"])

# related parties: the counterparty's sub-ledger and cash journals (the loan account's own cash is a loan movement, never cash paid or received)
R = G["relatedReport"]; f, t = day(R["from"]), day(R["to"])
company_refs = set(docs)
def party(name):
    refs = company_refs if name == G["relatedCompany"]["name"] else set()
    js = [j for j in journals if j[0] in refs]
    noncontrol = lambda ls: [(k, v) for k, v in ls if k not in CONTROL and k not in BANK and k not in LOANS]
    sales = -sum((v for (r, d, ls, s) in js if s in ("SALES_INVOICE", "SALES_CREDIT_NOTE") and f <= d <= t for k, v in noncontrol(ls)), D(0))
    purchases = sum((v for (r, d, ls, s) in js if s in ("PURCHASE_BILL", "PURCHASE_DEBIT_NOTE") and f <= d <= t for k, v in noncontrol(ls)), D(0))
    paid = sum((v for (r, d, ls, s) in js if s == "CASH_PAYMENT" and f <= d <= t for k, v in noncontrol(ls)), D(0))
    received = -sum((v for (r, d, ls, s) in js if s == "CASH_RECEIPT" and f <= d <= t for k, v in noncontrol(ls)), D(0))
    owed_by = sum((v for (r, d, ls, s) in js if d <= t for k, v in ls if k == "1100"), D(0))
    owed_to = -sum((v for (r, d, ls, s) in js if d <= t for k, v in ls if k == "2100"), D(0))
    sj = [r for (r, d, ls, s) in js if s in ("SALES_INVOICE", "SALES_CREDIT_NOTE")]; pj = [r for (r, d, ls, s) in js if s in ("PURCHASE_BILL", "PURCHASE_DEBIT_NOTE")]
    cj = [r for (r, d, ls, s) in js if s in ("CASH_PAYMENT", "CASH_RECEIPT")]
    return sales, purchases, paid, received, owed_by, owed_to, sj, pj, cj
for name, e in R["parties"].items():
    step = f"Related party {name}"
    sales, purchases, paid, received, owed_by, owed_to, sj, pj, cj = party(name)
    for item, got in [("sales", sales), ("purchases", purchases), ("cash paid", paid), ("cash received", received), ("owed by", owed_by), ("owed to", owed_to)]:
        check(step, item, got, e[{"cash paid": "cashPaid", "cash received": "cashReceived", "owed by": "owedBy", "owed to": "owedTo"}.get(item, item)])
    if "salesJournals" in e: check(step, "sales journals", sj, e["salesJournals"]); check(step, "purchase journals", pj, e["purchaseJournals"]); check(step, "cash journals", cj, e["cashJournals"])
    person = next((p for p in G["people"].values() if p["name"] == name), None)
    if person:
        _, _, c, lo, _, _ = statement(person["account"], R["from"], R["to"])
        check(step, "loan account closing", c, e["loanAccount"]["closing"]); check(step, "loan account lowest", lo, e["loanAccount"]["lowest"])
    else: check(step, "loan account", None, e["loanAccount"])

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:
    print("openpyxl missing"); sys.exit(1)
wb = Workbook(); ws = wb.active; ws.title = "Read me"
for line in ["V2-M4 golden director and shareholder loan accounts and related parties - workings (SYNTHETIC)", "", G["purpose"], "", G["disclaimer"], "",
             "Every figure is recomputed by this workbook's generator with Python Decimal and calendar dates, independently of the platform, and compared with the hand-written expected figure. The platform's API test matches the same figures to the cent.",
             "", "Rules recomputed (DEC-062): loan balances are credit-positive (negative = overdrawn); each movement takes the category of its other lines, VAT ignored, MIXED when they differ; the lowest balance and days overdrawn use day-end balances in the range; flags with no threshold set are 'not evaluated'.",
             "Related parties: sales and purchases are the non-control lines of the counterparty's documents; cash paid and received the non-control lines of its cash transactions (never a loan account's own lines); owed by and owed to the control-account lines of all its sub-ledger journals up to the end."]:
    ws.append([line])
ws.column_dimensions["A"].width = 150; ws["A1"].font = Font(bold=True, size=14)
for r in ws.iter_rows():
    for c in r: c.alignment = Alignment(wrap_text=True, vertical="top")
s2 = wb.create_sheet("Ledger"); s2.append(["Journal", "Date", "Source", "Account", "Debit", "Credit"])
for ref, d, ls, src in sorted(journals, key=lambda j: j[1]):
    for k, v in ls: s2.append([ref, d.isoformat(), src, k, float(max(v, D(0))), float(max(-v, D(0)))])
for i, w in enumerate([10, 12, 20, 10, 14, 14], 1): s2.column_dimensions["ABCDEF"[i - 1]].width = w
s3 = wb.create_sheet("Checks"); s3.append(["Step", "Figure", "Recomputed", "Expected (hand)", "Check"])
for r in rows: s3.append(r)
for i, w in enumerate([48, 40, 70, 70, 10], 1): s3.column_dimensions["ABCDE"[i - 1]].width = w
for sh in (s2, s3):
    for c in sh[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3864")
for row in s3.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
    row[4].fill = PatternFill("solid", fgColor="C6EFCE" if row[4].value == "OK" else "FFC7CE")
for row in s2.iter_rows(min_row=2):
    for c in row[4:6]: c.number_format = '#,##0.00'
out = root / "docs/architecture/v2-m4-golden-workings.xlsx"; wb.save(out)
print(f"wrote {out.relative_to(root)}; {len(rows)} figures; differences: {bad}")
for r in rows:
    if r[4] == "DIFFERS": print("DIFFERS", r)
sys.exit(1 if bad else 0)
