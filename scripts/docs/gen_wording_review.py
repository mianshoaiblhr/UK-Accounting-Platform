#!/usr/bin/env python3
"""Writes docs/architecture/v2-draft-wording-review.xlsx from content/standards/drafts/*.json (DEC-064): every draft rule's wording, the provision it
is meant to reflect and the note for the verifier, with empty Verdict and Comments columns for the named accountant, plus the statement lines and the
map of each pack. A review aid only: verification itself happens in the platform by a holder of accounts:rule-verify who is not the importer."""
import json, pathlib
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

root = pathlib.Path(__file__).resolve().parents[2]
drafts = [json.loads(p.read_text()) for p in sorted((root / "content/standards/drafts").glob("*.json"))]
packs = [d for d in drafts if "lines" in d]
thresholds = [d for d in drafts if "sets" in d]
wb = Workbook(); ws = wb.active; ws.title = "Read me"
for line in ["Draft wording packs - review workbook (DEC-064)", "",
             "Written by engineering from its knowledge on 2026-10-19 at the product owner's instruction. Nothing here was retrieved from or checked against the FRC's or legislation.gov.uk's text: both were unreachable from the build environment.",
             "How to use it: for each rule, compare the wording with the current text of the provision named, then set Verdict (Agree / Change / Remove) and write any change in Comments. Corrections are made in the platform while the version is UNVERIFIED; verification is then done in the platform by a holder of accounts:rule-verify who did not import the pack (DEC-052).",
             "Placeholders: {{line:KEY}} a statement line, {{line:KEY:prior}} its comparative, {{fact:KEY}} a fact (text.* is company wording the preparer enters), {{table:KEY}} a table from the ledger or registers, {{year:end}} and similar the year's dates, {{directors:names}} the directors in the year.",
             "Applies to: the entity sizes a rule is used for (the size chosen for the year). Omit when empty: left out when its tables have no rows and its lines are zero.",
             "Size thresholds (added 2026-10-20, DEC-066): the 'Size thresholds' sheet lists the draft company-law limits the size test reads, one row per set and class. Check each figure, the dates each set governs and the 'prior year on first year' transitional rule against the instrument named, then set a Verdict on the MICRO row of each set. Each set is imported UNVERIFIED and verified in the platform like a standard version.", ""]:
    ws.append([line])
for p in packs: ws.append([f"{p['pack']}: {p['standard']['name']} - {p['version']['label']}. {len(p['rules'])} wording rules, {len(p['lines'])} statement lines, {len(p['mapEntries'])} map entries."])
ws.column_dimensions["A"].width = 160; ws["A1"].font = Font(bold=True, size=14)
for r in ws.iter_rows():
    for c in r: c.alignment = Alignment(wrap_text=True, vertical="top")
head = lambda sh: [setattr(c, "font", Font(bold=True, color="FFFFFF")) or setattr(c, "fill", PatternFill("solid", fgColor="1F3864")) for c in sh[1]]
for p in packs:
    short = p["standard"]["code"]
    sh = wb.create_sheet(f"{short} wording")
    sh.append(["Section", "Key", "Title", "Applies to", "Omit when empty", "Wording (template)", "Provision it reflects", "Note for the verifier", "Verdict", "Comments"])
    for r in p["rules"]:
        sh.append([r["section"], r["key"], r["title"], ", ".join(r["appliesTo"]), "yes" if r.get("omitWhenEmpty") else "", r["body"], r["source"], r["verifierNote"], "", ""])
    dv = DataValidation(type="list", formula1='"Agree,Change,Remove"', allow_blank=True); sh.add_data_validation(dv); dv.add(f"I2:I{len(p['rules']) + 1}")
    for i, w in enumerate([22, 16, 30, 14, 10, 70, 38, 48, 10, 40]): sh.column_dimensions["ABCDEFGHIJ"[i]].width = w
    head(sh)
    sl = wb.create_sheet(f"{short} lines")
    sl.append(["Statement", "Key", "Name", "Kind", "Natural side", "Parent", "Total of", "Order"])
    for l in p["lines"]:
        sl.append([l["statement"], l["key"], l["name"], l["kind"], l.get("naturalSide", ""), l.get("parentKey", ""), " ".join(("+" if t["sign"] == 1 else "-") + t["line"] for t in l.get("formula", [])), l["sort"]])
    for i, w in enumerate([10, 12, 52, 16, 12, 10, 60, 8]): sl.column_dimensions["ABCDEFGH"[i]].width = w
    head(sl)
    sm = wb.create_sheet(f"{short} map")
    sm.append(["V1 reporting line", "Account subtype (if narrowed)", "Statutory line"])
    for a, b, t in p["mapEntries"]: sm.append([a, b or "", t])
    for i, w in enumerate([34, 32, 14]): sm.column_dimensions["ABC"[i]].width = w
    head(sm)
    for sh2 in (sh, sl, sm):
        for row in sh2.iter_rows(min_row=2):
            for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
# V2-M5b (DEC-066): the draft size thresholds, one row per set and class, with a Verdict per set
for t in thresholds:
    st = wb.create_sheet("Size thresholds")
    st.append(["Set", "Governs years beginning", "Prior year on first year", "Class", "Turnover", "Balance-sheet total", "Average employees", "Source", "Note for the verifier", "Verdict", "Comments"])
    for s_ in t["sets"]:
        for cls in ("MICRO", "SMALL", "MEDIUM"):
            l = s_["limits"][cls]
            st.append([s_["label"], f"{s_['effectiveFrom']} to {s_['effectiveTo'] or 'open'}", "yes" if s_["priorYearOnFirst"] else "no", cls, float(l["turnover"]), float(l["balanceSheetTotal"]), l["employees"],
                       f"{s_['source']['provision']} ({s_['source']['url']})" if cls == "MICRO" else "", s_["verifierNote"] if cls == "MICRO" else "", "", ""])
    dv = DataValidation(type="list", formula1='"Agree,Change,Remove"', allow_blank=True); st.add_data_validation(dv); dv.add(f"J2:J{3 * len(t['sets']) + 1}")
    for i, w in enumerate([40, 24, 12, 10, 16, 16, 12, 50, 60, 10, 40]): st.column_dimensions["ABCDEFGHIJK"[i]].width = w
    for row in st.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
        for c in row[4:6]: c.number_format = "#,##0.00"
    head(st)
out = root / "docs/architecture/v2-draft-wording-review.xlsx"; wb.save(out)
print(f"wrote {out.relative_to(root)}; " + "; ".join(f"{p['pack']}: {len(p['rules'])} rules" for p in packs) + "".join(f"; {t['pack']}: {len(t['sets'])} sets" for t in thresholds))
