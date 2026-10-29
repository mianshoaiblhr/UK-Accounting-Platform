#!/usr/bin/env python3
"""Writes docs/architecture/v4-m4-golden-workings.xlsx: the hand workings of tests/golden/v4-m4-computation.json (V4-M4, DEC-101), for review by the
named accountant (the DEC-054 pattern): a profitable 12-month period across two financial years (A), a 182-day loss period carried back (B), and A
again after the carry back. Every figure is a formula; the expected cells hold the golden the unit and API tests assert. The rates and limits are
synthetic fixtures, not tax advice. Usage: python3 scripts/docs/gen_ct_computation_golden.py (needs openpyxl)."""
import json
import pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

root = pathlib.Path(__file__).resolve().parents[2]
G = json.loads((root / "tests/golden/v4-m4-computation.json").read_text())
OUT = root / "docs/architecture/v4-m4-golden-workings.xlsx"
EXPECT = PatternFill("solid", fgColor="E8F1E4")
HEAD = Font(bold=True)
MONEY = "#,##0.00"
R = G["rates"]

wb = Workbook()
about = wb.active
about.title = "About"
for r, t in enumerate([
    "V4-M4 golden workings: losses, tax by financial year and marginal relief (DEC-101)", "", G["about"],
    "Trading profit = adjusted profit - capital allowances + balancing charges. Taxable total profits = trading profit - losses brought forward used - loss carried back from the next period.",
    "Profits are apportioned to financial years by days (first part rounded to the penny, the last part the remainder).",
    "Limits per part = annual limit x days of the part / (days of the period if 12 months, else 365) / (1 + associated companies), rounded to the penny.",
    "At or under the lower limit: small profits rate; at or over the upper: main rate; between: main rate less fraction x (upper limit - profits).",
    "Green cells are the golden the tests assert; each sheet's Check row must be TRUE.",
], start=1):
    about.cell(row=r, column=1, value=t).font = Font(bold=(r == 1))
about.column_dimensions["A"].width = 160

ru = wb.create_sheet("Rates")
for i, (k, v) in enumerate([("Main rate %", R["mainRate"]), ("Small profits rate %", R["smallProfitsRate"]), ("Lower limit", R["lowerLimit"]), ("Upper limit", R["upperLimit"]),
                            ("Marginal relief fraction", R["marginalReliefFraction"]["numerator"] / R["marginalReliefFraction"]["denominator"])], start=1):
    ru.cell(row=i, column=1, value=k); ru.cell(row=i, column=2, value=float(v))
ru.column_dimensions["A"].width = 26
MAIN, SMALL, LOW, UP, FRAC = "Rates!$B$1", "Rates!$B$2", "Rates!$B$3", "Rates!$B$4", "Rates!$B$5"


def sheet(name, key, days, twelve, parts, lines):
    """lines: (label, formula or value, expected key) for the profit build-up; rows referenced by label."""
    ex = G["expected"][key]
    ws = wb.create_sheet(name)
    ws.column_dimensions["A"].width = 44
    for c in "BCDEFGHIJ": ws.column_dimensions[c].width = 15
    ws.cell(row=1, column=1, value=name).font = HEAD
    ws.cell(row=2, column=1, value="Days in the period"); ws.cell(row=2, column=2, value=days)
    ws.cell(row=3, column=1, value="Limits divide by"); ws.cell(row=3, column=2, value="=B2" if twelve else 365)
    ws.cell(row=4, column=1, value="Associated companies"); ws.cell(row=4, column=2, value=0)
    row, at, checks = 6, {}, []
    for label, f, ek in lines:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=f(at) if callable(f) else f).number_format = MONEY
        if ek and ek in ex:
            ws.cell(row=row, column=3, value=float(ex[ek])).fill = EXPECT; checks.append(f"ROUND(B{row}-C{row},2)=0")
        at[label] = f"B{row}"; row += 1
    ttp = at["Taxable total profits"]
    row += 1
    hdr = ["Financial year", "Days", "Profit", "Lower limit", "Upper limit", "Tax at rate", "Marginal relief", "Tax", "Expected tax", "Expected relief"]
    for c, h in enumerate(hdr, start=1): ws.cell(row=row, column=c, value=h).font = HEAD
    row += 1
    first = row
    for k, (fyear, d) in enumerate(parts):
        e = ex["financialYears"][k]
        ws.cell(row=row, column=1, value=fyear); ws.cell(row=row, column=2, value=d)
        ws.cell(row=row, column=3, value=f"=ROUND({ttp}*B{row}/$B$2,2)" if k < len(parts) - 1 else f"={ttp}-SUM(C{first}:C{row - 1})" if k else f"={ttp}")
        ws.cell(row=row, column=4, value=f"=ROUND({LOW}*B{row}/$B$3/(1+$B$4),2)")
        ws.cell(row=row, column=5, value=f"=ROUND({UP}*B{row}/$B$3/(1+$B$4),2)")
        ws.cell(row=row, column=6, value=f"=IF(C{row}<=D{row},ROUND(C{row}*{SMALL}/100,2),ROUND(C{row}*{MAIN}/100,2))")
        ws.cell(row=row, column=7, value=f"=IF(AND(C{row}>D{row},C{row}<E{row}),ROUND({FRAC}*(E{row}-C{row}),2),0)")
        ws.cell(row=row, column=8, value=f"=F{row}-G{row}")
        ws.cell(row=row, column=9, value=float(e[8])).fill = EXPECT; ws.cell(row=row, column=10, value=float(e[7])).fill = EXPECT
        for c in range(3, 11): ws.cell(row=row, column=c).number_format = MONEY
        checks += [f"ROUND(H{row}-I{row},2)=0", f"ROUND(G{row}-J{row},2)=0", f"ROUND(C{row}-{float(e[2])},2)=0"]
        row += 1
    ws.cell(row=row, column=1, value="Corporation tax").font = HEAD
    ws.cell(row=row, column=8, value=f"=SUM(H{first}:H{row - 1})").number_format = MONEY
    ws.cell(row=row, column=9, value=float(ex["corporationTax"])).fill = EXPECT
    checks.append(f"ROUND(H{row}-I{row},2)=0")
    ws.cell(row=row + 2, column=1, value="Check: workings equal expected").font = HEAD
    ws.cell(row=row + 2, column=2, value="=AND(" + ",".join(checks) + ")")
    return ws


A = G["periods"][0]
common_a = [
    ("Sales (4000)", 200000.0, None), ("Wages (6000)", -50000.0, None), ("Adjusted profit", lambda a: f"={a['Sales (4000)']}+{a['Wages (6000)']}", "adjustedProfit"),
    ("Main pool opening", float(A["mainPoolOpening"]), None), ("Capital allowances (WDA 18%)", lambda a: f"=ROUND({a['Main pool opening']}*18/100,2)", "allowances"),
    ("Trading profit", lambda a: f"=MAX({a['Adjusted profit']}-{a['Capital allowances (WDA 18%)']},0)", "tradingProfit"),
    ("Losses brought forward", float(A["lossOpening"]), None), ("Losses brought forward used", lambda a: f"=MIN({a['Losses brought forward']},{a['Trading profit']})", "lossesBroughtForwardUsed"),
]
sheet("A before carry back", "A1", 365, True, [(2026, 90), (2027, 275)], common_a + [
    ("Loss carried back from B", 0.0, None),
    ("Taxable total profits", lambda a: f"={a['Trading profit']}-{a['Losses brought forward used']}-{a['Loss carried back from B']}", "taxableTotalProfits"),
])
sheet("B loss period", "B", 182, False, [(2027, 91), (2028, 91)], [
    ("Sales (4000)", 10000.0, None), ("Wages (6000)", -40000.0, None), ("Adjusted profit", lambda a: f"={a['Sales (4000)']}+{a['Wages (6000)']}", "adjustedProfit"),
    ("Main pool from A (closing)", float(A["mainPoolOpening"]) - 20000.0, None),
    ("Capital allowances (18% x 182/365)", lambda a: f"=ROUND({a['Main pool from A (closing)']}*18/100*182/365,2)", "allowances"),
    ("Trading loss", lambda a: f"=MAX({a['Capital allowances (18% x 182/365)']}-{a['Adjusted profit']},0)", "tradingLoss"),
    ("Carried back to A (claim)", lambda a: f"={a['Trading loss']}", "carryBack"),
    ("Losses carried forward", lambda a: f"={a['Trading loss']}-{a['Carried back to A (claim)']}", "lossesCarriedForward"),
    ("Taxable total profits", 0.0, "taxableTotalProfits"),
])
sheet("A after carry back", "A2", 365, True, [(2026, 90), (2027, 275)], common_a + [
    ("Loss carried back from B", lambda a: "='B loss period'!B12", "lossCarriedBackFromFollowing"),
    ("Taxable total profits", lambda a: f"={a['Trading profit']}-{a['Losses brought forward used']}-{a['Loss carried back from B']}", "taxableTotalProfits"),
])
wb.save(OUT)
print(OUT)
