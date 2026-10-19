#!/usr/bin/env python3
"""Writes docs/architecture/v2-m5a-golden-workings.xlsx from tests/golden/v2-m5a-document.json: the statements, the fixed-asset movement table, the
provision, advances and related-party tables, and every rendered sentence of the V2-M5a golden accounts document (DEC-063) recomputed here with
Python's Decimal and calendar dates, independently of the platform, next to the hand-written expected output with a check column. For the named
accountant's review (DEC-054). Exits 1 if anything differs."""
import datetime as dt, json, pathlib, re, sys
from decimal import Decimal as D

root = pathlib.Path(__file__).resolve().parents[2]
G = json.loads((root / "tests/golden/v2-m5a-document.json").read_text())
L = json.loads((root / G["linesFrom"]).read_text())["version"]
day = dt.date.fromisoformat
rows, bad = [], 0
def check(step, item, got, exp):
    global bad
    ok = got == exp; bad += not ok; rows.append([step, item, json.dumps(got) if not isinstance(got, str) else got, json.dumps(exp) if not isinstance(exp, str) else exp, "OK" if ok else "DIFFERS"])
long = lambda d: f"{d.day} {d.strftime('%B')} {d.year}"
num = lambda v: f"{v:,.2f}"

# ── the ledger, written out by hand from the scenario's postings ──
ledger = []                                          # (date, account, debit - credit, source)
def post(d, lines, source="MANUAL"): ledger.extend((day(d), k, D(v), source) for k, v in lines)
for s in G["seed"]:
    amt = sum(D(a) for _, a in s["lines"])
    post(s["date"], ([("1200", amt)] + [(k, -D(a)) for k, a in s["lines"]]) if s["direction"] == "RECEIPT" else ([(k, D(a)) for k, a in s["lines"]] + [("1200", -amt)]), "CASH")
for j in G["journals"]: post(j["date"], [(k, D(dr) - D(cr)) for k, dr, cr in j["lines"]])
fa = G["fixedAssets"]; c = fa["class"]
cost = {a["key"]: D(a["cost"]) for a in fa["assets"]}; acquired = {a["key"]: day(a["acquiredOn"]) for a in fa["assets"]}
per_month = {k: v / D(c["lifeMonths"]) for k, v in cost.items()}
charged = {k: D(0) for k in cost}; disposed = {}
def months(a, upto, disposal=False):                 # monthly convention: full month in the month of acquisition, none in the month of disposal
    s = acquired[a]; span = (upto.year - s.year) * 12 + upto.month - s.month
    return max(0, span) if disposal else span + 1
for st in fa["steps"]:
    d = day(st.get("asAt") or st["on"])
    if st["do"] == "run":
        tot = D(0)
        for a in cost:
            if a in disposed or acquired[a] > d: continue
            due = per_month[a] * min(months(a, d), c["lifeMonths"]) - charged[a]; charged[a] += due; tot += due
        post(st["asAt"], [(c["expenseAccount"], tot), (c["depreciationAccount"], -tot)], "FIXED_ASSET_DEPRECIATION")
        check("Fixed assets", f"run {st['asAt']} total", num(tot), num(D(st["total"])))
    else:
        a = st["asset"]; catch = per_month[a] * months(a, d, True) - charged[a]; charged[a] += catch
        carrying = cost[a] - charged[a]; proceeds = D(st["proceeds"])
        post(st["on"], [(c["expenseAccount"], catch), (c["depreciationAccount"], -catch), (c["depreciationAccount"], charged[a]), (c["disposalAccount"], carrying), (c["costAccount"], -cost[a]),
                        (st["proceedsAccount"], proceeds), (c["disposalAccount"], -proceeds)], "FIXED_ASSET_DISPOSAL")
        disposed[a] = d; check("Fixed assets", f"disposal of {a} catch-up", num(catch), num(D(st["catchUp"])))
p = G["provision"]; post(p["made"]["date"], [(p["expenseAccount"], D(p["made"]["amount"])), (p["account"], -D(p["made"]["amount"]))], "PROVISION_MOVEMENT")

# ── statements through the synthetic version's map ──
MAPPING = {"0030": "BS.FIXED_ASSETS.TANGIBLE", "0031": "BS.FIXED_ASSETS.TANGIBLE", "1110": "BS.CURRENT_ASSETS.DEBTORS", "1200": "BS.CURRENT_ASSETS.CASH",
           "2230": "BS.CREDITORS_WITHIN_ONE_YEAR", "2160": "BS.CREDITORS_WITHIN_ONE_YEAR", "3000": "BS.CAPITAL.SHARE_CAPITAL", "4000": "PL.TURNOVER",
           "6100": "PL.ADMIN_EXPENSES", "6400": "PL.ADMIN_EXPENSES", "6700": "PL.ADMIN_EXPENSES", "6900": "PL.DEPRECIATION", "6950": "PL.OTHER_EXPENSES"}
target = {src: t for src, sub, t in L["mapEntries"] if sub is None}
lines = {l["key"]: l for l in L["lines"]}
pl = lambda k: k[0] in "456789"
def statements(start, end):
    v = {k: D(0) for k, l in lines.items() if l["kind"] == "ACCOUNTS"}
    for (d, k, x, s) in ledger:
        if (pl(k) and start <= d <= end) or (not pl(k) and d <= end):
            t = target[MAPPING[k]]; v[t] += x if lines[t]["naturalSide"] == "DEBIT" else -x
    v["B.PLR"] = -sum((x for (d, k, x, s) in ledger if pl(k) and d <= end), D(0))
    def tot(k):
        if k in v: return v[k]
        v[k] = sum(tot(t["line"]) * t["sign"] for t in lines[k]["formula"]); return v[k]
    for k, l in lines.items():
        if l["kind"] == "TOTAL": tot(k)
    return v
cur, prior = statements(day("2026-01-01"), day("2026-12-31")), statements(day("2025-01-01"), day("2025-12-31"))
for k, (e, ep) in G["expected"]["statements"].items():
    check("Statements", f"{k} 2026", f"{cur[k]:.2f}", e); check("Statements", f"{k} 2025", f"{prior[k]:.2f}", ep)

# ── facts ──
ys, ye, pe = day("2026-01-01"), day("2026-12-31"), day("2025-12-31")
bal = lambda acct, d, sources=None, side=None: sum((x for (dd, k, x, s) in ledger if k == acct and dd <= d and (sources is None or s in sources) and (side is None or (x > 0) == (side == "D"))), D(0))
inyear = lambda acct, sources, side: sum((x for (dd, k, x, s) in ledger if k == acct and ys <= dd <= ye and s in sources and (x > 0) == (side == "D")), D(0))
cbf, ccf = bal("0030", ys - dt.timedelta(days=1)), bal("0030", ye)
disposals = inyear("0030", {"FIXED_ASSET_DISPOSAL"}, "C"); additions = ccf - cbf - disposals
dbf, dcf = -bal("0031", ys - dt.timedelta(days=1)), -bal("0031", ye)
charge = -inyear("0031", {"FIXED_ASSET_DEPRECIATION", "FIXED_ASSET_DISPOSAL"}, "C"); elim = inyear("0031", {"FIXED_ASSET_DISPOSAL"}, "D")
other = dcf - dbf - charge + elim
fa_rows = [[f"Cost at {long(ys)}", cbf], ["Additions", additions], ["Disposals", disposals], [f"Cost at {long(ye)}", ccf], [f"Depreciation at {long(ys)}", dbf], ["Charge for the year", charge],
           ["Eliminated on disposals", -elim], ["Other movements", other], [f"Depreciation at {long(ye)}", dcf], [f"Net book value at {long(ye)}", ccf - dcf], [f"Net book value at {long(pe)}", cbf - dbf]]
check("Fixed assets", "net book value equals B.FA", f"{ccf - dcf:.2f}", f"{cur['B.FA']:.2f}")
tables = {
  "fixedAssets.movements": (["", c["name"], "Total"], [[r, num(v), num(v)] for r, v in fa_rows]),
  "provisions.movements": (["Provision", f"At {long(ys)}", "Made", "Used", "Released", f"At {long(ye)}"], [[p["name"], num(D(0)), num(D(p["made"]["amount"])), num(D(0)), num(D(0)), num(D(p["made"]["amount"]))]]),
}
loan = sorted([(d, -x) for (d, k, x, s) in ledger if k == "2230" and ys <= d <= ye])
opening = -bal("2230", ys - dt.timedelta(days=1)); b = opening; low = opening
for d, m in loan: b += m; low = min(low, b)
adv = -sum((m for d, m in loan if m < 0), D(0)); cred = sum((m for d, m in loan if m > 0), D(0))
tables["directors.advances"] = (["Director", f"At {long(ys)}", "Advanced", "Credited or repaid", f"At {long(ye)}", "Most overdrawn"], [[G["company"]["director"], num(opening), num(adv), num(cred), num(b), num(low)]])
tables["relatedParties.transactions"] = (["Party", "Relationship", "Sales", "Purchases", "Cash paid", "Cash received", "Owed by", "Owed to", "Loan account"],
  [[G["company"]["director"], G["relatedParty"]["relationship"].replace("_", " ").capitalize(), *[num(D(0))] * 6, num(b)]])
facts = {f["key"]: f["value"] for f in G["facts"]}
facts["approval.date"] = long(day(facts["approval.date"])); facts["approval.director"] = G["company"]["director"]

# ── render the wording ──
TOKEN = re.compile(r"\{\{([a-zA-Z]+):([A-Za-z0-9_.\-]+)(?::(prior))?\}\}")
def render(body, fx):
    blocks, missing = [], []
    for para in body.split("\n"):
        m = re.fullmatch(r"\{\{table:([A-Za-z0-9_.]+)\}\}", para)
        if m: cols, rs = tables[m.group(1)]; blocks.append(["table", m.group(1), cols, rs]); continue
        def sub(t):
            kind, key, pr = t.groups()
            if kind == "line": return num((prior if pr else cur)[key])
            if kind == "fact":
                if key not in fx: missing.append(key); return f"[missing: {key}]"
                return fx[key]
            if kind == "year": return long({"start": ys, "end": ye, "priorEnd": pe}[key])
            if kind == "company": return {"name": G["company"]["name"], "number": G["company"]["number"]}[key]
            if kind == "standard": return G["standard"]["name"]
            if kind == "version": return G["version"]["label"]
            if kind == "directors": return G["company"]["director"]
            raise ValueError(t.group(0))
        blocks.append(["text", TOKEN.sub(sub, para)])
    return blocks, missing
def empty(rule):
    keys = [m.group(1) for m in re.finditer(r"\{\{table:([A-Za-z0-9_.]+)\}\}", rule["body"])]; lk = [m.group(2) for m in TOKEN.finditer(rule["body"]) if m.group(1) == "line"]
    return all(len(tables[k][1]) == 0 for k in keys) and all(cur[k] == 0 for k in lk)
ORDER = ["POLICIES", "NOTES", "DIRECTORS_REPORT", "BALANCE_SHEET_STATEMENTS"]
def document(fx):
    sections, omitted, problems = {}, [], []
    for r in sorted(G["rules"], key=lambda r: (ORDER.index(r["section"]), r["sort"], r["key"])):
        if "MICRO" not in r["appliesTo"]: omitted.append([r["key"], "NOT_APPLICABLE"]); continue
        if r.get("omitWhenEmpty") and empty(r): omitted.append([r["key"], "EMPTY"]); continue
        blocks, missing = render(r["body"], fx)
        sections.setdefault(r["section"], []).append({"key": r["key"], "title": r["title"], "blocks": blocks, "missing": missing})
        problems += [f"{r['key']}: missing fact {k}" for k in missing]
    return sections, omitted, problems
E = G["expected"]
sections, omitted, problems = document(facts)
check("Document", "sections in order", list(sections), [s["section"] for s in E["sections"]])
for s in E["sections"]:
    got = {i["key"]: i for i in sections[s["section"]]}
    check("Document", f"{s['section']} items", list(got), [i["key"] for i in s["items"]])
    for i in s["items"]:
        check(f"{s['section']} {i['key']}", "title", got[i["key"]]["title"], i["title"])
        check(f"{s['section']} {i['key']}", "blocks", got[i["key"]]["blocks"], i["blocks"])
        check(f"{s['section']} {i['key']}", "missing facts", got[i["key"]]["missing"], i.get("missing", []))
check("Document", "omitted", omitted, E["omitted"]); check("Document", "problems", problems, E["problemsBefore"]); check("Document", "complete", not problems, E["completeBefore"])
fx2 = {**facts, G["lateFact"]["key"]: G["lateFact"]["value"]}
s2, _, p2 = document(fx2)
gc = next(i for i in s2["NOTES"] if i["key"] == "NOTE.GC"); exp_gc = next(i for i in E["sections"][1]["items"] if i["key"] == "NOTE.GC")["afterLateFact"]
check("Document after the going-concern fact", "NOTE.GC blocks", gc["blocks"], exp_gc); check("Document after the going-concern fact", "complete", not p2, E["completeAfter"])

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:
    print("openpyxl missing"); sys.exit(1)
wb = Workbook(); ws = wb.active; ws.title = "Read me"
for line in ["V2-M5a golden accounts document - workings (SYNTHETIC)", "", G["purpose"], "", G["disclaimer"], "",
             "Every figure, table and sentence is recomputed by this workbook's generator with Python Decimal and calendar dates, independently of the platform, and compared with the hand-written expected output. The platform's API test matches the same output exactly.",
             "", E["fixedAssetWorkings"], E["advancesWorkings"], *[f"{k}: {v}" for k, v in E["workings"].items()]]:
    ws.append([line])
ws.column_dimensions["A"].width = 150; ws["A1"].font = Font(bold=True, size=14)
for r in ws.iter_rows():
    for cc in r: cc.alignment = Alignment(wrap_text=True, vertical="top")
s1 = wb.create_sheet("Ledger"); s1.append(["Date", "Account", "Debit", "Credit", "Source"])
for d, k, x, s in sorted(ledger, key=lambda r: r[0]): s1.append([d.isoformat(), k, float(max(x, D(0))), float(max(-x, D(0))), s])
s2s = wb.create_sheet("Document"); s2s.append(["Section", "Key", "Title", "Block"])
for sec, items in sections.items():
    for i in items:
        for blk in i["blocks"]:
            if blk[0] == "text": s2s.append([sec, i["key"], i["title"], blk[1]])
            else:
                s2s.append([sec, i["key"], i["title"], " | ".join(blk[2])])
                for rr in blk[3]: s2s.append(["", "", "", " | ".join(rr)])
s3 = wb.create_sheet("Checks"); s3.append(["Step", "Item", "Recomputed", "Expected (hand)", "Check"])
for r in rows: s3.append(r)
for sh, widths in ((s1, [12, 10, 14, 14, 28]), (s2s, [26, 16, 40, 120]), (s3, [40, 34, 80, 80, 10])):
    for i, w in enumerate(widths): sh.column_dimensions["ABCDE"[i]].width = w
    for cc in sh[1]: cc.font = Font(bold=True, color="FFFFFF"); cc.fill = PatternFill("solid", fgColor="1F3864")
    for row in sh.iter_rows(min_row=2):
        for cc in row: cc.alignment = Alignment(wrap_text=True, vertical="top")
for row in s3.iter_rows(min_row=2): row[4].fill = PatternFill("solid", fgColor="C6EFCE" if row[4].value == "OK" else "FFC7CE")
out = root / "docs/architecture/v2-m5a-golden-workings.xlsx"; wb.save(out)
print(f"wrote {out.relative_to(root)}; {len(rows)} checks; differences: {bad}")
for r in rows:
    if r[4] == "DIFFERS": print("DIFFERS", r)
sys.exit(1 if bad else 0)
