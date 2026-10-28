#!/usr/bin/env python3
"""Writes docs/architecture/v4-m3-golden-workings.xlsx: the hand workings of the two synthetic CT periods in
tests/golden/v4-m3-capital-allowances.json (V4-M3, DEC-100), for review by the named accountant (the DEC-054 pattern). Every allowance is a formula
over the asset sheet, and the second period opens from the first period's closing cells; the expected columns hold the golden the unit and API
tests assert. The rates and limits are synthetic fixtures, not tax advice. Usage: python3 scripts/docs/gen_capital_allowances_golden.py (needs openpyxl)."""
import datetime as dt
import json
import pathlib
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

root = pathlib.Path(__file__).resolve().parents[2]
G = json.loads((root / "tests/golden/v4-m3-capital-allowances.json").read_text())
OUT = root / "docs/architecture/v4-m3-golden-workings.xlsx"
HEAD = PatternFill("solid", fgColor="1F3A5F")
EXPECT = PatternFill("solid", fgColor="E8F1E4")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
D = dt.date.fromisoformat
MONEY = "#,##0.00"


def header(ws, row, titles, widths):
    for i, (t, w) in enumerate(zip(titles, widths), start=1):
        c = ws.cell(row=row, column=i, value=t)
        c.font, c.fill, c.alignment, c.border = Font(bold=True, color="FFFFFF"), HEAD, Alignment(wrap_text=True, vertical="top"), BOX
        ws.column_dimensions[get_column_letter(i)].width = w


wb = Workbook()
about = wb.active
about.title = "About"
for r, t in enumerate([
    "V4-M3 golden workings: capital allowances of two synthetic CT periods (DEC-100)", "", G["about"],
    "Fraction: 1 for a 12-month period, else days / 365. AIA: the limit x fraction, given to AIA additions in date order; the excess goes into the main pool.",
    "Disposal value: the lower of the proceeds and the cost. Pool: opening + additions - disposal values = available.",
    "Below zero: balancing charge. A single asset pool disposed of: balancing allowance. Main or special pool at or under the small pools limit x fraction: all claimed.",
    "Otherwise WDA = available x rate x fraction, rounded to the penny. Green cells are the golden the tests assert; the Check rows must be TRUE.",
], start=1):
    about.cell(row=r, column=1, value=t).font = Font(bold=(r == 1))
about.column_dimensions["A"].width = 150

ru = wb.create_sheet("Rules")
header(ru, 1, ["Rule", "Value"], [24, 14])
for i, (k, v) in enumerate([("AIA annual limit", G["rules"]["aiaLimit"]), ("Main rate %", G["rules"]["mainRate"]), ("Special rate %", G["rules"]["specialRate"]), ("Small pools limit", G["rules"]["smallPoolLimit"])], start=2):
    ru.cell(row=i, column=1, value=k); ru.cell(row=i, column=2, value=float(v))
AIA, MAINR, SPECR, SMALL = "Rules!$B$2", "Rules!$B$3", "Rules!$B$4", "Rules!$B$5"

at = wb.create_sheet("Assets")
header(at, 1, ["Key", "Description", "Category", "Acquired", "Cost", "Disposed", "Proceeds", "Disposal value"], [6, 30, 20, 12, 14, 12, 12, 14])
row_of = {}
disp = {x["asset"]: x for x in G["disposals"]}
for i, a in enumerate(G["assets"], start=2):
    row_of[a["key"]] = i
    at.cell(row=i, column=1, value=a["key"]); at.cell(row=i, column=2, value=a["description"]); at.cell(row=i, column=3, value=a["category"])
    at.cell(row=i, column=4, value=D(a["acquiredOn"])).number_format = "yyyy-mm-dd"; at.cell(row=i, column=5, value=float(a["cost"])).number_format = MONEY
    if a["key"] in disp:
        at.cell(row=i, column=6, value=D(disp[a["key"]]["disposedOn"])).number_format = "yyyy-mm-dd"
        at.cell(row=i, column=7, value=float(disp[a["key"]]["proceeds"])).number_format = MONEY
        at.cell(row=i, column=8, value=f"=MIN(G{i},E{i})").number_format = MONEY
cost = lambda k: f"Assets!$E${row_of[k]}"
dval = lambda k: f"Assets!$H${row_of[k]}"

closing_cells = {}
for n, y in enumerate(G["years"]):
    ex = G["expected"][n]
    ws = wb.create_sheet(f"Period {n + 1}")
    start, end = D(y["yearStart"]), D(y["yearEnd"])
    inp = lambda d: start <= D(d) <= end
    ws.cell(row=1, column=1, value=f"CT period {y['yearStart']} to {y['yearEnd']}").font = Font(bold=True)
    ws.cell(row=2, column=1, value="Days"); ws.cell(row=2, column=2, value=f"=DATE({end.year},{end.month},{end.day})-DATE({start.year},{start.month},{start.day})+1")
    ws.cell(row=3, column=1, value="Fraction"); ws.cell(row=3, column=2, value="=1" if n == 0 else "=B2/365")
    ws.cell(row=4, column=1, value="AIA available"); ws.cell(row=4, column=2, value=f"=ROUND({AIA}*B3,2)").number_format = MONEY
    ws.cell(row=4, column=3, value=float(ex["aia"]["available"])).fill = EXPECT
    ws.cell(row=5, column=1, value="Small pools limit"); ws.cell(row=5, column=2, value=f"=ROUND({SMALL}*B3,2)").number_format = MONEY
    # AIA claims in date order
    header(ws, 7, ["AIA addition", "Cost", "AIA", "To main pool", "Expected AIA", "Expected to main"], [34, 14, 14, 14, 14, 14])
    aia_assets = sorted([a for a in G["assets"] if a["category"] == "AIA" and inp(a["acquiredOn"])], key=lambda a: a["acquiredOn"])
    r = 8
    first = r
    for k, a in enumerate(aia_assets):
        e = ex["aia"]["claims"][k]
        ws.cell(row=r, column=1, value=f'{a["key"]} {a["description"]}'); ws.cell(row=r, column=2, value=f"={cost(a['key'])}")
        ws.cell(row=r, column=3, value=f"=MIN(B{r},$B$4-SUM(C${first}:C{r - 1}))" if r > first else f"=MIN(B{r},$B$4)")
        ws.cell(row=r, column=4, value=f"=B{r}-C{r}")
        ws.cell(row=r, column=5, value=float(e[1])).fill = EXPECT; ws.cell(row=r, column=6, value=float(e[2])).fill = EXPECT
        r += 1
    aia_last = r - 1
    # pools
    r += 1
    ph = r
    header(ws, ph, ["Pool", "Opening", "Additions", "Disposal values", "Available", "WDA (incl. small pool)", "Balancing allowance", "Balancing charge", "Closing",
                    "Exp. WDA", "Exp. BA", "Exp. BC", "Exp. closing"], [34, 14, 14, 14, 14, 16, 14, 14, 14, 12, 12, 12, 12])
    r += 1
    pool_rows = {}
    for prow in ex["pools"]:
        pool, key = prow[0], prow[1]
        label = pool if not key else f"SINGLE_ASSET {key}"
        ws.cell(row=r, column=1, value=label)
        if n == 0:
            op = G["openings"]["MAIN"] if pool == "MAIN" else G["openings"]["SPECIAL_RATE"] if pool == "SPECIAL_RATE" else next((o["amount"] for o in G["openings"]["single"] if o["asset"] == key), "0")
            ws.cell(row=r, column=2, value=float(op))
        else:
            ws.cell(row=r, column=2, value=f"={closing_cells.get((pool, key), '0')}")
        cat = {"MAIN": ["MAIN_POOL"], "SPECIAL_RATE": ["SPECIAL_RATE_POOL"], "SINGLE_ASSET": ["SINGLE_ASSET_POOL"]}[pool]
        adds = [cost(a["key"]) for a in G["assets"] if a["category"] in cat and inp(a["acquiredOn"]) and (pool != "SINGLE_ASSET" or a["key"] == key)]
        if pool == "MAIN" and aia_assets: adds.append(f"SUM(D{first}:D{aia_last})")
        ws.cell(row=r, column=3, value="=" + ("+".join(adds) if adds else "0"))
        dcat = cat + (["AIA"] if pool == "MAIN" else [])
        ds = [dval(x["asset"]) for x in G["disposals"] if inp(x["disposedOn"]) and next(a for a in G["assets"] if a["key"] == x["asset"])["category"] in dcat and (pool != "SINGLE_ASSET" or x["asset"] == key)]
        ws.cell(row=r, column=4, value="=" + ("+".join(ds) if ds else "0"))
        disposed = pool == "SINGLE_ASSET" and any(x["asset"] == key and inp(x["disposedOn"]) for x in G["disposals"])
        rate = SPECR if pool == "SPECIAL_RATE" else MAINR
        ws.cell(row=r, column=5, value=f"=B{r}+C{r}-D{r}")
        if pool == "SINGLE_ASSET":
            ws.cell(row=r, column=6, value=f"=IF(OR(E{r}<0,{str(disposed).upper()}),0,ROUND(E{r}*{rate}/100*$B$3,2))")
        else:
            ws.cell(row=r, column=6, value=f"=IF(E{r}<0,0,IF(AND(E{r}>0,E{r}<=$B$5),E{r},ROUND(E{r}*{rate}/100*$B$3,2)))")
        ws.cell(row=r, column=7, value=f"=IF(AND({str(disposed).upper()},E{r}>=0),E{r},0)")
        ws.cell(row=r, column=8, value=f"=IF(E{r}<0,-E{r},0)")
        ws.cell(row=r, column=9, value=f"=IF(E{r}<0,0,E{r}-F{r}-G{r})")
        for c, v in zip(range(10, 14), (prow[5], prow[6], prow[7], prow[8])): ws.cell(row=r, column=c, value=float(v)).fill = EXPECT
        for c in range(2, 14): ws.cell(row=r, column=c).number_format = MONEY
        pool_rows[(pool, key)] = r
        closing_cells[(pool, key)] = f"'Period {n + 1}'!I{r}"
        r += 1
    last = r - 1
    r += 1
    t = ex["totals"]
    header(ws, r, ["Totals", "Workings", "Expected"], [34, 14, 14]); r += 1
    tot = [("AIA", f"=SUM(C{first}:C{aia_last})" if aia_assets else "=0", t["aia"]), ("Writing-down allowances", f"=SUM(F{ph + 1}:F{last})", t["writingDownAllowances"]),
           ("Balancing allowances", f"=SUM(G{ph + 1}:G{last})", t["balancingAllowances"]), ("Allowances", f"=B{r}+B{r + 1}+B{r + 2}", t["allowances"]), ("Balancing charges", f"=SUM(H{ph + 1}:H{last})", t["balancingCharges"])]
    t0 = r
    for a, b, c in tot:
        ws.cell(row=r, column=1, value=a); ws.cell(row=r, column=2, value=b).number_format = MONEY; ws.cell(row=r, column=3, value=float(c)).fill = EXPECT; r += 1
    checks = [f"ROUND(B{i}-C{i},2)=0" for i in range(t0, r)] + [f"ROUND(B4-C4,2)=0"] + [f"ROUND(C{i}-E{i},2)=0" for i in range(first, aia_last + 1)] \
        + [f"ROUND({col}{i}-{ecol}{i},2)=0" for i in range(ph + 1, last + 1) for col, ecol in (("F", "J"), ("G", "K"), ("H", "L"), ("I", "M"))]
    ws.cell(row=r + 1, column=1, value="Check: workings equal expected").font = Font(bold=True)
    ws.cell(row=r + 1, column=2, value="=AND(" + ",".join(checks) + ")")
wb.save(OUT)
print(OUT)
