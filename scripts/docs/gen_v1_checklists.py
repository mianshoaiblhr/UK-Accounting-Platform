#!/usr/bin/env python3
"""Writes docs/architecture/v1-reviewer-checklists.xlsx: one working sheet per reviewer named in v1-readiness-report.md section 5,
with the open matrix rows and the legal references. Re-run after editing the lists here. Needs openpyxl."""
import json, pathlib
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

root = pathlib.Path(__file__).resolve().parents[2]
matrix = json.loads((root / "docs/architecture/v1-compliance-matrix.json").read_text())

CHECKLISTS = {
    "Tax reviewer": [
        ("Enter and verify a rate version for each tax code the firm will use (source URL, provision, retrieval date; verifier not the author)", "DEC-021; docs/legal/m4b-vat-sources.xlsx; Settings > VAT codes (API: POST tax-codes/{id}/rates, POST tax-rates/{id}/verify)", "V1-MOD-12, V1-VAT-04", "VAT Notice 700; VAT Act 1994 s.2"),
        ("Confirm the net and VAT box of each rate version", "Rate version boxes; GET vat/workpaper", "V1-VAT-04", "VAT Notice 700/12"),
        ("Confirm per-line half-up rounding (net, or the VAT fraction of the gross)", "DEC-025; tests/unit/vat-math.test.ts", "V1-VAT-04", "VAT Notice 700 (rounding)"),
        ("Confirm the tax-point rule (explicit or issue date; supply date flagged) or specify the derivation to automate", "DEC-028; sales and purchase document screens (taxPointReview)", "V1-VAT-07", "VAT Act 1994 s.6; HMRC VATTOS3600; Notice 700 ss.14-15"),
        ("Confirm credit and debit notes take the original document's rate version and are reported at their own tax point", "DEC-029; tests/platform/sales.test.ts, purchases.test.ts", "V1-VAT-05", "VAT Regs 1995 reg 38; HMRC VATREC13040, VATREC13120"),
        ("Decide which reverse-charge cases the firm sets up as purchase codes", "VAT codes with reverseCharge = true; tests/platform/vat.test.ts", "V1-VAT-10", "VAT Act 1994 s.8, s.55A; SI 2019/892; HMRC VATREVCON13000"),
        ("Give the rule for VAT on foreign-currency supplies (the M5c specification depends on it)", "DEC-031; plan M5c", "V1-VAT-15", "VAT Act 1994 Sch 6 para 11; HMRC VATVAL09500"),
        ("Review the nine-box workpaper of a test company against hand-calculated figures", "GET vat/workpaper; tests/platform/vat.test.ts", "V1-VAT-09", "VAT Notice 700/12"),
    ],
    "Privacy (DEC-013)": [
        ("Read the device-metadata investigation (every write and read of IP address and user agent)", "docs/legal/device-metadata-investigation.md; tests/api/device-metadata.test.ts", "V1-PRIV-01", "UK GDPR art 5"),
        ("Decide the purpose and lawful basis for collecting session and challenge IP address and user agent", "Investigation, proposal section", "V1-PRIV-01", "UK GDPR art 6"),
        ("Decide the retention period and approve the deletion approach (frozen until approved, S4(b))", "docs/legal/retention-verification-schedule.xlsx", "V1-PRIV-01", "UK GDPR art 5(1)(e); ICO storage limitation guidance"),
        ("Note IP addresses can be personal data for a website operator; rate-limit keys hold raw IP addresses", "Investigation, Redis section", "V1-PRIV-01", "CJEU C-582/14 Breyer (19 Oct 2016)"),
    ],
    "Security reviewer": [
        ("Review the protection class of every table and the row-level security and privilege tests", "docs/architecture/security-architecture.md; tests/db/privileges.test.ts", "", ""),
        ("Penetration test API and web tier on a non-production environment (sessions, MFA, Origin checks, idempotency, rate limits)", "apps/api, apps/web; tests/api/*", "", ""),
        ("Review field encryption of bank details and key handling", "contact-roles and bank services; tests/db/contact-roles.test.ts, tests/db/bank.test.ts", "", ""),
        ("Review the document pipeline (AV scan before use) and the CI security gates (CodeQL, audit, secret scan, Trivy)", "CI runs on the release commit", "", ""),
    ],
    "Infrastructure (AWS)": [
        ("Plan and apply infra/terraform in a non-production account; verify encryption, backups, private networking, IAM, logging", "infra/terraform", "", ""),
        ("Set the S3 bucket CORS policy for presigned uploads (statement import)", "infra/terraform storage; M9c acceptance, limitation 4", "", ""),
        ("Run the clean-clone and e2e suites against the deployed environment", "scripts; tests/e2e", "", ""),
        ("Add a Terraform provider lock file (Dependabot's Terraform job fails without one); needs access to registry.terraform.io", "infra/terraform/versions.tf", "", ""),
    ],
    "Performance": [
        ("Load realistic volumes and measure posting, trial balance, statements, alerts and reconciliation suggestions", "packages/accounting (reports, alerts, reconciliation)", "", ""),
        ("Choose indexes, caching or materialised balances from the measurements", "", "", ""),
    ],
    "Product owner": [
        ("Accept V1 as built and verified for its scope, or list what is missing", "docs/architecture/v1-readiness-report.md", "", ""),
        ("Name the qualified tax reviewer and the DPO / legal reviewer", "", "", ""),
        ("Decide MAN-ACC-09 (reporting mappings effective-dated in the database) with V2 planning", "docs/architecture/v2-proposal.md", "MAN-ACC-09", ""),
        ("Approve or amend the V2 proposal before V2 work starts", "docs/architecture/v2-proposal.md", "", ""),
    ],
}
SOURCES = [
    ("VAT Act 1994, s.6 (time of supply)", "https://www.legislation.gov.uk/ukpga/1994/23/section/6", "confirmed by search 2026-10-17"),
    ("VAT Act 1994, s.8 (reverse charge on services from abroad)", "https://www.legislation.gov.uk/ukpga/1994/23/section/8", "standard address; open before relying on it"),
    ("VAT Act 1994, s.55A (domestic reverse charge)", "https://www.legislation.gov.uk/ukpga/1994/23/section/55A", "standard address; open before relying on it"),
    ("HMRC VATREVCON13000 (domestic reverse charge: law)", "https://www.gov.uk/hmrc-internal-manuals/vat-reverse-charge-for-building-and-construction-services-manual/vatrevcon13000", "confirmed by search 2026-10-17"),
    ("VAT Act 1994, Sch 6 para 11 (foreign currency)", "https://www.legislation.gov.uk/ukpga/1994/23/schedule/6/paragraph/11", "standard address; open before relying on it"),
    ("HMRC VATVAL09500 (foreign currency conversion)", "https://www.gov.uk/hmrc-internal-manuals/vat-valuation/vatval09500", "confirmed by search 2026-10-17"),
    ("VAT Regulations 1995 (SI 1995/2518), reg 38", "https://www.legislation.gov.uk/uksi/1995/2518/regulation/38", "standard address; open before relying on it"),
    ("HMRC VATREC13040 (conditions of a valid credit note)", "https://www.gov.uk/hmrc-internal-manuals/vat-trader-records/vatrec13040", "confirmed by search 2026-10-17"),
    ("HMRC VATREC13070 (reg 38: change in consideration)", "https://www.gov.uk/hmrc-internal-manuals/vat-trader-records/vatrec13070", "confirmed by search 2026-10-17"),
    ("HMRC VATREC13120 (credit notes: case law)", "https://www.gov.uk/hmrc-internal-manuals/vat-trader-records/vatrec13120", "confirmed by search 2026-10-17"),
    ("Case: British United Shoe Machinery Co Ltd (1977) VATTR 187", "https://www.gov.uk/hmrc-internal-manuals/vat-trader-records/vatrec13120", "as summarised by HMRC"),
    ("Case: Silvermere Golf and Equestrian Centre Ltd (1981) VATTR 106", "https://www.gov.uk/hmrc-internal-manuals/vat-trader-records/vatrec13120", "as summarised by HMRC"),
    ("Case: Highsize Ltd (LON/90/945)", "https://www.gov.uk/hmrc-internal-manuals/vat-trader-records/vatrec13120", "as summarised by HMRC"),
    ("Case: Kwik Fit (GB) Ltd (1992) VATTR 427", "https://www.gov.uk/hmrc-internal-manuals/vat-trader-records/vatrec13120", "as summarised by HMRC"),
    ("HMRC VATTOS3600 (actual tax points)", "https://www.gov.uk/hmrc-internal-manuals/vat-time-of-supply/vattos3600", "confirmed by search 2026-10-17"),
    ("VAT Notice 700/12 (filling in the VAT return)", "https://www.gov.uk/guidance/how-to-fill-in-and-submit-your-vat-return-vat-notice-70012", "confirmed by search 2026-10-17"),
    ("UK GDPR art 5 (principles)", "https://www.legislation.gov.uk/eur/2016/679/article/5", "standard address; open before relying on it"),
    ("UK GDPR art 6 (lawful basis)", "https://www.legislation.gov.uk/eur/2016/679/article/6", "standard address; open before relying on it"),
    ("ICO guidance: storage limitation", "https://ico.org.uk/for-organisations/uk-gdpr-guidance-and-resources/data-protection-principles/a-guide-to-the-data-protection-principles/storage-limitation/", "confirmed by search 2026-10-17"),
    ("Case: CJEU C-582/14 Breyer v Bundesrepublik Deutschland (19 Oct 2016)", "https://curia.europa.eu/juris/liste.jsf?num=C-582/14", "judgment confirmed by search; standard address"),
]

HEAD = Font(bold=True, color="FFFFFF"); FILL = PatternFill("solid", fgColor="1F3864"); WRAP = Alignment(wrap_text=True, vertical="top")
def sheet(ws, header, rows, widths):
    ws.append(header)
    for c in ws[1]: c.font = HEAD; c.fill = FILL
    for r in rows: ws.append(list(r))
    for i, wdt in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = wdt
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = WRAP
    ws.freeze_panes = "A2"

wb = Workbook()
s = wb.active; s.title = "Summary"
counts = {k: sum(1 for r in matrix["rows"] if r["s"] == k) for k in "IPMN"}
sheet(s, ["Item", "Value"], [
    ("Report", "docs/architecture/v1-readiness-report.md"),
    ("Recommendation", "V1 built and verified for its approved scope; NOT production-approved (DEC-006)"),
    ("Last code commit verified", "c0cad12"), ("CI", "https://github.com/mianshoaiblhr/uk-accounting-platform/actions/runs/38973214119 (success)"),
    ("Clean clone", "1881 tests passed; e2e 21/21; OpenAPI drift 0"),
    ("Matrix", f"{len(matrix['rows'])} rows: Implemented {counts['I']}, Partially {counts['P']}, Missing {counts['M']}"),
    ("Legal references", "Pointers, not legal advice; see the Legal sources sheet for which were confirmed"),
], [28, 110])
o = wb.create_sheet("Open rows")
sheet(o, ["ID", "Requirement", "Status", "Owner", "Acceptance criteria"],
      [(r["id"], r["req"], matrix["statuses"][r["s"]], r.get("owner", ""), r.get("accept", "")) for r in matrix["rows"] if r["s"] in "PM"], [14, 50, 22, 34, 90])
for name, items in CHECKLISTS.items():
    ws = wb.create_sheet(name[:31])
    sheet(ws, ["#", "Check", "Where to look", "Matrix row", "Legal reference", "Result (accept / amend / reject)", "Reviewer", "Date", "Notes"],
          [(i + 1, *it, "", "", "", "") for i, it in enumerate(items)], [5, 60, 46, 14, 34, 20, 18, 12, 30])
l = wb.create_sheet("Legal sources")
sheet(l, ["Reference", "Link", "Check status"], SOURCES, [60, 100, 40])
for row in l.iter_rows(min_row=2, min_col=2, max_col=2):
    for c in row: c.hyperlink = c.value; c.font = Font(color="0563C1", underline="single")
out = root / "docs/architecture/v1-reviewer-checklists.xlsx"
wb.save(out)
print(out, sum(len(v) for v in CHECKLISTS.values()), "checks,", len(SOURCES), "sources")
