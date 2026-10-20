#!/usr/bin/env python3
"""Writes docs/architecture/v2-m5b-golden-workings.xlsx from tests/golden/v2-m5b-size.json: the figures, adjusted limits, conditions, two-year
statuses and results of the V2-M5b golden size test (DEC-066) recomputed here with Python's Decimal, Fraction and calendar dates, independently of
the platform, next to the hand-written expected values with a check column. For the named accountant's review (DEC-054). Exits 1 if anything
differs."""
import datetime as dt, json, pathlib, sys
from decimal import Decimal as D, ROUND_HALF_UP
from fractions import Fraction as F
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

root = pathlib.Path(__file__).resolve().parents[2]
G = json.loads((root / "tests/golden/v2-m5b-size.json").read_text())
day = dt.date.fromisoformat
CLASSES = ["MICRO", "SMALL", "MEDIUM"]
ORDER = {c: i for i, c in enumerate(CLASSES + ["LARGE"])}
rows, bad = [], 0
def check(step, item, got, exp):
    global bad
    ok = got == exp
    bad += not ok; rows.append([step, item, json.dumps(got), json.dumps(exp), "OK" if ok else "DIFFERS"])
c2 = lambda x: str(D(x).quantize(D("0.01"), rounding=ROUND_HALF_UP))

# ── the ledger: cash receipts and payments through 1200; turnover is account 4000 only (4100 lands on other income) ──
ledger = []                                              # (date, code, debit - credit)
for p in G["seed"]:
    for code, amt in p["lines"]:
        a = D(amt); sign = 1 if p["direction"] == "RECEIPT" else -1
        ledger += [(day(p["date"]), "1200", sign * a), (day(p["date"]), code, -sign * a)]
fw = {f["yearStart"]: f for f in G["frameworks"]}
def figures(y):
    s, e = day(y), day(fw[y]["yearEnd"])
    turnover = -sum(x for d, k, x in ledger if k == "4000" and s <= d <= e)        # closing journals are not in this ledger: they are left out of P&L lines
    assets = sum(x for d, k, x in ledger if k == "1200" and d <= e)                # the only debit balance-sheet line with a balance
    return {"turnover": c2(turnover), "balanceSheetTotal": c2(assets), "employees": int(fw[y]["employees"])}

sets = G["sets"]
def governing(y): return next(st for st in sets if day(st["effectiveFrom"]) <= day(y) and (st["effectiveTo"] is None or day(y) <= day(st["effectiveTo"])))
def length(y):
    s, e = day(y), day(fw[y]["yearEnd"])
    d = (e - s).days + 1
    nxt = s.replace(year=s.year + 1)
    return d, (nxt - s).days
def conditions(fig, st, cls, y):
    lim = st["limits"][cls]; d, d12 = length(y)
    tl = F(D(lim["turnover"])) * d / d12
    cond = {"turnover": F(D(fig["turnover"])) <= tl, "balanceSheetTotal": D(fig["balanceSheetTotal"]) <= D(lim["balanceSheetTotal"]), "employees": fig["employees"] <= lim["employees"]}
    shown = {"turnover": c2(D(tl.numerator) / D(tl.denominator)), "balanceSheetTotal": lim["balanceSheetTotal"], "employees": lim["employees"]}
    return shown, cond, sum(cond.values()) >= 2
def prev(y):
    e = day(y) - dt.timedelta(days=1)
    return next((k for k, f in fw.items() if day(f["yearEnd"]) == e), None)

memo = {}
def test(y):
    if y in memo: return memo[y]
    st = governing(y); fig = figures(y); d, d12 = length(y); p = prev(y)
    out = {"start": y, "end": fw[y]["yearEnd"], "days": d, "daysInTwelveMonths": d12, "twelveMonths": d == d12, "thresholdSet": st["label"], "figures": fig, "classes": []}
    if p:
        pst = st if st["priorYearOnFirst"] and governing(p) is not st else governing(p)
        out["priorTestedAgainst"] = pst["label"]; out["prior"] = {"source": "PLATFORM", "start": p, "end": fw[p]["yearEnd"]}
    else:
        f = G["earliestFacts"]; out["priorTestedAgainst"] = None; out["prior"] = {"source": "FACTS", "met": f["size.priorMet"], "status": f["size.priorStatus"]}
    used = []
    for cls in CLASSES:
        shown, cond, met = conditions(fig, st, cls, y)
        if p: prior_met = conditions(figures(p), pst, cls, p)[2]
        else: prior_met = ORDER[cls] >= ORDER[out["prior"]["met"]]
        if met == prior_met: status, basis = met, "BOTH_YEARS"
        elif p: status, basis = test(p)["classes"][CLASSES.index(cls)]["status"], "CARRIED_FORWARD"; used.append(p)
        else: status, basis = ORDER[cls] >= ORDER[out["prior"]["status"]], "CARRIED_FORWARD"
        out["classes"].append({"class": cls, "limits": shown, "conditions": cond, "met": met, "priorMet": prior_met, "status": status, "basis": basis})
    out["_chain"] = [y] + (memo_chain(p) if used else [])
    memo[y] = out
    return out
def memo_chain(p): return test(p)["_chain"]

for y, exp in G["yearResults"].items():
    got = {k: v for k, v in test(y).items() if k != "_chain"}
    for k in exp:
        if k == "classes":
            for gc, ec in zip(got["classes"], exp["classes"]):
                for kk in ec: check(f"Year {y}", f"{ec['class']} {kk}", gc[kk], ec[kk])
        else: check(f"Year {y}", k, got[k], exp[k])
for e in G["expected"]:
    t = test(e["yearStart"])
    q = next((c["class"] for c in t["classes"] if c["status"]), "LARGE")
    check(f"Result {e['yearStart']}", "qualifiesAs", q, e["qualifiesAs"])
    check(f"Result {e['yearStart']}", "chosen", fw[e["yearStart"]]["entitySize"], e["chosen"])
    check(f"Result {e['yearStart']}", "fits", ORDER[fw[e["yearStart"]]["entitySize"]] >= ORDER[q], e["fits"])
    check(f"Result {e['yearStart']}", "years used", t["_chain"], e["years"])
# the claims in the workings: without the adjustment the 2027 micro conditions would be met; without the transitional rule 2026 would be SMALL
fig27 = figures("2027-01-01"); B = sets[1]["limits"]["MICRO"]
check("Workings", "2027 micro met with an unadjusted turnover limit", sum([D(fig27["turnover"]) <= D(B["turnover"]), D(fig27["balanceSheetTotal"]) <= D(B["balanceSheetTotal"]), fig27["employees"] <= B["employees"]]) >= 2, True)
A = sets[0]
met25A = conditions(figures("2025-01-01"), A, "MICRO", "2025-01-01")[2]
check("Workings", "2026 qualifies as SMALL if 2025 were tested against set A", "SMALL" if (not met25A and not test("2025-01-01")["classes"][0]["status"]) else "MICRO", "SMALL")
check("Workings", "2026 turnover leaves out other income", figures("2026-01-01")["turnover"], "28000.00")

wb = Workbook(); ws = wb.active; ws.title = "Read me"
for line in ["V2-M5b golden size test - workings (SYNTHETIC)", "", G["purpose"], "", G["disclaimer"], "",
             "Every figure, limit, condition and status is recomputed by this workbook's generator with Python Decimal, exact fractions and calendar dates, independently of the platform, and compared with the hand-written expected output. The platform's API test matches the same output exactly.",
             "", *[f"{k}: {v}" for k, v in G["workings"].items()]]:
    ws.append([line])
ws.column_dimensions["A"].width = 150; ws["A1"].font = Font(bold=True, size=14)
for r in ws.iter_rows():
    for cc in r: cc.alignment = Alignment(wrap_text=True, vertical="top")
s1 = wb.create_sheet("Ledger"); s1.append(["Date", "Account", "Debit", "Credit"])
for d, k, x in ledger: s1.append([d.isoformat(), k, float(max(x, D(0))), float(max(-x, D(0)))])
s2 = wb.create_sheet("Size test"); s2.append(["Year", "Set", "Prior tested against", "Class", "Turnover / limit", "Balance-sheet total / limit", "Employees / limit", "Met", "Prior met", "Status", "Basis"])
for y in G["yearResults"]:
    t = test(y)
    for c in t["classes"]:
        f, l, cd = t["figures"], c["limits"], c["conditions"]
        mark = lambda k: f"{f[k]} / {l[k]} {'within' if cd[k] else 'EXCEEDS'}"
        s2.append([f"{y} to {t['end']} ({t['days']} days)", t["thresholdSet"], t["priorTestedAgainst"] or "facts", c["class"], mark("turnover"), mark("balanceSheetTotal"), mark("employees"), c["met"], c["priorMet"], c["status"], c["basis"]])
s3 = wb.create_sheet("Checks"); s3.append(["Step", "Item", "Recomputed", "Expected (hand)", "Check"])
for r in rows: s3.append(r)
for sh, widths in ((s1, [12, 10, 14, 14]), (s2, [30, 20, 20, 10, 30, 30, 18, 8, 10, 8, 18]), (s3, [24, 50, 40, 40, 10])):
    for i, w in enumerate(widths): sh.column_dimensions["ABCDEFGHIJK"[i]].width = w
    for cc in sh[1]: cc.font = Font(bold=True, color="FFFFFF"); cc.fill = PatternFill("solid", fgColor="1F3864")
    for row in sh.iter_rows(min_row=2):
        for cc in row: cc.alignment = Alignment(wrap_text=True, vertical="top")
for row in s3.iter_rows(min_row=2): row[4].fill = PatternFill("solid", fgColor="C6EFCE" if row[4].value == "OK" else "FFC7CE")
out = root / "docs/architecture/v2-m5b-golden-workings.xlsx"; wb.save(out)
print(f"wrote {out.relative_to(root)}; {len(rows)} checks; differences: {bad}")
for r in rows:
    if r[4] == "DIFFERS": print("DIFFERS", r)
sys.exit(1 if bad else 0)
