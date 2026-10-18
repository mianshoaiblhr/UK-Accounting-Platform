#!/usr/bin/env python3
"""Writes docs/architecture/v2-m1-golden-workings.xlsx from tests/golden/v2-m1-micro-synthetic.json: the postings, the trial balance of each year,
and every statutory line of both synthetic layouts recomputed here independently of the platform (from the postings, the default chart in
packages/contracts/src/ledger.ts and the map in the golden file), next to the hand-written expected figure, with a check column.
For the named accountant's review of the expected outputs (DEC-054). Usage: python3 scripts/docs/gen_v2_golden_workbook.py (needs openpyxl)."""
import json, pathlib, re, sys
from collections import defaultdict
from decimal import Decimal as D

root = pathlib.Path(__file__).resolve().parents[2]
G = json.loads((root / "tests/golden/v2-m1-micro-synthetic.json").read_text())
chart = {m[0]: {"name": m[1], "type": m[2], "subtype": m[3], "mapping": m[4]}
         for m in re.findall(r"A\('(\d{4})', '((?:[^'\\]|\\')*)', '(\w+)', '(\w+)', '([\w.]+)'", (root / "packages/contracts/src/ledger.ts").read_text())}
YEARS = {"2025": ("2025-01-01", "2025-12-31"), "2026": ("2026-01-01", "2026-12-31")}

# ── the ledger from the postings: (date, account, debit, credit) ──
lines = []
for p in G["postings"]:
    if p["kind"] == "cash":
        for code, amt in p["lines"]:
            a = D(amt)
            lines += [(p["date"], p["bank"], a, D(0)), (p["date"], code, D(0), a)] if p["direction"] == "RECEIPT" else [(p["date"], code, a, D(0)), (p["date"], p["bank"], D(0), a)]
    else:
        lines += [(p["date"], code, D(dr), D(cr)) for code, dr, cr in p["lines"]]
assert sum(l[2] for l in lines) == sum(l[3] for l in lines), "postings do not balance"
PL = lambda code: chart[code]["type"] in ("INCOME", "EXPENSE")
def net(code, start, end):          # debit-positive; P&L accounts over the year, balance-sheet accounts to its end
    return sum((d - c) for dt, a, d, c in lines if a == code and (dt >= start if PL(code) else True) and dt <= end)

def layout(version, start, end):
    m = {(s, sub): t for s, sub, t in version["mapEntries"]}
    defs = {l["key"]: l for l in version["lines"]}
    vals, unplaced = defaultdict(D), []
    for k, l in defs.items():
        if l["kind"] == "ACCOUNTS": vals[k] = D(0)
    for code in sorted({l[1] for l in lines}):
        n = net(code, start, end)
        if n == 0: continue
        acc = chart[code]
        target = m.get((acc["mapping"], acc["subtype"])) or m.get((acc["mapping"], None))
        if not target: unplaced.append((code, n)); continue
        vals[target] += n if defs[target]["naturalSide"] == "DEBIT" else -n
    retained = -sum(net(c, "0001-01-01", end) for c in chart if PL(c) and any(l[1] == c for l in lines))
    for k, l in defs.items():
        if l["kind"] == "RETAINED_PROFIT": vals[k] = retained if l["naturalSide"] == "CREDIT" else -retained
    def total(k):
        if k in vals: return vals[k]
        vals[k] = sum(t["sign"] * total(t["line"]) for t in defs[k]["formula"]); return vals[k]
    for k in defs: total(k)
    return vals, unplaced

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print("openpyxl missing"); sys.exit(1)
wb = Workbook()
HEAD = PatternFill("solid", fgColor="1F3864"); OK = PatternFill("solid", fgColor="C6EFCE"); BAD = PatternFill("solid", fgColor="FFC7CE")
def sheet(ws, header, rows, widths):
    ws.append(header)
    for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = HEAD
    for r in rows: ws.append(r)
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(c.value, float): c.number_format = '#,##0.00;[Red]-#,##0.00'
    ws.freeze_panes = "A2"

ws = wb.active; ws.title = "Read me"
for t in ["V2-M1 golden trial balance - workings (SYNTHETIC)", "", G["purpose"], "", G["disclaimer"], "",
          "Every figure on the layout sheets is recomputed by this workbook's generator from the postings and the default chart, independently of the platform, and compared with the hand-written expected figure in tests/golden/v2-m1-micro-synthetic.json. The platform's API test matches the same expected figures to the cent.",
          "", "Sheets: Postings; Trial balance; Layout 1 (2026 and 2025); Layout 2 (2026 and 2025); Scenarios.", "", f"Recorded correction: {G['expected']['brokenMap']['comment']}"]:
    ws.append([t])
ws.column_dimensions["A"].width = 140
for r in ws.iter_rows():
    for c in r: c.alignment = Alignment(wrap_text=True, vertical="top")
ws["A1"].font = Font(bold=True, size=14)

rows = []
for i, p in enumerate(G["postings"], 1):
    if p["kind"] == "cash":
        for code, amt in p["lines"]:
            dr, cr = (p["bank"], code) if p["direction"] == "RECEIPT" else (code, p["bank"])
            rows += [[i, p["date"], f"Cash {p['direction'].lower()}", p["description"], dr, chart[dr]["name"], float(D(amt)), 0.0], [i, p["date"], "", "", cr, chart[cr]["name"], 0.0, float(D(amt))]]
    else:
        for code, d, c in p["lines"]: rows.append([i, p["date"], "Manual journal", p["description"], code, chart[code]["name"], float(D(d)), float(D(c))])
sheet(wb.create_sheet("Postings"), ["#", "Date", "Kind", "Description", "Account", "Account name", "Debit", "Credit"], rows, [5, 12, 16, 40, 9, 40, 14, 14])

rows = []
for code in sorted({l[1] for l in lines}):
    a = chart[code]
    rows.append([code, a["name"], a["type"], a["subtype"], a["mapping"]] + [float(net(code, *YEARS[y])) for y in ("2025", "2026")])
s = wb.create_sheet("Trial balance")
sheet(s, ["Account", "Name", "Type", "Subtype", "V1 reporting line", "2025 (debit +; P&L for the year, BS at 31.12.2025)", "2026 (debit +; P&L for the year, BS at 31.12.2026)"], rows, [9, 40, 10, 28, 30, 26, 26])
s.append(["", "Check: P&L and BS together net to zero at each year end", "", "", "",
          float(sum(net(c, "0001-01-01", "2025-12-31") for c in {l[1] for l in lines})), float(sum(net(c, "0001-01-01", "2026-12-31") for c in {l[1] for l in lines}))])

def layout_sheet(title, version, expected):
    cur, _ = layout(version, *YEARS["2026"]); prev, _ = layout(version, *YEARS["2025"])
    exp = {**expected["profitAndLoss"], **expected["balanceSheet"]}
    rows, bad = [], 0
    for l in version["lines"]:
        e = exp[l["key"]]; ok = D(e[0]) == cur[l["key"]] and D(e[1]) == prev[l["key"]]; bad += not ok
        rows.append([l["statement"], l["key"], ("    " if l.get("parentKey") else "") + l["name"], l["kind"], l.get("naturalSide") or "", " ".join(f"{'+' if t['sign'] > 0 else '-'}{t['line']}" for t in l.get("formula") or []),
                     float(cur[l["key"]]), float(D(e[0])), float(prev[l["key"]]), float(D(e[1])), "OK" if ok else "DIFFERS"])
    ws = wb.create_sheet(title)
    sheet(ws, ["Statement", "Key", "Line", "Kind", "Natural side", "Total formula", "2026 recomputed", "2026 expected (hand)", "2025 recomputed", "2025 expected (hand)", "Check"], rows, [9, 10, 46, 16, 10, 40, 16, 16, 16, 16, 9])
    for row in ws.iter_rows(min_row=2): row[10].fill = OK if row[10].value == "OK" else BAD
    m = wb.create_sheet(title + " map")
    sheet(m, ["V1 reporting line", "Account subtype (narrows)", "Statutory line"], [[s_, sub or "(any)", t] for s_, sub, t in version["mapEntries"]], [32, 30, 16])
    return bad

bad = layout_sheet("Layout 1", G["version"], G["expected"]) + layout_sheet("Layout 2", G["alternative"]["version"], G["alternative"]["expected"])

# scenarios: override and broken map, recomputed
v = G["version"]; o = G["expected"]["override"]; b = G["expected"]["brokenMap"]
ov = json.loads(json.dumps(v)); acc = chart[o["account"]]
cur, _ = layout(v, *YEARS["2026"]); moved = cur[ "P.OI"] - (-net(o["account"], *YEARS["2026"]))
ovT, ovOI = cur["P.T"] + (-net(o["account"], *YEARS["2026"])), moved
br = json.loads(json.dumps(v)); br["mapEntries"] = [e for e in br["mapEntries"] if [e[0], e[1]] != b["remove"]]
bc, bu = layout(br, *YEARS["2026"]); _, bpu = layout(br, *YEARS["2025"])
rows = [
    ["Override", f"{o['account']} {acc['name']} placed on {o['to']} for this company", "P.T", float(ovT), float(D(o["P.T"]))], ["", "", "P.OI", float(ovOI), float(D(o["P.OI"]))], ["", "", "P.PROFIT", float(cur["P.PROFIT"]), float(D(o["P.PROFIT"]))],
    ["Broken map", f"map entry {b['remove'][0]} removed in a superseding version", "B.CA1", float(bc["B.CA1"]), float(D(b["B.CA1"]))], ["", "", "B.NA", float(bc["B.NA"]), float(D(b["B.NA"]))], ["", "", "B.CR", float(bc["B.CR"]), float(D(b["B.CR"]))],
    ["", "unplaced 2026 (debit +)", ", ".join(c for c, _ in bu), float(sum(n for _, n in bu)), float(sum(D(n) for _, n in b["unplaced"]))],
    ["", "unplaced 2025 (debit +)", ", ".join(c for c, _ in bpu), float(sum(n for _, n in bpu)), float(sum(D(n) for _, n in b["comparativeUnplaced"]))],
    ["", "net assets less capital and reserves (validation NA_EQ_CR fails)", "", float(bc["B.NA"] - bc["B.CR"]), float(D(b["balanceSheetDifference"]))],
    ["Earlier year", "2025 on the version it was prepared under, after a later version replaced it", "B.NA", float(layout(v, *YEARS["2025"])[0]["B.NA"]), float(D(G["expected"]["earlierYear"]["B.NA"]))],
]
for r in rows: r.append("OK" if D(str(r[3])) == D(str(r[4])) else "DIFFERS"); bad += r[-1] != "OK"
ws = wb.create_sheet("Scenarios")
sheet(ws, ["Scenario", "What changes", "Line", "Recomputed", "Expected (hand)", "Check"], rows, [14, 60, 14, 16, 16, 9])
for row in ws.iter_rows(min_row=2): row[5].fill = OK if row[5].value == "OK" else BAD
out = root / "docs/architecture/v2-m1-golden-workings.xlsx"
wb.save(out)
print(f"wrote {out.relative_to(root)}; differences: {bad}")
sys.exit(1 if bad else 0)
