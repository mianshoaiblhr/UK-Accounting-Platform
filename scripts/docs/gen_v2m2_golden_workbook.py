#!/usr/bin/env python3
"""Writes docs/architecture/v2-m2-golden-workings.xlsx from tests/golden/v2-m2-fixed-assets.json: every depreciation charge, the disposal, the
carrying amounts, the reconciliation and the ledger totals of the V2-M2 golden scenario recomputed here with Python's Decimal and datetime,
independently of the platform, next to the hand-written expected figure with a check column. For the named accountant's review (DEC-054).
Usage: python3 scripts/docs/gen_v2m2_golden_workbook.py (needs openpyxl). Exits 1 if any figure differs."""
import calendar, datetime as dt, json, pathlib, sys
from decimal import Decimal as D, ROUND_HALF_UP

root = pathlib.Path(__file__).resolve().parents[2]
G = json.loads((root / "tests/golden/v2-m2-fixed-assets.json").read_text())
day = dt.date.fromisoformat
C2 = lambda x: x.quantize(D("0.01"), rounding=ROUND_HALF_UP)
def add_months(d, n):
    m = d.month - 1 + n; y = d.year + m // 12; m = m % 12 + 1
    return dt.date(y, m, min(d.day, calendar.monthrange(y, m)[1]))
def months_counted(start, d, disposal=False):
    if d < start: return 0
    span = (d.year - start.year) * 12 + d.month - start.month
    if disposal: return max(0, span)
    return span + 1 if d.day == calendar.monthrange(d.year, d.month)[1] else span
def due(b, d, disposal=False):
    """Depreciation due since the basis start at d, DEC-058 / v2-plan 4.2."""
    if d < b["from"]: return D(0)
    cap = max(D(0), b["carrying"] - b["residual"])
    if b["method"] == "STRAIGHT_LINE":
        if b["convention"] == "DAILY":
            life = (add_months(b["from"], b["life"]) - b["from"]).days
            return cap * D(min((d - b["from"]).days + 1, life)) / D(life)
        return cap * D(min(months_counted(b["from"], d, disposal), b["life"])) / D(b["life"])
    r = b["rate"] / 100
    if b["convention"] == "DAILY":
        nxt = d + dt.timedelta(days=1); n = 0
        while add_months(b["from"], 12 * (n + 1)) <= nxt: n += 1
        ys, ye = add_months(b["from"], 12 * n), add_months(b["from"], 12 * (n + 1))
        f = D((nxt - ys).days) / D((ye - ys).days)
    else:
        m = months_counted(b["from"], d, disposal); n, f = m // 12, D(m % 12) / 12
    carrying = b["carrying"] * (1 - r) ** n * (1 - r * f)
    return min(cap, b["carrying"] - max(carrying, b["residual"]))

classes = {c["code"]: c for c in G["classes"]}
state = {}
for a in G["assets"]:
    k = classes[a["class"]]; add = a["addition"]
    method = a.get("method", k["method"]); take = D(add.get("accumulatedDepreciation", "0")); cost = D(a["cost"])
    frm = day(add["takeOnDate"]) + dt.timedelta(days=1) if add["kind"] == "TAKE_ON" else day(a.get("inUseFrom", a["acquiredOn"]))
    state[a["key"]] = {"asset": a, "class": k, "cost": cost, "charged": take, "status": "ACTIVE", "inUse": day(a.get("inUseFrom", a["acquiredOn"])), "last": None,
      "basis": {"from": frm, "carrying": cost - take, "before": take, "method": method, "life": a.get("lifeMonths", k.get("lifeMonths")), "rate": D(a.get("ratePercent", k.get("ratePercent", "0"))),
                "residual": D(a.get("residual", "0")), "convention": a.get("convention", k["convention"])}}
acc_due = lambda s, d, disposal=False: C2(s["basis"]["before"] + due(s["basis"], d, disposal))
ledger = []          # (date, account, debit - credit)
rows, bad = [], 0
def check(step, item, got, exp):
    global bad
    ok = D(str(got)) == D(str(exp)); bad += not ok
    rows.append([step, item, float(D(str(got))), float(D(str(exp))), "OK" if ok else "DIFFERS"])
for st in G["steps"]:
    if st["do"] == "run":
        d = day(st["asAt"]); per_class = {}
        for key, s in state.items():
            if s["status"] != "ACTIVE" or s["inUse"] > d: continue
            ch = acc_due(s, d) - s["charged"]
            if ch > 0:
                check(f"Run {st['asAt']}", f"{key} charge", ch, st["expect"]["lines"].get(key, "0"))
                s["charged"] += ch; s["last"] = d; per_class.setdefault(s["class"]["code"], D(0)); per_class[s["class"]["code"]] += ch
        for code, amt in per_class.items():
            ledger += [(d, classes[code]["expenseAccount"], amt), (d, classes[code]["depreciationAccount"], -amt)]
        check(f"Run {st['asAt']}", "total", sum(per_class.values()), st["expect"]["total"])
    elif st["do"] == "dispose":
        s = state[st["asset"]]; d = day(st["on"]); k = s["class"]
        catch = max(D(0), acc_due(s, d, True) - s["charged"]); acc = s["charged"] + catch; carrying = s["cost"] - acc; proceeds = D(st["proceeds"])
        check(f"Dispose {st['asset']}", "catch-up", catch, st["expect"]["catchUp"]); check(f"Dispose {st['asset']}", "carrying amount", carrying, st["expect"]["carrying"])
        check(f"Dispose {st['asset']}", "gain or loss", proceeds - carrying, st["expect"]["gainOrLoss"])
        ledger += [(d, k["expenseAccount"], catch), (d, k["depreciationAccount"], -catch), (d, k["depreciationAccount"], acc), (d, k["disposalAccount"], carrying), (d, k["costAccount"], -s["cost"]),
                   (d, st["proceedsAccount"], proceeds), (d, k["disposalAccount"], -proceeds)]
        s.update(status="DISPOSED", charged=acc, carrying=carrying)
    elif st["do"] == "register":
        d = day(st["asAt"])
        for key, (cost, accd, carrying, status) in st["expect"]["assets"].items():
            s = state[key]; check(f"Register {st['asAt']}", f"{key} accumulated", s["charged"], accd); check(f"Register {st['asAt']}", f"{key} carrying", s["cost"] - s["charged"], carrying)
        live = [s for s in state.values() if s["status"] == "ACTIVE"]
        for code, (reg, led) in st["expect"]["reconciliation"].items():
            kind = "cost" if any(c["costAccount"] == code for c in classes.values()) else "dep"
            r = sum((s["cost"] if kind == "cost" else s["charged"]) for s in live if s["class"]["costAccount" if kind == "cost" else "depreciationAccount"] == code)
            # the ledger: what the register posted plus the opening postings it links (the purchases and the take-on journal)
            opening = D(0)
            for p in G["seed"]:
                for ln in p["lines"]:
                    if ln[0] == code: opening += D(ln[1]) if p["kind"] == "cash" else D(ln[1]) - D(ln[2])
            for a in G["assets"]:
                if a["addition"]["kind"] == "POSTED" and kind == "cost" and classes[a["class"]]["costAccount"] == code: opening += D(a["cost"])
            l = opening + sum(v for (dd, ac, v) in ledger if ac == code and dd <= d)
            check(f"Register {st['asAt']}", f"{code} register", r, reg); check(f"Register {st['asAt']}", f"{code} ledger", l if kind == "cost" else -l, led)
    elif st["do"] == "ledger":
        for code, years in st["expect"].items():
            for y, v in years.items():
                check("Ledger", f"{code} {y}", sum(val for (dd, ac, val) in ledger if ac == code and dd.year == int(y)), v)
    elif st["do"] == "estimate":
        s = state[st["asset"]]; frm = s["last"] + dt.timedelta(days=1)
        check(f"Estimate {st['asset']}", "carrying at new basis", s["cost"] - s["charged"], st["expect"]["carryingAtBasis"])
        check(f"Estimate {st['asset']}", "effective from (days after 1 Jan 2027)", (frm - dt.date(2027, 1, 1)).days, (day(st["expect"]["effectiveFrom"]) - dt.date(2027, 1, 1)).days)
        s["basis"] = {**s["basis"], "from": frm, "carrying": s["cost"] - s["charged"], "before": s["charged"], "life": st["lifeMonths"], "residual": D(st["residual"])}

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:
    print("openpyxl missing"); sys.exit(1)
wb = Workbook(); ws = wb.active; ws.title = "Read me"
for t in ["V2-M2 golden fixed assets - workings (SYNTHETIC)", "", G["purpose"], "", G["disclaimer"], "",
          "Every figure is recomputed by this workbook's generator with Python Decimal and calendar dates, independently of the platform, and compared with the hand-written expected figure. The platform's API test matches the same figures to the cent.",
          "", "Sheets: Steps (the scenario and why each figure is what it is); Checks (recomputed against expected)."]: ws.append([t])
ws.column_dimensions["A"].width = 140; ws["A1"].font = Font(bold=True, size=14)
for r in ws.iter_rows():
    for c in r: c.alignment = Alignment(wrap_text=True, vertical="top")
HEAD = PatternFill("solid", fgColor="1F3864")
s2 = wb.create_sheet("Steps"); s2.append(["#", "Step", "Date / asset", "Why (hand calculation)"])
for i, st in enumerate(G["steps"], 1): s2.append([i, st["do"], st.get("asAt") or st.get("on") or st.get("asset") or "", st.get("why", "")])
for i, w in enumerate([5, 12, 14, 150], 1): s2.column_dimensions["ABCD"[i - 1]].width = w
s3 = wb.create_sheet("Checks"); s3.append(["Step", "Figure", "Recomputed", "Expected (hand)", "Check"])
for r in rows: s3.append(r)
for i, w in enumerate([22, 34, 16, 16, 10], 1): s3.column_dimensions["ABCDE"[i - 1]].width = w
for sh in (s2, s3):
    for c in sh[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = HEAD
    for row in sh.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(c.value, float): c.number_format = '#,##0.00;[Red]-#,##0.00'
for row in s3.iter_rows(min_row=2): row[4].fill = PatternFill("solid", fgColor="C6EFCE" if row[4].value == "OK" else "FFC7CE")
out = root / "docs/architecture/v2-m2-golden-workings.xlsx"; wb.save(out)
print(f"wrote {out.relative_to(root)}; {len(rows)} figures; differences: {bad}")
sys.exit(1 if bad else 0)
