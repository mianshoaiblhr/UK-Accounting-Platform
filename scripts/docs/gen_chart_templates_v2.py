#!/usr/bin/env python3
"""Writes docs/architecture/chart-templates-v2.xlsx: the draft Sage-style chart-of-accounts templates of DEC-080, one sheet per business
type (limited company, sole trader, partnership, landlord), for the product owner's review before any code. Every HMRC category and
SA103F/SA105 box reference is UNVERIFIED (DEC-022). Usage: python3 scripts/docs/gen_chart_templates_v2.py (needs openpyxl)."""
import pathlib
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

root = pathlib.Path(__file__).resolve().parents[2]
OUT = root / "docs/architecture/chart-templates-v2.xlsx"

# Self-employment (SE) and UK property (UP) categories, as named in HMRC's MTD APIs, with the SA103F / SA105 box they feed. UNVERIFIED.
SE = {
    "turnover": "15", "other": "16", "costOfGoods": "17", "paymentsToSubcontractors": "18", "wagesAndStaffCosts": "19",
    "carVanTravelExpenses": "20", "premisesRunningCosts": "21", "maintenanceCosts": "22", "adminCosts": "23", "advertisingCosts": "24",
    "businessEntertainmentCosts": "24 (disallowable)", "interestOnBankOtherLoans": "25", "financeCharges": "26", "irrecoverableDebts": "27",
    "professionalFees": "28", "depreciation": "29 (disallowable)", "otherExpenses": "30",
}
UP = {
    "periodAmount": "20", "premiumsOfLeaseGrant": "22", "reversePremiums": "23", "otherIncome": "20", "premisesRunningCosts": "24",
    "repairsAndMaintenance": "25", "financialCosts": "26", "professionalFees": "27", "costOfServices": "28", "other": "29",
    "travelCosts": "29", "residentialFinancialCost": "44",
}
NOT_TRADE = "not business income"

# (templates, code, name, type, subtype, control, VAT treatment, reporting line, HMRC category, note)
# templates: L limited company, S sole trader, P partnership, R landlord (UK property). Subtypes and lines marked NEW do not exist yet.
ROWS = [
    ("LSP", "0010", "Freehold property", "ASSET", "FIXED_ASSET_TANGIBLE", "", "NOT_APPLICABLE", "BS.FIXED_ASSETS.TANGIBLE", "", ""),
    ("R", "0010", "Let property at cost", "ASSET", "FIXED_ASSET_TANGIBLE", "", "NOT_APPLICABLE", "BS.FIXED_ASSETS.TANGIBLE", "", "Only if the landlord keeps a balance sheet; SA105 does not need one"),
    ("LSP", "0011", "Leasehold property", "ASSET", "FIXED_ASSET_TANGIBLE", "", "NOT_APPLICABLE", "BS.FIXED_ASSETS.TANGIBLE", "", ""),
    ("LSP", "0020", "Plant and machinery", "ASSET", "FIXED_ASSET_TANGIBLE", "", "NOT_APPLICABLE", "BS.FIXED_ASSETS.TANGIBLE", "", "Capital allowances, not depreciation, for tax"),
    ("LSP", "0021", "Plant and machinery depreciation", "ASSET", "FIXED_ASSET_DEPRECIATION", "", "NOT_APPLICABLE", "BS.FIXED_ASSETS.TANGIBLE", "", ""),
    ("LSP", "0030", "Office equipment", "ASSET", "FIXED_ASSET_TANGIBLE", "", "NOT_APPLICABLE", "BS.FIXED_ASSETS.TANGIBLE", "", ""),
    ("LSP", "0031", "Office equipment depreciation", "ASSET", "FIXED_ASSET_DEPRECIATION", "", "NOT_APPLICABLE", "BS.FIXED_ASSETS.TANGIBLE", "", ""),
    ("LSP", "0040", "Furniture and fixtures", "ASSET", "FIXED_ASSET_TANGIBLE", "", "NOT_APPLICABLE", "BS.FIXED_ASSETS.TANGIBLE", "", ""),
    ("LSP", "0041", "Furniture and fixtures depreciation", "ASSET", "FIXED_ASSET_DEPRECIATION", "", "NOT_APPLICABLE", "BS.FIXED_ASSETS.TANGIBLE", "", ""),
    ("LSP", "0050", "Motor vehicles", "ASSET", "FIXED_ASSET_TANGIBLE", "", "NOT_APPLICABLE", "BS.FIXED_ASSETS.TANGIBLE", "", ""),
    ("LSP", "0051", "Motor vehicles depreciation", "ASSET", "FIXED_ASSET_DEPRECIATION", "", "NOT_APPLICABLE", "BS.FIXED_ASSETS.TANGIBLE", "", ""),
    ("LSP", "0060", "Computer equipment", "ASSET", "FIXED_ASSET_TANGIBLE", "", "NOT_APPLICABLE", "BS.FIXED_ASSETS.TANGIBLE", "", ""),
    ("LSP", "0061", "Computer equipment depreciation", "ASSET", "FIXED_ASSET_DEPRECIATION", "", "NOT_APPLICABLE", "BS.FIXED_ASSETS.TANGIBLE", "", ""),
    ("LSP", "0070", "Goodwill", "ASSET", "FIXED_ASSET_INTANGIBLE", "", "NOT_APPLICABLE", "BS.FIXED_ASSETS.INTANGIBLE", "", ""),
    ("LSP", "0071", "Goodwill amortisation", "ASSET", "FIXED_ASSET_DEPRECIATION", "", "NOT_APPLICABLE", "BS.FIXED_ASSETS.INTANGIBLE", "", ""),
    ("LSP", "1001", "Stock", "ASSET", "CURRENT_ASSET_STOCK", "", "NOT_APPLICABLE", "BS.CURRENT_ASSETS.STOCK", "", ""),
    ("LSP", "1002", "Work in progress", "ASSET", "CURRENT_ASSET_STOCK", "", "NOT_APPLICABLE", "BS.CURRENT_ASSETS.STOCK", "", ""),
    ("LSPR", "1100", "Trade debtors", "ASSET", "CURRENT_ASSET_RECEIVABLE", "TRADE_RECEIVABLES", "NOT_APPLICABLE", "BS.CURRENT_ASSETS.DEBTORS", "", "System account. For a landlord: rent due from tenants"),
    ("LSPR", "1102", "Other debtors", "ASSET", "CURRENT_ASSET_RECEIVABLE", "", "NOT_APPLICABLE", "BS.CURRENT_ASSETS.DEBTORS", "", ""),
    ("LSPR", "1103", "Prepayments", "ASSET", "CURRENT_ASSET_PREPAYMENT", "", "NOT_APPLICABLE", "BS.CURRENT_ASSETS.DEBTORS", "", ""),
    ("LSP", "1104", "Accrued income", "ASSET", "CURRENT_ASSET_RECEIVABLE", "", "NOT_APPLICABLE", "BS.CURRENT_ASSETS.DEBTORS", "", ""),
    ("LSPR", "1200", "Bank current account", "ASSET", "CURRENT_ASSET_BANK", "BANK", "NOT_APPLICABLE", "BS.CURRENT_ASSETS.CASH", "", ""),
    ("LSPR", "1210", "Bank deposit account", "ASSET", "CURRENT_ASSET_BANK", "BANK", "NOT_APPLICABLE", "BS.CURRENT_ASSETS.CASH", "", ""),
    ("LSP", "1220", "Building society account", "ASSET", "CURRENT_ASSET_BANK", "BANK", "NOT_APPLICABLE", "BS.CURRENT_ASSETS.CASH", "", ""),
    ("LSP", "1230", "Petty cash", "ASSET", "CURRENT_ASSET_CASH", "CASH", "NOT_APPLICABLE", "BS.CURRENT_ASSETS.CASH", "", "Cannot go below zero (DEC-036)"),
    ("LSPR", "1240", "Credit card", "ASSET", "CURRENT_ASSET_BANK", "BANK", "NOT_APPLICABLE", "BS.CURRENT_ASSETS.CASH", "", "Runs overdrawn, with the DEC-036 warning"),
    ("LSP", "1250", "Card and online payment clearing", "ASSET", "CURRENT_ASSET_OTHER", "", "NOT_APPLICABLE", "BS.CURRENT_ASSETS.DEBTORS", "", "Takings held by a payment provider before payout"),
    ("LSPR", "2100", "Trade creditors", "LIABILITY", "CURRENT_LIABILITY_PAYABLE", "TRADE_PAYABLES", "NOT_APPLICABLE", "BS.CREDITORS_WITHIN_ONE_YEAR", "", "System account"),
    ("LSPR", "2109", "Accruals", "LIABILITY", "CURRENT_LIABILITY_ACCRUAL", "", "NOT_APPLICABLE", "BS.CREDITORS_WITHIN_ONE_YEAR", "", ""),
    ("LSPR", "2110", "Deferred income", "LIABILITY", "CURRENT_LIABILITY_OTHER", "", "NOT_APPLICABLE", "BS.CREDITORS_WITHIN_ONE_YEAR", "", "For a landlord: rent received in advance"),
    ("R", "2120", "Tenant deposits held", "LIABILITY", "CURRENT_LIABILITY_OTHER", "", "NOT_APPLICABLE", "BS.CREDITORS_WITHIN_ONE_YEAR", "", "Only deposits the landlord holds outside a protection scheme's custody"),
    ("LSPR", "2200", "VAT control account", "LIABILITY", "CURRENT_LIABILITY_VAT", "VAT_CONTROL", "VAT_CONTROL", "BS.CREDITORS_WITHIN_ONE_YEAR", "", "System account. One control account, not Sage's 2200/2201/2202: the VAT engine posts to one (M4b)"),
    ("LSP", "2210", "PAYE and National Insurance", "LIABILITY", "CURRENT_LIABILITY_PAYE", "", "NOT_APPLICABLE", "BS.CREDITORS_WITHIN_ONE_YEAR", "", ""),
    ("LSP", "2220", "Net wages", "LIABILITY", "CURRENT_LIABILITY_OTHER", "", "NOT_APPLICABLE", "BS.CREDITORS_WITHIN_ONE_YEAR", "", ""),
    ("LSP", "2230", "Pension contributions payable", "LIABILITY", "CURRENT_LIABILITY_OTHER", "", "NOT_APPLICABLE", "BS.CREDITORS_WITHIN_ONE_YEAR", "", ""),
    ("LSP", "2240", "CIS deductions payable", "LIABILITY", "CURRENT_LIABILITY_PAYE", "", "NOT_APPLICABLE", "BS.CREDITORS_WITHIN_ONE_YEAR", "", "Tax deducted from subcontractors, paid to HMRC"),
    ("L", "2250", "Directors' loan accounts", "LIABILITY", "CURRENT_LIABILITY_OTHER", "", "NOT_APPLICABLE", "BS.CREDITORS_WITHIN_ONE_YEAR", "", "Read by the director-loan register (DEC-062)"),
    ("L", "2320", "Corporation tax", "LIABILITY", "CURRENT_LIABILITY_TAX", "", "NOT_APPLICABLE", "BS.CREDITORS_WITHIN_ONE_YEAR", "", ""),
    ("LSPR", "2300", "Bank loans", "LIABILITY", "LONG_TERM_LIABILITY", "", "NOT_APPLICABLE", "BS.CREDITORS_AFTER_ONE_YEAR", "", "The part due within a year is reclassified at the year end"),
    ("LSP", "2310", "Hire purchase", "LIABILITY", "LONG_TERM_LIABILITY", "", "NOT_APPLICABLE", "BS.CREDITORS_AFTER_ONE_YEAR", "", ""),
    ("LSPR", "2330", "Mortgages", "LIABILITY", "LONG_TERM_LIABILITY", "", "NOT_APPLICABLE", "BS.CREDITORS_AFTER_ONE_YEAR", "", ""),
    ("L", "3000", "Ordinary shares", "EQUITY", "SHARE_CAPITAL", "", "NOT_APPLICABLE", "BS.CAPITAL.SHARE_CAPITAL", "", ""),
    ("L", "3010", "Preference shares", "EQUITY", "SHARE_CAPITAL", "", "NOT_APPLICABLE", "BS.CAPITAL.SHARE_CAPITAL", "", ""),
    ("L", "3100", "Other reserves", "EQUITY", "RESERVES", "", "NOT_APPLICABLE", "BS.CAPITAL.RESERVES", "", ""),
    ("SR", "3000", "Capital introduced", "EQUITY", "OWNER_CAPITAL (NEW)", "", "NOT_APPLICABLE", "BS.CAPITAL.PROPRIETOR (NEW)", "", ""),
    ("P", "3000", "Partner 1 capital account", "EQUITY", "PARTNER_CAPITAL (NEW)", "", "NOT_APPLICABLE", "BS.CAPITAL.PARTNERS (NEW)", "", "One per partner: 3000 to 3099"),
    ("P", "3100", "Partner 1 current account", "EQUITY", "PARTNER_CURRENT (NEW)", "", "NOT_APPLICABLE", "BS.CAPITAL.PARTNERS (NEW)", "", "One per partner: 3100 to 3199; profit shares credited here"),
    ("LSPR", "3200", "Profit and loss account", "EQUITY", "RETAINED_EARNINGS", "RETAINED_EARNINGS", "NOT_APPLICABLE", "BS.CAPITAL.RESERVES", "", "System account; receives the year-end closing journal (DEC-060)"),
    ("L", "3300", "Dividends", "EQUITY", "DIVIDENDS", "", "NOT_APPLICABLE", "BS.CAPITAL.RESERVES", "", ""),
    ("SR", "3260", "Drawings", "EQUITY", "DRAWINGS (NEW)", "", "NOT_APPLICABLE", "BS.CAPITAL.PROPRIETOR (NEW)", "", "Owner's wages are drawings, not wages"),
    ("P", "3260", "Partner 1 drawings", "EQUITY", "DRAWINGS (NEW)", "", "NOT_APPLICABLE", "BS.CAPITAL.PARTNERS (NEW)", "", "One per partner: 3260 to 3299"),
    ("LSP", "4000", "Sales - standard rate", "INCOME", "SALES", "", "VATABLE", "PL.TURNOVER", "SE: turnover", ""),
    ("LSP", "4001", "Sales - reduced rate", "INCOME", "SALES", "", "VATABLE", "PL.TURNOVER", "SE: turnover", ""),
    ("LSP", "4002", "Sales - zero rate", "INCOME", "SALES", "", "VATABLE", "PL.TURNOVER", "SE: turnover", ""),
    ("LSP", "4003", "Sales - exempt", "INCOME", "SALES", "", "EXEMPT", "PL.TURNOVER", "SE: turnover", ""),
    ("LSP", "4004", "Sales - outside the scope of VAT", "INCOME", "SALES", "", "OUT_OF_SCOPE", "PL.TURNOVER", "SE: turnover", ""),
    ("LSP", "4009", "Discounts allowed", "INCOME", "SALES", "", "VATABLE", "PL.TURNOVER", "SE: turnover", "Debit balance: reduces turnover"),
    ("LSP", "4900", "Miscellaneous income", "INCOME", "OTHER_INCOME", "", "OUT_OF_SCOPE", "PL.OTHER_INCOME", "SE: other", ""),
    ("LSP", "4902", "Commissions received", "INCOME", "OTHER_INCOME", "", "VATABLE", "PL.OTHER_INCOME", "SE: other", ""),
    ("LSP", "4903", "Insurance claims received", "INCOME", "OTHER_INCOME", "", "OUT_OF_SCOPE", "PL.OTHER_INCOME", "SE: other", ""),
    ("LSP", "4904", "Rent received", "INCOME", "OTHER_INCOME", "", "EXEMPT", "PL.OTHER_INCOME", f"SE: {NOT_TRADE}", "For an individual this is property income: report it on SA105, not SA103"),
    ("LSP", "4910", "Grants received", "INCOME", "OTHER_INCOME", "", "OUT_OF_SCOPE", "PL.OTHER_INCOME", "SE: other", ""),
    ("LSP", "4920", "Bank interest received", "INCOME", "INTEREST_INCOME", "", "OUT_OF_SCOPE", "PL.INTEREST_RECEIVABLE", f"SE: {NOT_TRADE}", "Savings income on the SA100 for an individual"),
    ("R", "4000", "Rents received", "INCOME", "SALES", "", "EXEMPT", "PL.TURNOVER", "UP: periodAmount", "Exempt unless the landlord opted to tax a commercial property"),
    ("R", "4010", "Lease premiums", "INCOME", "SALES", "", "EXEMPT", "PL.TURNOVER", "UP: premiumsOfLeaseGrant", ""),
    ("R", "4020", "Reverse premiums and inducements", "INCOME", "OTHER_INCOME", "", "OUT_OF_SCOPE", "PL.OTHER_INCOME", "UP: reversePremiums", ""),
    ("R", "4030", "Other property income", "INCOME", "OTHER_INCOME", "", "EXEMPT", "PL.OTHER_INCOME", "UP: otherIncome", ""),
    ("LSP", "5000", "Materials and goods purchased", "EXPENSE", "COST_OF_SALES", "", "VATABLE", "PL.COST_OF_SALES", "SE: costOfGoods", ""),
    ("LSP", "5001", "Goods imported", "EXPENSE", "COST_OF_SALES", "", "VATABLE", "PL.COST_OF_SALES", "SE: costOfGoods", ""),
    ("LSP", "5009", "Discounts taken", "EXPENSE", "COST_OF_SALES", "", "VATABLE", "PL.COST_OF_SALES", "SE: costOfGoods", "Credit balance: reduces cost of sales"),
    ("LSP", "5100", "Carriage inwards", "EXPENSE", "COST_OF_SALES", "", "VATABLE", "PL.COST_OF_SALES", "SE: costOfGoods", ""),
    ("LSP", "5200", "Stock movement", "EXPENSE", "COST_OF_SALES", "", "OUT_OF_SCOPE", "PL.COST_OF_SALES", "SE: costOfGoods", "Opening stock less closing stock"),
    ("LSP", "6000", "Direct labour", "EXPENSE", "COST_OF_SALES", "", "OUT_OF_SCOPE", "PL.COST_OF_SALES", "SE: wagesAndStaffCosts", ""),
    ("LSP", "6100", "Subcontractors", "EXPENSE", "COST_OF_SALES", "", "VATABLE", "PL.COST_OF_SALES", "SE: paymentsToSubcontractors", "Construction Industry Scheme payments"),
    ("LSP", "7000", "Wages and salaries", "EXPENSE", "OVERHEADS", "", "OUT_OF_SCOPE", "PL.ADMIN_EXPENSES", "SE: wagesAndStaffCosts", ""),
    ("L", "7001", "Directors' remuneration", "EXPENSE", "OVERHEADS", "", "OUT_OF_SCOPE", "PL.ADMIN_EXPENSES", "", "Disclosed in the accounts notes"),
    ("LSP", "7006", "Employer's National Insurance", "EXPENSE", "OVERHEADS", "", "OUT_OF_SCOPE", "PL.ADMIN_EXPENSES", "SE: wagesAndStaffCosts", ""),
    ("LSP", "7007", "Employer's pension contributions", "EXPENSE", "OVERHEADS", "", "OUT_OF_SCOPE", "PL.ADMIN_EXPENSES", "SE: wagesAndStaffCosts", ""),
    ("LSP", "7008", "Staff welfare and training", "EXPENSE", "OVERHEADS", "", "VATABLE", "PL.ADMIN_EXPENSES", "SE: wagesAndStaffCosts", ""),
    ("LSP", "7100", "Rent", "EXPENSE", "OVERHEADS", "", "VATABLE", "PL.ADMIN_EXPENSES", "SE: premisesRunningCosts", ""),
    ("LSP", "7103", "Business rates", "EXPENSE", "OVERHEADS", "", "OUT_OF_SCOPE", "PL.ADMIN_EXPENSES", "SE: premisesRunningCosts", ""),
    ("LSP", "7104", "Premises insurance", "EXPENSE", "OVERHEADS", "", "EXEMPT", "PL.ADMIN_EXPENSES", "SE: premisesRunningCosts", ""),
    ("SP", "7105", "Use of home as office", "EXPENSE", "OVERHEADS", "", "OUT_OF_SCOPE", "PL.ADMIN_EXPENSES", "SE: premisesRunningCosts", "Flat rate or an apportionment of household costs"),
    ("LSP", "7200", "Electricity", "EXPENSE", "OVERHEADS", "", "VATABLE", "PL.ADMIN_EXPENSES", "SE: premisesRunningCosts", ""),
    ("LSP", "7201", "Gas", "EXPENSE", "OVERHEADS", "", "VATABLE", "PL.ADMIN_EXPENSES", "SE: premisesRunningCosts", ""),
    ("LSP", "7202", "Water rates", "EXPENSE", "OVERHEADS", "", "OUT_OF_SCOPE", "PL.ADMIN_EXPENSES", "SE: premisesRunningCosts", ""),
    ("LSP", "7300", "Motor fuel", "EXPENSE", "OVERHEADS", "", "VATABLE", "PL.ADMIN_EXPENSES", "SE: carVanTravelExpenses", "Private use is disallowable for a sole trader"),
    ("LSP", "7301", "Motor repairs and servicing", "EXPENSE", "OVERHEADS", "", "VATABLE", "PL.ADMIN_EXPENSES", "SE: carVanTravelExpenses", ""),
    ("LSP", "7302", "Motor insurance and road tax", "EXPENSE", "OVERHEADS", "", "EXEMPT", "PL.ADMIN_EXPENSES", "SE: carVanTravelExpenses", ""),
    ("LSP", "7303", "Vehicle leasing", "EXPENSE", "OVERHEADS", "", "VATABLE", "PL.ADMIN_EXPENSES", "SE: carVanTravelExpenses", "Car lease restrictions apply for tax and VAT"),
    ("LSP", "7400", "Travel", "EXPENSE", "OVERHEADS", "", "VATABLE", "PL.ADMIN_EXPENSES", "SE: carVanTravelExpenses", "Wholly and exclusively (see Samadian)"),
    ("LSP", "7401", "Hotels and subsistence", "EXPENSE", "OVERHEADS", "", "VATABLE", "PL.ADMIN_EXPENSES", "SE: carVanTravelExpenses", ""),
    ("LSP", "7402", "Staff entertaining", "EXPENSE", "OVERHEADS", "", "VATABLE", "PL.ADMIN_EXPENSES", "SE: wagesAndStaffCosts", "Allowable when for staff only"),
    ("LSP", "7403", "Client entertaining", "EXPENSE", "OVERHEADS", "", "VATABLE", "PL.ADMIN_EXPENSES", "SE: businessEntertainmentCosts", "Disallowed for tax; input VAT blocked"),
    ("LSP", "7500", "Postage and courier", "EXPENSE", "OVERHEADS", "", "EXEMPT", "PL.ADMIN_EXPENSES", "SE: adminCosts", ""),
    ("LSP", "7501", "Stationery and printing", "EXPENSE", "OVERHEADS", "", "VATABLE", "PL.ADMIN_EXPENSES", "SE: adminCosts", ""),
    ("LSP", "7502", "Telephone and internet", "EXPENSE", "OVERHEADS", "", "VATABLE", "PL.ADMIN_EXPENSES", "SE: adminCosts", ""),
    ("LSP", "7503", "Software and subscriptions", "EXPENSE", "OVERHEADS", "", "VATABLE", "PL.ADMIN_EXPENSES", "SE: adminCosts", ""),
    ("LSP", "7504", "Advertising and marketing", "EXPENSE", "OVERHEADS", "", "VATABLE", "PL.ADMIN_EXPENSES", "SE: advertisingCosts", ""),
    ("LSP", "7505", "Website costs", "EXPENSE", "OVERHEADS", "", "VATABLE", "PL.ADMIN_EXPENSES", "SE: advertisingCosts", "Development may be capital"),
    ("LSP", "7600", "Legal fees", "EXPENSE", "OVERHEADS", "", "VATABLE", "PL.ADMIN_EXPENSES", "SE: professionalFees", "Fees on capital items are disallowable"),
    ("LSP", "7601", "Accountancy fees", "EXPENSE", "OVERHEADS", "", "VATABLE", "PL.ADMIN_EXPENSES", "SE: professionalFees", ""),
    ("LSP", "7602", "Consultancy fees", "EXPENSE", "OVERHEADS", "", "VATABLE", "PL.ADMIN_EXPENSES", "SE: professionalFees", ""),
    ("LSP", "7700", "Equipment hire", "EXPENSE", "OVERHEADS", "", "VATABLE", "PL.ADMIN_EXPENSES", "SE: otherExpenses", ""),
    ("LSP", "7800", "Repairs and maintenance", "EXPENSE", "OVERHEADS", "", "VATABLE", "PL.ADMIN_EXPENSES", "SE: maintenanceCosts", "Improvements are capital (see Odeon v Jones)"),
    ("LSP", "7900", "Bank interest paid", "EXPENSE", "INTEREST_EXPENSE", "", "OUT_OF_SCOPE", "PL.INTEREST_PAYABLE", "SE: interestOnBankOtherLoans", ""),
    ("LSP", "7901", "Bank charges", "EXPENSE", "OVERHEADS", "", "EXEMPT", "PL.ADMIN_EXPENSES", "SE: financeCharges", ""),
    ("LSP", "7902", "Card and payment processing fees", "EXPENSE", "OVERHEADS", "", "EXEMPT", "PL.ADMIN_EXPENSES", "SE: financeCharges", ""),
    ("LSP", "7903", "Loan interest", "EXPENSE", "INTEREST_EXPENSE", "", "OUT_OF_SCOPE", "PL.INTEREST_PAYABLE", "SE: interestOnBankOtherLoans", ""),
    ("LSP", "7904", "Hire purchase interest", "EXPENSE", "INTEREST_EXPENSE", "", "OUT_OF_SCOPE", "PL.INTEREST_PAYABLE", "SE: interestOnBankOtherLoans", ""),
    ("R", "7100", "Rent, rates, insurance and ground rents", "EXPENSE", "OVERHEADS", "", "EXEMPT", "PL.ADMIN_EXPENSES", "UP: premisesRunningCosts", ""),
    ("R", "7400", "Travel for the property business", "EXPENSE", "OVERHEADS", "", "VATABLE", "PL.ADMIN_EXPENSES", "UP: travelCosts", ""),
    ("R", "7600", "Legal and professional fees", "EXPENSE", "OVERHEADS", "", "VATABLE", "PL.ADMIN_EXPENSES", "UP: professionalFees", "Fees on buying a property are capital"),
    ("R", "7601", "Accountancy fees", "EXPENSE", "OVERHEADS", "", "VATABLE", "PL.ADMIN_EXPENSES", "UP: professionalFees", ""),
    ("R", "7603", "Letting agent and management fees", "EXPENSE", "OVERHEADS", "", "VATABLE", "PL.ADMIN_EXPENSES", "UP: professionalFees", ""),
    ("R", "7800", "Property repairs and maintenance", "EXPENSE", "OVERHEADS", "", "VATABLE", "PL.ADMIN_EXPENSES", "UP: repairsAndMaintenance", "Improvements are capital (see Law Shipping)"),
    ("R", "7810", "Replacement of domestic items", "EXPENSE", "OVERHEADS", "", "VATABLE", "PL.ADMIN_EXPENSES", "UP: annual adjustment (check)", "Relief for replacing furnishings in a dwelling; claimed annually, not quarterly"),
    ("R", "7820", "Costs of services provided, including wages", "EXPENSE", "OVERHEADS", "", "VATABLE", "PL.ADMIN_EXPENSES", "UP: costOfServices", ""),
    ("R", "7901", "Bank charges", "EXPENSE", "OVERHEADS", "", "EXEMPT", "PL.ADMIN_EXPENSES", "UP: other", ""),
    ("R", "7910", "Mortgage interest - residential property", "EXPENSE", "INTEREST_EXPENSE", "", "OUT_OF_SCOPE", "PL.INTEREST_PAYABLE", "UP: residentialFinancialCost", "Not deducted: relief is a basic-rate tax reduction"),
    ("R", "7911", "Finance costs - non-residential property", "EXPENSE", "INTEREST_EXPENSE", "", "OUT_OF_SCOPE", "PL.INTEREST_PAYABLE", "UP: financialCosts", ""),
    ("LSP", "8000", "Depreciation", "EXPENSE", "DEPRECIATION", "", "OUT_OF_SCOPE", "PL.DEPRECIATION", "SE: depreciation", "Added back for tax; capital allowances instead"),
    ("L", "8001", "Amortisation", "EXPENSE", "DEPRECIATION", "", "OUT_OF_SCOPE", "PL.DEPRECIATION", "", ""),
    ("LSP", "8002", "Profit or loss on disposal of fixed assets", "EXPENSE", "OTHER_EXPENSE", "", "OUT_OF_SCOPE", "PL.OTHER_EXPENSES", "SE: depreciation", "Disallowable; balancing allowances instead"),
    ("LSP", "8100", "Bad debts written off", "EXPENSE", "OVERHEADS", "", "OUT_OF_SCOPE", "PL.ADMIN_EXPENSES", "SE: irrecoverableDebts", "General provisions are disallowable"),
    ("R", "8100", "Bad debts written off", "EXPENSE", "OVERHEADS", "", "OUT_OF_SCOPE", "PL.ADMIN_EXPENSES", "UP: other", "Rent that will not be collected"),
    ("LSP", "8200", "Donations", "EXPENSE", "OTHER_EXPENSE", "", "OUT_OF_SCOPE", "PL.OTHER_EXPENSES", "SE: otherExpenses", "Usually disallowable; charity gifts relieved elsewhere"),
    ("LSP", "8201", "Trade subscriptions", "EXPENSE", "OVERHEADS", "", "EXEMPT", "PL.ADMIN_EXPENSES", "SE: otherExpenses", ""),
    ("LSP", "8204", "General insurance", "EXPENSE", "OVERHEADS", "", "EXEMPT", "PL.ADMIN_EXPENSES", "SE: otherExpenses", ""),
    ("LSP", "8205", "Sundry expenses", "EXPENSE", "OTHER_EXPENSE", "", "VATABLE", "PL.OTHER_EXPENSES", "SE: otherExpenses", ""),
    ("R", "8205", "Other allowable property expenses", "EXPENSE", "OTHER_EXPENSE", "", "VATABLE", "PL.OTHER_EXPENSES", "UP: other", ""),
    ("L", "8900", "Corporation tax charge", "EXPENSE", "TAX_EXPENSE", "", "OUT_OF_SCOPE", "PL.TAXATION", "", ""),
    ("LSPR", "9998", "Suspense", "ASSET", "CURRENT_ASSET_OTHER", "SUSPENSE", "NOT_APPLICABLE", "BS.CURRENT_ASSETS.DEBTORS", "", "System account; a balance raises an alert (DEC-044)"),
    ("LSPR", "9999", "Mispostings", "ASSET", "CURRENT_ASSET_OTHER", "", "NOT_APPLICABLE", "BS.CURRENT_ASSETS.DEBTORS", "", "Must be cleared before the year end"),
]

TEMPLATES = [("L", "Limited company"), ("S", "Sole trader"), ("P", "Partnership"), ("R", "Landlord")]
RANGES = [
    ("0010-0099", "Fixed assets and their depreciation", "Each asset class has the next code for its accumulated depreciation"),
    ("1000-1999", "Current assets", "Stock 1001; debtors 1100; bank and cash 1200-1299"),
    ("2000-2999", "Liabilities", "Creditors 2100; VAT 2200; payroll 2210-2240; long-term 2300-2399"),
    ("3000-3999", "Capital and reserves", "Shares or owners' and partners' accounts; profit and loss account 3200; drawings 3260"),
    ("4000-4999", "Sales and other income", "Sales 4000-4099; other income 4900-4999"),
    ("5000-5999", "Purchases and direct costs", ""),
    ("6000-6999", "Direct expenses", "Direct labour, subcontractors"),
    ("7000-7999", "Overheads", "Staff 7000; premises 7100-7299; motor and travel 7300-7499; office 7500; professional 7600; repairs 7800; finance 7900"),
    ("8000-8999", "Depreciation and other expenses", "Corporation tax charge 8900"),
    ("9000-9999", "Suspense and mispostings", "9998 suspense; 9999 mispostings"),
]
CHANGES = [
    ("Chart template version", "DEFAULT_CHART_VERSION 2: one template per business type, chosen when the chart is initialised. Version 1 stays for every existing company and for the V2 golden workings, which read it."),
    ("Legal form", "Templates follow the company's legal form (LTD, SOLE_TRADER, PARTNERSHIP). A landlord needs a form the platform lacks (for example INDIVIDUAL); LLP and CHARITY need members' interests and fund templates and are not in this draft."),
    ("Equity subtypes (NEW)", "OWNER_CAPITAL, DRAWINGS, PARTNER_CAPITAL, PARTNER_CURRENT."),
    ("Reporting lines (NEW)", "BS.CAPITAL.PROPRIETOR (capital account) and BS.CAPITAL.PARTNERS (partners' accounts). The reporting-line registry is versioned, so this is REPORT_MAPPING_VERSION 2, never an edit of version 1."),
    ("HMRC category", "A new account attribute: the self-employment or UK property category of HMRC's MTD APIs, used by V13 to build the quarterly updates. Verified through tax:rule-verify before use."),
    ("One VAT control account", "2200 only. Sage charts split VAT into 2200, 2201 and 2202; the platform's VAT engine posts to one control account and reports the boxes from the VAT facts."),
    ("Product wording", "Screens call these 'UK standard nominal codes'. 'Sage' is a trade mark and is not used in the product."),
]
LAW = [
    ("Income Tax (Trading and Other Income) Act 2005, s25 and s25A", "Accounts on generally accepted practice; the cash basis", "https://www.legislation.gov.uk/ukpga/2005/5/section/25A"),
    ("Income Tax (Trading and Other Income) Act 2005, s33 and s34", "No deduction for capital; wholly and exclusively", "https://www.legislation.gov.uk/ukpga/2005/5/section/34"),
    ("Income Tax (Trading and Other Income) Act 2005, s45", "Business entertainment disallowed", "https://www.legislation.gov.uk/ukpga/2005/5/section/45"),
    ("Income Tax (Trading and Other Income) Act 2005, s272A", "Residential property finance costs: no deduction", "https://www.legislation.gov.uk/ukpga/2005/5/section/272A"),
    ("Income Tax (Trading and Other Income) Act 2005, s311A", "Replacement of domestic items relief", "https://www.legislation.gov.uk/ukpga/2005/5/section/311A"),
    ("Corporation Tax Act 2009, s46, s53 and s54", "Company trading profits: accounting basis, capital, wholly and exclusively", "https://www.legislation.gov.uk/ukpga/2009/4/section/54"),
    ("Corporation Tax Act 2009, s1298", "Business entertainment disallowed for companies", "https://www.legislation.gov.uk/ukpga/2009/4/section/1298"),
    ("Capital Allowances Act 2001, s51A", "Annual investment allowance in place of depreciation", "https://www.legislation.gov.uk/ukpga/2001/2/section/51A"),
    ("Value Added Tax (Input Tax) Order 1992, SI 1992/3222, art 5", "Input VAT on business entertainment blocked", "https://www.legislation.gov.uk/uksi/1992/3222/article/5"),
    ("Companies Act 2006, s386", "Adequate accounting records", "https://www.legislation.gov.uk/ukpga/2006/46/section/386"),
    ("Mallalieu v Drummond [1983] 2 AC 861 (HL)", "Duality of purpose defeats 'wholly and exclusively'", "https://scholar.google.com/scholar_case?q=Mallalieu+v+Drummond"),
    ("Samadian v HMRC [2014] UKUT 13 (TCC)", "A sole trader's travel: only the business part is allowable", "https://www.bailii.org/uk/cases/UKUT/TCC/2014/13.html"),
    ("Odeon Associated Theatres Ltd v Jones [1973] Ch 288 (CA)", "Repairs are revenue when they restore, not improve", "https://scholar.google.com/scholar_case?q=Odeon+Associated+Theatres+v+Jones"),
    ("Law Shipping Co Ltd v IRC (1923) 12 TC 621", "Repairs to make a newly bought asset usable are capital", "https://scholar.google.com/scholar_case?q=Law+Shipping+v+IRC"),
    ("Interfish Ltd v HMRC [2014] EWCA Civ 876", "Payments with a non-trade purpose fail 'wholly and exclusively'", "https://www.bailii.org/ew/cases/EWCA/Civ/2014/876.html"),
]

INK, PENCIL = "172231", "566170"
HEAD = PatternFill("solid", fgColor=INK)
FONT = Font(name="Arial", size=10, color=INK)
BOLD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
TITLE = Font(name="Arial", size=14, bold=True, color=INK)
NOTE = Font(name="Arial", size=9, italic=True, color=PENCIL)
NEW = PatternFill("solid", fgColor="FFFBEB")
RULE = Border(bottom=Side(style="thin", color="E1E6EC"))
WRAP = Alignment(wrap_text=True, vertical="top")


def table(ws, title, note, headers, rows, widths, start=4):
    ws["A1"], ws["A1"].font = title, TITLE
    ws["A2"], ws["A2"].font = note, NOTE
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=start, column=c, value=h)
        cell.font, cell.fill, cell.alignment = BOLD, HEAD, WRAP
    for r, row in enumerate(rows, start + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font, cell.alignment, cell.border = FONT, WRAP, RULE
            if isinstance(v, str) and "(NEW)" in v:
                cell.fill = NEW
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=start + 1, column=1)
    ws.auto_filter.ref = f"A{start}:{get_column_letter(len(headers))}{start + len(rows)}"


def box(category):
    """SA103F or SA105 box for an 'SE: x' or 'UP: x' category, UNVERIFIED."""
    kind, _, name = category.partition(": ")
    ref = (SE if kind == "SE" else UP if kind == "UP" else {}).get(name)
    return f"{'SA103F' if kind == 'SE' else 'SA105'} box {ref}" if ref else ""


wb = Workbook()
ws = wb.active
ws.title = "Read me"
readme = [
    ("What this is", "The draft chart-of-accounts templates of DEC-080: Sage-style four-digit nominal codes, one template per business type. Nothing is built yet."),
    ("What to do", "On each template sheet, put Approve, Change or Remove in 'Your verdict' and say what to change in 'Your comment'. Add any account you want in an empty row."),
    ("What is unverified", "Every HMRC category and SA103F/SA105 box is UNVERIFIED (DEC-022): HMRC's documents could not be read in this environment. They are verified in the platform with tax:rule-verify before V13 uses them."),
    ("Yellow cells", "Subtypes and reporting lines marked (NEW) do not exist in the platform yet; see 'Platform changes'."),
    ("Codes", "Sage 50 layout of ranges (see 'Code ranges'). The platform keeps one VAT control account, 2200."),
    ("Not legal advice", "The 'Legal references' sheet lists the provisions and cases behind the tax notes. Links were not opened from this environment; check each one."),
]
table(ws, "Chart templates version 2 (DEC-080)", "Generated by scripts/docs/gen_chart_templates_v2.py. Draft for the product owner's review.",
      ["Item", "Detail"], readme, [24, 120])
table(wb.create_sheet("Code ranges"), "Code ranges", "Sage 50 layout, used by every template.", ["Range", "Holds", "Notes"], RANGES, [14, 36, 90])

for key, name in TEMPLATES:
    # A company's figures go to the CT computation (V4), not to an HMRC income category.
    cat = (lambda r: "") if key == "L" else (lambda r: r[8])
    rows = sorted(((r[1], r[2], r[3], r[4], r[5], r[6], r[7], cat(r), box(cat(r)), r[9], "", "") for r in ROWS if key in r[0]), key=lambda x: x[0])
    codes = [r[0] for r in rows]
    assert len(codes) == len(set(codes)), f"duplicate code in {name}"
    s = wb.create_sheet(name)
    hmrc = "HMRC category (MTD)" if key != "L" else "HMRC category (none: CT computation, V4)"
    table(s, f"{name} template", f"{len(rows)} accounts. HMRC categories and boxes are UNVERIFIED.",
          ["Code", "Name", "Type", "Subtype", "Control", "VAT treatment", "Reporting line", hmrc, "SA103F / SA105 box (UNVERIFIED)", "Note", "Your verdict", "Your comment"],
          rows, [8, 38, 11, 28, 20, 16, 30, 30, 22, 50, 14, 40])
    dv = DataValidation(type="list", formula1='"Approve,Change,Remove"', allow_blank=True)
    s.add_data_validation(dv)
    dv.add(f"K5:K{4 + len(rows) + 20}")

table(wb.create_sheet("Platform changes"), "What the platform needs before these templates can be used",
      "Recorded in DEC-080; built only after the owner's review.", ["Change", "Detail"], CHANGES, [26, 120])
law = wb.create_sheet("Legal references")
table(law, "Legal references", "Not checked from this environment and not legal advice. Older cases link to a search.",
      ["Provision or case", "Why it matters for the chart", "Link"], LAW, [52, 60, 70])
for r in range(5, 5 + len(LAW)):
    law.cell(row=r, column=3).hyperlink = law.cell(row=r, column=3).value
    law.cell(row=r, column=3).font = Font(name="Arial", size=10, color="1D5FBF", underline="single")

wb.save(OUT)
print(OUT)
